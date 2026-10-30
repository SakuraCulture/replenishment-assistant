"""Repeatable isolated forms, original-source, return versioning and compatibility checks."""
import argparse,copy,hashlib,json,sqlite3,sys,threading
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path[:0]=[str(ROOT/'src'),str(ROOT/'tools')]
import work_repository as wr,normal_orders as no,normal_exports as ne,merchant_feedback as mf
import history_forms as hf,source_archive as sources,import_templates,storage_maintenance as sm
import openpyxl
from parallel_qa import diff

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--fixture',type=Path,required=True);ap.add_argument('--output',type=Path,required=True);a=ap.parse_args();out=a.output.resolve();out.mkdir(parents=True,exist_ok=False);checks=[]
    p=wr.loads(a.fixture.read_text(encoding='utf-8'));p['day']=wr.today()
    import warehouse_catalog
    p['catalogs']={wh:warehouse_catalog.import_catalog(a.fixture.resolve().parent/(wh+'.xlsx'),wh) for wh in ['河北仓','山东仓']}
    work=wr.Repository(out/'settings');repo=no.Repository(work);codes=list(p['results']);repo.confirm(p,codes,None,{s:None for s in codes});original=[repo.get(p['day'],s) for s in codes];saved=work.get_job(p['day']);heads=[{k:v for k,v in r.items() if k!='payload'} for r in original]
    def unchanged():
        assert not diff(original,[repo.get(p['day'],s,1) for s in codes]);assert not diff(saved,work.get_job(p['day']))
    def excel(path,rows):
        b=openpyxl.Workbook();w=b.active;w.append(['补货单号','商品编码','门店确认数量','备注']);[w.append(r) for r in rows];b.save(path);b.close();return path
    h=heads[0];row=next(d for d in original[0]['payload']['result']['rows'] if d['E']=='new_must');sku=row['E'];num=h['number']
    book=out/'returned.xlsx';ne.export(original[0],book)
    b=openpyxl.load_workbook(book)
    for ws in b:
        for cells in list(ws.rows)[1:]:
            if cells[4].value==sku:cells[16].value=5
    b.save(book);b.close();scanned=mf.scan([book]);plan=mf.prepare(work.path,scanned,{});assert plan['plans'][0]['warnings'];receipt=mf.commit(work.path,plan);assert receipt[0]['feedback_version']==1;unchanged();checks.append('exported workbook custom revision auto binds; non-box quantity 5 retained with warning, original issued 6 untouched')
    again=mf.prepare(work.path,mf.scan([book]),{});assert not again['plans'][0]['changed'];assert mf.commit(work.path,again)[0]['feedback_version']==1;unchanged();checks.append('same returned values idempotent: no duplicate feedback version')
    blank=excel(out/'blank.xlsx',[[num,sku,None,'']]);scan=mf.scan([blank])
    try:mf.prepare(work.path,scan,{})
    except ValueError:pass
    else:raise AssertionError('missing revision silently inferred')
    plan=mf.prepare(work.path,scan,{num:1});assert not plan['plans'][0]['changed'] and plan['plans'][0]['payload']['values'][sku]==5;checks.append('missing revision requires explicit binding; blank preserves prior 5')
    zero=excel(out/'zero.xlsx',[[num,sku,0,'拒绝']]);plan=mf.prepare(work.path,mf.scan([zero]),{num:1});assert mf.commit(work.path,plan)[0]['feedback_version']==2;unchanged();checks.append('explicit zero saved as rejection in version2; signed quantities unchanged')
    dup=excel(out/'duplicate.xlsx',[[num,sku,1,''],[num,sku,1,'']]);plan=mf.prepare(work.path,mf.scan([dup]),{num:1});assert plan['plans'][0]['duplicates']==1 and len(plan['plans'][0]['payload']['incoming'][sku]['sources'])==2
    conflicting=excel(out/'conflict.xlsx',[[num,sku,1,''],[num,sku,2,'']])
    try:mf.prepare(work.path,mf.scan([conflicting]),{num:1})
    except ValueError:pass
    else:raise AssertionError('duplicate conflicts merged')
    unknown=excel(out/'unknown.xlsx',[[num,'unknown',1,'']])
    try:mf.prepare(work.path,mf.scan([unknown]),{num:1})
    except ValueError:pass
    else:raise AssertionError('unknown SKU inserted')
    checks.append('identical duplicates retain sources; conflicting duplicates and unknown SKU blocked')
    for index,bad in enumerate([-1,'=1+1','bad',True]):
        f=excel(out/f'invalid-{index}.xlsx',[[num,sku,bad,'']])
        try:mf.scan([f])
        except ValueError:pass
        else:raise AssertionError('invalid quantity accepted')
    checks.append('negative, formula, text and boolean return quantities rejected')

    # Competing reply writer invalidates the first preview.
    stale=mf.prepare(work.path,mf.scan([dup]),{num:1});mf.commit(work.path,stale)
    try:mf.commit(work.path,stale)
    except wr.Conflict:pass
    else:raise AssertionError('stale feedback committed')
    checks.append('feedback optimistic concurrency rejects stale preview')
    # All-order transaction rollback on second insert.
    two=excel(out/'two.xlsx',[[heads[0]['number'],'new_must',2,''],[heads[1]['number'],'moving',2,'']]);plan=mf.prepare(work.path,mf.scan([two]),{x['number']:1 for x in heads})
    with work.connect() as db:
        before=list(db.execute('SELECT * FROM merchant_feedback ORDER BY rowid'));db.execute("CREATE TRIGGER qa_fail_reply BEFORE INSERT ON merchant_feedback WHEN NEW.oid='"+heads[1]['oid']+"' BEGIN SELECT RAISE(ABORT,'injected rollback'); END")
    try:mf.commit(work.path,plan)
    except sqlite3.IntegrityError:pass
    else:raise AssertionError('rollback failed')
    with work.connect() as db:assert list(db.execute('SELECT * FROM merchant_feedback ORDER BY rowid'))==before;db.execute('DROP TRIGGER qa_fail_reply')
    receipts=mf.commit(work.path,plan);assert [x['number'] for x in receipts]==[x['number'] for x in heads];unchanged();checks.append('multi-order reply atomic rollback and retry; selected order retained')
    # Original files are byte-exact and independent from working paths.
    catalog=next(iter(p['catalogs'].values()));file=Path(catalog['source_path']);sha=catalog['source_sha256'];sources.capture(work.path,file,sha);sources.capture(work.path,file,sha)
    with work.connect() as db:
        assert db.execute('SELECT COUNT(*) FROM source_files WHERE sha=?',(sha,)).fetchone()[0]==1;sources.export(db,sha,out/'catalog-original.xlsx')
    assert hashlib.sha256((out/'catalog-original.xlsx').read_bytes()).hexdigest()==sha
    for scope in ['单据概况','当时参数','商品建议与销量库存','商品完整计算依据','当时商品资料快照']:
        form=hf.form(repo,original[0],scope,sku);assert form and all(len(r)==3 for r in form)
    snapshots=hf.form(repo,original[0],'当时商品资料快照',sku);assert any('原表全部列' in x[0] for x in snapshots)
    assert any('历史原文件未留存' in x[1] for x in snapshots);checks.append('business forms cover historical parameters, quantities, raw/effective stockout, catalog; byte-exact source dedup; uncaptured warehouse explicitly missing')
    templates=[]
    for kind in import_templates.TEMPLATES:
        path=out/(kind+'-template.xlsx');import_templates.write(kind,path);b=openpyxl.load_workbook(path,read_only=True);title,sheet,req,opt=import_templates.TEMPLATES[kind];assert b.sheetnames==[sheet,'填写说明'];assert list(next(b[sheet].values))==req+opt;b.close();templates.append(kind)
    checks.append('all six minimum templates: required headers, empty data area, instructions')
    import business_excel
    from runtime_support import Cancelled
    from openpyxl.worksheet._writer import ALL_TEMP_FILES
    stop=threading.Event();temp_before=set(ALL_TEMP_FILES)
    try:business_excel.write(out/'cancelled.xlsx',[('测试',['值'],([i] for i in range(2000)))],stop,lambda _:stop.set())
    except Cancelled:pass
    else:raise AssertionError('business export cancellation ignored')
    assert not (out/'cancelled.xlsx').exists() and set(ALL_TEMP_FILES)==temp_before
    checks.append('cancel new business Excel cleans its own XML temporary files and publishes no partial workbook')

    # Legacy conversion preserves added independent tables and exact original snapshots.
    for mode in ['legacy','upgrade']:
        target=out/(mode+'.sqlite3');sm.convert(work.path,target,mode)
        with sqlite3.connect(target) as db,work.connect() as src:
            for table in ['source_files','merchant_feedback']:assert list(db.execute('SELECT * FROM '+table+' ORDER BY rowid'))==list(src.execute('SELECT * FROM '+table+' ORDER BY rowid'))
    unchanged();checks.append('legacy/upgrade copy preserves returned versions and original file BLOBs, original signed payload exact')
    # Signing revision race rejects feedback plan, historical version remains independently readable.
    stale=mf.prepare(work.path,mf.scan([zero]),{num:1});changed=copy.deepcopy(p);changed['results'][codes[0]]['qa_note']='new signed revision';repo.confirm(changed,[codes[0]],saved['revision'],{codes[0]:1})
    try:mf.commit(work.path,stale)
    except wr.Conflict:pass
    else:raise AssertionError('signed revision changed during reply preview')
    assert mf.view(work.path,heads[0])['record'];h2=repo.head(p['day'],codes[0]);assert mf.view(work.path,h2)['record'] is None;checks.append('new signing invalidates reply preview; V1 return never auto-applies to V2')
    report={'passed':True,'checks':checks,'templates':templates,'business_payload_differences':0,'new_tables':['source_files','merchant_feedback'],'scope':'return and snapshot forms use isolated QA database; no user paths written'};(out/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8',newline='\n');print(json.dumps(report,ensure_ascii=False))
if __name__=='__main__':main()
