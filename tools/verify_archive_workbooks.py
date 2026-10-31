"""Real frozen-result archive export check; independent ordered field/cell oracle."""
import argparse,hashlib,json,sys,time
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path[:0]=[str(ROOT/'src'),str(ROOT/'tools')]
from parallel_qa import read,js
import work_repository as wr,normal_orders as no,normal_exports as ne,archive_exports as ae,archive_orders as ao
import result_semantics as sem,new_product_focus as focus,openpyxl
from version_info import MODEL_VERSION
from contextlib import closing


def leaves(value,path=()):
    if isinstance(value,dict):
        if not value:yield path,'空记录'
        for k,v in value.items():yield from leaves(v,path+(str(k),))
    elif isinstance(value,(list,tuple)):
        if not value:yield path,'空清单'
        for i,v in enumerate(value):yield from leaves(v,path+(str(i+1),))
    else:yield path,value


def main():
    ap=argparse.ArgumentParser();ap.add_argument('--results',type=Path,required=True);ap.add_argument('--output',type=Path,required=True);a=ap.parse_args();out=a.output.resolve();out.mkdir(parents=True,exist_ok=False);report=[]
    for name in ('real_None_True','real_3000_True'):
        folder=out/name;folder.mkdir();r=read(a.results/(name+'.pklz'));work=wr.Repository(folder/'settings');repo=no.Repository(work)
        p={'schema':1,'model_version':MODEL_VERSION,'day':wr.today(),'input_dirty':False,'states':{r['store']:{'name':r['store_name'],'state':'已完成'}},'results':{r['store']:r},'new_focus':focus.empty(wr.today()[:7])}
        order=repo.confirm(p,[r['store']],None,{r['store']:None})[0];before=wr.dumps(order['payload'])
        for wh in ao.scopes(r['rows']):
            h={k:v for k,v in order.items() if k!='payload'};h.update(warehouse=wh,number=ao.number(order['number'],wh));target=folder/(wh+'.xlsx');start=time.perf_counter();ae.export(work.path,h,target);elapsed=time.perf_counter()-start
            b=openpyxl.load_workbook(target,read_only=True);assert b.sheetnames==ae.SHEETS
            it=b.worksheets[0].iter_rows();headers=[c.value for c in next(it)];actual=list(it);expected=[d for d in r['rows'] if ao.warehouse(d)==wh];assert len(actual)==len(expected);checked=0
            for cells,d in zip(actual,expected):
                for path,v in leaves(d):
                    title=ae.label(path);c=cells[headers.index(title)]
                    # Excel numeric serialization has its documented 15/16 digit
                    # representation, not a relaxed business comparison tolerance.
                    if isinstance(v,float):
                        from openpyxl.compat.strings import safe_string
                        want=float(safe_string(v));assert c.value==want,(name,wh,d['E'],path,c.value,want)
                    else:assert c.value==(None if v=='' else v),(name,wh,d['E'],path,c.value,v)
                    if isinstance(v,str) and v:assert c.data_type=='s'
                    if isinstance(v,bool):assert c.data_type=='b'
                    checked+=1
                assert cells[-3].value==sem.order_quantity(d)
            assert '未留存' in next(b.worksheets[1].iter_rows(min_row=2,values_only=True))[0] # deliberately legacy mixed fixture, no invented warehouse source
            signed=list(b.worksheets[2].values);want=[]
            for group,items in ne.groups(r,p['new_focus']).items():
                for d in items:
                    if ao.warehouse(d)==wh:want.append((group,d['E'],sem.order_quantity(d) if sem.orderable(d) else 0))
            assert [(x[0],x[5],x[17]) for x in signed[1:]]==want
            b.close();assert wr.dumps(repo.get(p['day'],r['store'])['payload'])==before
            report.append({'case':name,'warehouse':wh,'rows':len(expected),'verified_input_intermediate_audit_cells':checked,'signed_rows':len(want),'seconds':elapsed,'differences':0,'sha256':hashlib.sha256(target.read_bytes()).hexdigest()})
    js(out/'report.json',{'passed':True,'cases':report,'scope':'all existing row leaves incl empty structures, types, exact row order, derived execution and signed group rows; archived source absent explicitly','numeric_excel_encoding':'openpyxl safe_string, not widened float tolerance','snapshot_changes':0})
    print('passed',len(report),'warehouse archives')

if __name__=='__main__':main()
