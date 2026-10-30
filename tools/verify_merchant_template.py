"""Independent read-only contract comparison against the user supplied workbook."""
import argparse,hashlib,json
from pathlib import Path
import openpyxl


def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()


def compare(template,candidate):
    before=sha(template);old=openpyxl.load_workbook(template);new=openpyxl.load_workbook(candidate);diffs=[]
    def check(label,a,b):
        if a!=b:diffs.append({'field':label,'template':str(a),'candidate':str(b)})
    check('sheet order',old.sheetnames,new.sheetnames)
    for name in old.sheetnames:
        a,b=old[name],new[name]
        check(name+'.headers',[c.value for c in a[1][:20]],[c.value for c in b[1][:20]])
        for field in ['baseColWidth','defaultColWidth','defaultRowHeight']:check(name+'.'+field,getattr(a.sheet_format,field),getattr(b.sheet_format,field))
        for field in ['freeze_panes']:check(name+'.'+field,getattr(a,field),getattr(b,field))
        check(name+'.filter',a.auto_filter.ref,b.auto_filter.ref)
        check(name+'.merged',str(a.merged_cells),str(b.merged_cells))
        for row in b.iter_rows(max_col=20):
            for cell in row:
                source=a.cell(1 if cell.row==1 else 2,cell.column)
                for field in ['font','fill','alignment','border','number_format','protection']:
                    check(name+'.'+cell.coordinate+'.'+field,str(getattr(source,field)),str(getattr(cell,field)))
    old.close();new.close();check('source unchanged',before,sha(template))
    return {'passed':not diffs,'template':str(template),'template_sha256':before,'candidate':str(candidate),'candidate_sha256':sha(candidate),'scope':'Exact 3-sheet order, 20 visible headers, all exported visible cell font/fill/alignment/border/format/protection, default row/column dimensions, freeze/filter/merge. Values intentionally come from new signed result, never template records/formulas. Hidden audit columns, comments, print identity are declared additions.','differences':diffs}


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--template',type=Path,required=True);p.add_argument('--candidate',type=Path,required=True);p.add_argument('--report',type=Path,required=True);a=p.parse_args();r=compare(a.template,a.candidate);a.report.write_text(json.dumps(r,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps({'passed':r['passed'],'differences':len(r['differences'])}));assert r['passed']
