"""Frozen 0.9.2 export oracle + warehouse replies/transaction integration. Private QA only."""
import argparse,copy,hashlib,importlib.util,json,pickle,sys,threading,zlib
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'src'))
import openpyxl
import work_repository as wr,normal_orders as no,normal_delivery as nd,normal_exports as ne,warehouse_orders as wo,merchant_feedback as mf,new_product_focus as focus
from runtime_support import Cancelled

def read(p):return pickle.loads(zlib.decompress(p.read_bytes()))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def cells(ws):
    return [[(c.value,c.data_type,str(c.font),str(c.fill),str(c.border),str(c.alignment),c.number_format,str(c.protection),c.comment.text if c.comment else None) for c in row] for row in ws]

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--baseline',type=Path,required=True);ap.add_argument('--fixture',type=Path,required=True);ap.add_argument('--results',type=Path,required=True);ap.add_argument('--output',type=Path,required=True);a=ap.parse_args();out=a.output.resolve();out.mkdir(parents=True,exist_ok=False);checks=[];comparisons=[]
    spec=importlib.util.spec_from_file_location('frozen092exports',a.baseline/'src/normal_exports.py');old=importlib.util.module_from_spec(spec);spec.loader.exec_module(old)
    # Dependencies shared by the oracle are byte-identical, so only exporter differs.
    deps=['model.py','one_box_exception.py','allocation.py','default_config.json','default_must_sell.json','result_semantics.py','new_products.py','new_product_focus.py','export_progress.py']
    assert all(sha(ROOT/'src'/n)==sha(a.baseline/'src'/n) for n in deps)
    p=wr.loads(a.fixture.read_bytes());p['day']=wr.today();original=copy.deepcopy(p)
    # Explicit synthetic routing fixture: alternate eligible SKUs, not a real routing claim.
    for r in p['results'].values():
        eligible=[d for rows in ne.groups(r,p.get('new_focus')).values() for d in rows]
        for i,d in enumerate(eligible):d['AU']='河北仓' if i%2==0 else '山东仓'
    work=wr.Repository(out/'settings');repo=no.Repository(work);batch=nd.prepare_batch(repo,p,list(p['results']));folder=out/'batch';folder.mkdir();receipt=nd.execute_batch(repo,batch,folder)
    orders=[repo.get(p['day'],s) for s in p['results']];signed=[wr.dumps(o['payload']) for o in orders];job=wr.dumps(work.get_job(p['day']));checks.append('synthetic two-store split: one atomic parent per store; per-store ordered files')
    assert len(receipt['paths'])==sum(len(wo.partitions(o['payload']['result'],o['payload'].get('new_focus'))) for o in orders)
    real=[]
    for f in sorted(a.results.glob('real_*.pklz')):
        r=read(f)
        if 'rows' in r:real.append((f.stem,{'store':r['store'],'day':p['day'],'number':'BH-'+p['day'].replace('-','')+'-'+r['store']+'-ZC','revision':1,'current_revision':1,'payload':{'result':r,'new_focus':focus.empty(p['day'][:7])}}))
    for label,o in [(f'synthetic_{i}',o) for i,o in enumerate(orders)]+real:
        d=out/label;d.mkdir();before=wr.dumps(o);old.export(o,d/'old.xlsx');b=openpyxl.load_workbook(d/'old.xlsx');expected={ws.title:cells(ws) for ws in b};whs=wo.partitions(o['payload']['result'],o['payload'].get('new_focus'));counts=0;qty=amount=0;keys=[]
        for wh in whs:
            target=d/ne.filename(o,wh);ne.export(o,target,warehouse=wh);n=openpyxl.load_workbook(target);assert n.sheetnames==b.sheetnames;props={x.name:x.value for x in n.custom_doc_props};assert props['补货单号']==wo.child_number(o['number'],wh) and props['门店正常单号']==o['number'] and props['修订']=='V001'
            for ws in n:
                actual=cells(ws);want=[expected[ws.title][0]]+[row for row in expected[ws.title][1:] if row[18][0]==wh]
                assert len(actual)==len(want),(label,wh,ws.title)
                for rn,(x,y) in enumerate(zip(actual,want)):
                    y=list(y)
                    if rn==0:
                        # Header comment contains child number instead of parent.
                        t=list(y[0]);t[-1]=t[-1].replace(o['number'],wo.child_number(o['number'],wh));y[0]=tuple(t)
                    else:
                        t=list(y[20]);t[0]=wo.child_number(o['number'],wh);y[20]=tuple(t);keys.append((ws.title,x[4][0]));qty+=x[16][0];amount+=x[16][0]*x[15][0] if isinstance(x[15][0],(int,float)) else 0
                    assert x==y,(label,wh,ws.title,rn);counts+=len(x)
                assert ws.sheet_format==b[ws.title].sheet_format
                assert {k:str(v) for k,v in ws.column_dimensions.items()}=={k:str(v) for k,v in b[ws.title].column_dimensions.items()}
                assert ws.oddHeader.center.text==b[ws.title].oddHeader.center.text.replace(o['number'],wo.child_number(o['number'],wh))
            n.close()
        allrows=[row for rows in expected.values() for row in rows[1:]]
        assert len(keys)==len(set(keys))==len(allrows);assert qty==sum(row[16][0] for row in allrows);assert round(amount,6)==round(sum(row[16][0]*row[15][0] for row in allrows if isinstance(row[15][0],(int,float))),6)
        comparisons.append({'case':label,'files':len(whs),'rows':len(allrows),'cells':counts,'quantity':qty,'amount':round(amount,2),'differences':0});b.close();assert wr.dumps(o)==before
    checks.append('independent frozen exporter: all cell values/types/styles/comments, sheet/column/SKU order; exact row union, quantities and money')
    # Both warehouse return files in one call create one feedback version, not colliding plans.
    o=orders[0];paths=[Path(f) for f in receipt['paths'] if o['payload']['result']['store_name'] in Path(f).name]
    assert len(paths)==2
    scan=mf.scan(paths);opts=mf.choices(work.path,list(scan['orders']));assert all(opts.values());plan=mf.prepare(work.path,scan,{});assert len(plan['plans'])==1;mf.commit(work.path,plan)
    record=mf.view(work.path,{k:v for k,v in o.items() if k!='payload'})['record'];assert len(record['payload']['warehouses'])==2
    again=mf.prepare(work.path,mf.scan(paths),{});assert not again['plans'][0]['changed'];checks.append('two child returns merge atomically into one version; same return is idempotent')
    target=paths[0];wb=openpyxl.load_workbook(target);changed=None
    for ws in wb:
        if ws.max_row>1:changed=ws.cell(2,5).value;ws.cell(2,17).value=0;break
    wb.save(out/'one-return.xlsx');wb.close();plan=mf.prepare(work.path,mf.scan([out/'one-return.xlsx']),{});mf.commit(work.path,plan);values=plan['plans'][0]['payload']['values'];assert values[changed]==0
    for sku,value in record['payload']['values'].items():
        if sku!=changed:assert values[sku]==value
    checks.append('one warehouse return changes only its SKU; other warehouse replies retained; signed values immutable')
    # Swap a SKU from the other warehouse: rejected even when it is in the parent.
    eligible=[d for rows in ne.groups(o['payload']['result'],o['payload'].get('new_focus')).values() for d in rows];wh=next(d['AU'] for d in eligible if d['E']==changed);wrong=next(d['E'] for d in eligible if d['AU']!=wh)
    def book(name,number,sku,q):
        b=openpyxl.Workbook();w=b.active;w.append(['补货单号','商品编码','门店确认数量']);w.append([number,sku,q]);b.save(out/name);b.close();return out/name
    f=book('wrong.xlsx',wo.child_number(o['number'],wh),wrong,1)
    try:mf.prepare(work.path,mf.scan([f]),{wo.child_number(o['number'],wh):1})
    except ValueError as exc:assert '不属于' in str(exc)
    else:raise AssertionError('wrong warehouse accepted')
    legacy=book('legacy.xlsx',o['number'],wrong,None);plan=mf.prepare(work.path,mf.scan([legacy]),{o['number']:1});assert not plan['plans'][0]['changed'];checks.append('cross-warehouse injection rejected; legacy parent-number return still binds exact revision; blank preserves previous values')
    # Stale reply preview rejected; current revision never substitutes old export.
    f=book('race.xlsx',wo.child_number(o['number'],wh),changed,1);stale=mf.prepare(work.path,mf.scan([f]),{wo.child_number(o['number'],wh):1});mf.commit(work.path,stale)
    try:mf.commit(work.path,stale)
    except wr.Conflict:pass
    else:raise AssertionError('stale reply accepted')
    for i,order in enumerate(orders):assert wr.dumps(repo.get(p['day'],order['store'],1)['payload'])==signed[i]
    assert wr.dumps(work.get_job(p['day']))==job
    # Cancellation before and after signing, exact revision retry and occupied filename protection.
    stop=threading.Event();stop.set();dest=out/'cancel';dest.mkdir();fresh=wr.Repository(out/'cancel-settings');fr=no.Repository(fresh)
    try:nd.prepare_batch(fr,p,list(p['results']),cancel=stop)
    except Cancelled:pass
    else:raise AssertionError('cancel ignored')
    assert not fr.headers();stop.clear();batch=nd.prepare_batch(fr,p,list(p['results']))
    def progress(message):
        if str(message).startswith('文件 1/'):stop.set()
    try:nd.execute_batch(fr,batch,dest,progress=progress,cancel=stop)
    except RuntimeError as exc:assert '归档' in str(exc) and '未生成' in str(exc)
    else:raise AssertionError('postcommit cancel ignored')
    assert len(fr.headers())==len(p['results']) and not list(dest.glob('*.xlsx'));stop.clear()
    batch=nd.prepare_batch(fr,p,list(p['results']));r=nd.execute_batch(fr,batch,dest,cancel=stop);assert not r['signed'] and all(h['revision']==1 for h in fr.headers())
    try:nd.execute_batch(fr,batch,dest)
    except FileExistsError:pass
    else:raise AssertionError('overwrite accepted')
    checks.append('pre-sign cancel, post-sign cancel receipt, retry same V001, existing files rejected before signing')
    # Immutable old revision re-export and subset warehouse selection.
    changed_payload=copy.deepcopy(p);changed_payload['results'][o['store']]['qa_note']='new revision';repo.confirm(changed_payload,[o['store']],nd.job_head(work,p['day'])['revision'],{o['store']:1})
    assert repo.head(p['day'],o['store'])['revision']==2
    hdir=out/'historical';hdir.mkdir();nd.export_signed(repo,o,hdir,[wh]);assert wr.dumps(repo.get(p['day'],o['store'],1)['payload'])==signed[0]
    checks.append('old V001 exported by warehouse after V002 exists, without rewriting original snapshot')
    # A receives V001 while another writer commits V002 before A reads it back.
    from unittest.mock import patch
    racework=wr.Repository(out/'race-settings');first=no.Repository(racework);second=no.Repository(racework);store=orders[0]['store'];racebatch=nd.prepare_batch(first,p,[store]);committed=threading.Event();replaced=threading.Event();errors=[];readbacks=[];get=first.get
    newer=copy.deepcopy(p);newer['results'][store]['qa_note']='concurrent V002'
    def writer():
        if not committed.wait(30):errors.append('no first commit');return
        try:second.confirm(newer,[store],nd.job_head(racework,p['day'])['revision'],{store:1})
        except Exception as exc:errors.append(str(exc))
        finally:replaced.set()
    def readback(day,code,revision=None,**kw):
        if not readbacks:committed.set();assert replaced.wait(30)
        value=get(day,code,revision,**kw);readbacks.append(value['revision']);return value
    t=threading.Thread(target=writer);t.start();racefolder=out/'race-exports';racefolder.mkdir()
    with patch.object(first,'get',side_effect=readback):race=nd.execute_batch(first,racebatch,racefolder)
    t.join(30);assert not errors and not t.is_alive() and readbacks==[1];assert first.head(p['day'],store)['revision']==2
    events=[e for e in first.events() if e['action']=='带号导出'];assert len(events)==2 and all(e['revision']==1 for e in events)
    for path in race['paths']:
        b=openpyxl.load_workbook(path);props={v.name:v.value for v in b.custom_doc_props};assert props['修订']=='V001';b.close()
    checks.append('concurrent B commits V002 before A readback: both A warehouse files and logs remain bound to V001')
    # Select one warehouse: same whole-store parent, no second budget/archive.
    sw=wr.Repository(out/'subset-settings');sr=no.Repository(sw);sub=nd.prepare_owned_batch(sr,copy.deepcopy(p),[store],warehouses={store:['山东仓']});subdir=out/'subset';subdir.mkdir();sent=nd.execute_batch(sr,sub,subdir);assert len(sent['paths'])==1 and '山东仓' in sent['paths'][0];assert sr.get(p['day'],store)['payload']['result']==p['results'][store]
    checks.append('selective one-warehouse export still archives entire store result with one parent revision')
    report={'passed':True,'baseline_commit':'6894be792e99f9aa179c1fc5ac8a8e24ffe134b4','checks':checks,'comparisons':comparisons,'allowed_changes':['warehouse file boundaries and filename','child order number in U/header/comment/properties','parent/warehouse custom properties','feedback warehouse provenance/status','store default save metadata'],'business_differences':0,'source_hashes':{str(f.relative_to(ROOT)):sha(f) for f in (ROOT/'src').glob('*') if f.is_file()}}
    (out/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps({'passed':True,'checks':len(checks),'comparison_cases':len(comparisons)}))
if __name__=='__main__':main()
