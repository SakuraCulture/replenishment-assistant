"""Alternating isolated baseline/candidate export timings and strict ordered Excel comparison."""
import os
os.environ['OPENPYXL_LXML']='False'
import argparse,copy,hashlib,json,pickle,statistics,subprocess,sys,time,zlib,threading
from pathlib import Path
from itertools import zip_longest
ROOT=Path(__file__).resolve().parents[1]


def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def write(p,d):p.write_text(json.dumps(d,ensure_ascii=False,indent=2),encoding='utf-8')


def worker(source,case,out):
    sys.path.insert(0,str(source/'src'))
    import xlsx_exports as x,result_semantics as sem,psutil
    import openpyxl
    result=pickle.loads(zlib.decompress((ROOT/'qa/normal_070/rc3-candidate/real_None_True.pklz').read_bytes()))
    if case=='subset':result['rows']=result['rows'][:2000]
    statuses={};results={}
    if case=='five':
        for i in range(5):
            r=copy.deepcopy(result);r['rows']=r['rows'][:500];r['store']=f'SYNTH_{i}';r['store_name']=f'合成店{i}'
            for d in r['rows']:d['B']=r['store'];d['A']=r['store_name']
            r['summary']=sem.summarize(r['rows']);results[r['store']]=r;statuses[r['store']]={'name':r['store_name'],'state':'已完成'}
    messages=[];peak=[0];stop=threading.Event();proc=psutil.Process()
    def sample():
        while not stop.wait(.02):peak[0]=max(peak[0],proc.memory_info().rss)
    thread=threading.Thread(target=sample);thread.start();start=time.perf_counter()
    if case=='five':x.export_batch(results,statuses,out/'files',per_store=True,progress=messages.append)
    else:x.export_xlsx(result,out/'export.xlsx',progress=messages.append)
    elapsed=time.perf_counter()-start;stop.set();thread.join()
    write(out/'measurement.json',{'seconds':elapsed,'peak_rss_mib':peak[0]/1024**2,'progress_events':len(messages),'row_progress_events':sum(hasattr(m,'current') for m in messages),'openpyxl':openpyxl.__version__,'lxml':openpyxl.LXML})


def compare(a,b):
    import openpyxl,zipfile,xml.etree.ElementTree as ET
    aa=openpyxl.load_workbook(a,read_only=True);bb=openpyxl.load_workbook(b,read_only=True);assert aa.sheetnames==bb.sheetnames
    count=0;diff=[]
    whitelist={('运行概览','B2'),('运行概览','B4'),('版本及参数','B3')}
    for left,right in zip(aa,bb):
        for ar,br in zip_longest(left.iter_rows(),right.iter_rows()):
            assert ar is not None and br is not None and len(ar)==len(br)
            for ac,bc in zip(ar,br):
                count+=1
                assert type(ac) is type(bc), (left.title, 'cell presence/type')
                for field in ['value','data_type','number_format','font','fill','alignment','border','protection']:
                    if field=='value' and (left.title,getattr(ac,'coordinate','')) in whitelist:continue
                    av,bv=getattr(ac,field,None),getattr(bc,field,None)
                    if av!=bv or field=='value' and type(av)!=type(bv):diff.append([left.title,getattr(ac,'coordinate',''),field,str(av),str(bv)])
    aa.close();bb.close()
    # Compare all worksheet-level XML outside cell data too: widths, filters,
    # freezes, print settings, column order and dimensions (not just cell styles).
    with zipfile.ZipFile(a) as za,zipfile.ZipFile(b) as zb:
        names=[n for n in za.namelist() if n.startswith('xl/worksheets/sheet') and n.endswith('.xml')]
        assert names==[n for n in zb.namelist() if n.startswith('xl/worksheets/sheet') and n.endswith('.xml')]
        for name in names:
            trees=[ET.fromstring(z.read(name)) for z in [za,zb]]
            for tree in trees:
                for node in list(tree):
                    if node.tag.endswith('}sheetData'):tree.remove(node)
            if ET.tostring(trees[0])!=ET.tostring(trees[1]):diff.append([name,'worksheet metadata'])
    return {'cells':count,'differences':diff}


def main(a):
    out=a.output.resolve();out.mkdir(parents=True,exist_ok=False);q=ROOT/'qa/export_071';baseline=q/'baseline';measurements=[];comparisons=[]
    for case in ['subset','five','full']:
        rounds=1 if case=='full' else 3
        for repeat in range(rounds):
            dirs={}
            for mode in (['old','new'] if repeat%2==0 else ['new','old']):
                dest=out/f'{case}-{repeat}-{mode}';dest.mkdir();dirs[mode]=dest
                subprocess.run([sys.executable,__file__,'--worker','--source',str(baseline if mode=='old' else ROOT),'--case',case,'--output',str(dest)],check=True,cwd=ROOT)
                m=json.loads((dest/'measurement.json').read_text());measurements.append({'case':case,'repeat':repeat,'mode':mode,**m});print(json.dumps(measurements[-1]),flush=True)
            if repeat==0:
                files=lambda d:sorted(d.rglob('*.xlsx'),key=lambda p:p.name.split('_')[1] if p.name.startswith('SYNTH_') else p.name)
                old,new=files(dirs['old']),files(dirs['new']);assert len(old)==len(new)
                for p,r in zip(old,new):comparisons.append({'case':case,'file':p.name,**compare(p,r)})
    manifest=json.loads((q/'baseline_manifest.json').read_text(encoding='utf-8'))
    protected=[p for p,h in manifest['protected'].items() if not Path(p).exists() or sha(p)!=h]
    core={p:{'before':h,'after':sha(ROOT/p)} for p,h in manifest['core'].items()}
    summary={}
    for case in ['subset','five','full']:
        summary[case]={}
        for mode in ['old','new']:
            rows=[r for r in measurements if r['case']==case and r['mode']==mode];v=[r['seconds'] for r in rows]
            summary[case][mode]={'n':len(v),'median_seconds':statistics.median(v),'min_seconds':min(v),'max_seconds':max(v),'max_peak_rss_mib':max(r['peak_rss_mib'] for r in rows)}
    differences=sum(len(c['differences']) for c in comparisons)
    r={'passed':not differences and not protected and all(v['before']==v['after'] for v in core.values()),'version':json.loads((ROOT/'src/version.json').read_text())['version'],'source_commit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),'baseline_commit':'850d31ded8595df72d221c7dc87358d522432ca7','candidate_sources':{str(p.relative_to(ROOT)):sha(p) for p in sorted((ROOT/'src').glob('*')) if p.is_file()},'differences':differences,'comparisons':comparisons,'measurements':measurements,'summary':summary,'core':core,'changed_protected':protected,'whitelist':['运行概览!B2 application version','运行概览!B4 build info','版本及参数!B3 export wall-clock time','OOXML ZIP/core-property timestamps'],'limits':'Same Windows host/Python/openpyxl, etree (same as EXE), sequential alternating trials; warm OS file cache uncontrolled. Includes rows/styles/XML/ZIP/progress callbacks, excludes input/calculation and signing. RSS sampled every20ms within export worker, not total OS usage. subset=first2000 real rows; full=14244 real rows once; five=5 synthetic stores copied from first500 real rows, per-store files plus existing full summary, not user actual five-store run. Merchant requested schema/filter changes are tested separately, not claimed byte-equivalent.'}
    write(out/'export_report.json',r);print(json.dumps(summary),flush=True);assert r['passed']


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--worker',action='store_true');p.add_argument('--source',type=Path);p.add_argument('--case');p.add_argument('--output',type=Path,required=True);a=p.parse_args()
    if a.worker:worker(a.source,a.case,a.output)
    else:main(a)
