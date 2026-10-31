"""Two real application launches on generated data; no user settings touched."""
import argparse,copy,hashlib,json,os,sqlite3,subprocess,sys,time
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT/'src'),str(ROOT/'tests'),str(ROOT/'qa/export_071/qa-deps')]
import openpyxl,psutil,normal_orders as no,normal_exports as ne,work_repository as wr,lifecycle_io as io
from test_archive_orders import seed

def fixture(out):
    work,repo,p=seed(out)
    for sku in ('0003','0004'):
        d=copy.deepcopy(p['result']['rows'][0]);d['E']=sku;d['D']='合成商品'+sku;p['result']['rows'].append(d)
    digest=no.fingerprint(p['result'])
    with sqlite3.connect(work.path) as db:db.execute('UPDATE normal_revisions SET digest=?,payload=?',(digest,no.encode(p)))
    order=repo.get(wr.today(),'T')
    for wh,tag in [('河北仓','hb'),('山东仓','sd')]:
        original=out/('signed-'+tag+'.xlsx');ne.export(order,original,warehouse=wh)
        for suffix in ('','-v2') if tag=='hb' else ('',):
            b=openpyxl.load_workbook(original)
            for ws in b:
                for n in range(ws.max_row,1,-1):
                    sku=ws.cell(n,5).value
                    if sku=='0004':ws.delete_rows(n)
                    else:ws.cell(n,17).value=None if sku=='0003' else 8 if suffix else 6
            b.save(out/('return-'+tag+suffix+'.xlsx'));b.close()
    b=openpyxl.Workbook();s=b.active;s.title='原付款明细';s.append(io.PAYMENT_HEADERS);s.append(['1298602','S0000001','测试店','001588','2026-10-05 15:01:24','线下店铺','0001','合成商品',24,4,1.3,5.2,'']);b.save(out/'payment.xlsx');b.close()
    for name,quantity in [('receipt.xlsx',None),('receipt-invalid.xlsx','invalid')]:
        b=openpyxl.Workbook();s=b.active;s.append(io.fields('receipt'));r={'补货单号':'BH-T-ZC-HB','签发修订':1,'供货仓':'河北仓','商品编码':'0001','数量单位':'库存单位','实际收货数量':quantity,'收货单号':'R001','收货日期':wr.today()};s.append([r.get(f) for f in io.fields('receipt')]);b.save(out/name);b.close()

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--exe',type=Path);ap.add_argument('--output',required=True,type=Path);a=ap.parse_args();out=a.output.resolve();out.mkdir(parents=True,exist_ok=False);fixture(out)
    env=os.environ.copy()
    for k in ('PYTHONPATH','PYTHONHOME','QT_PLUGIN_PATH','QT_QPA_PLATFORM'):env.pop(k,None)
    env['REPLENISHMENT_DATA_DIR']=str(out/'settings');env['PATH']=env['WINDIR']+';'+str(Path(env['WINDIR'])/'System32')
    if not a.exe:env['QT_QPA_PLATFORM']='offscreen'
    runs=[]
    for phase in ('exercise','reopen'):
        cmd=[str(a.exe.resolve())] if a.exe else [sys.executable,str(ROOT/'src/app.py')]
        with (out/(phase+'.log')).open('w',encoding='utf-8') as log:
            p=subprocess.Popen(cmd+['--reconciliation-smoke',str(out),phase],cwd=out,env=env,stdout=log,stderr=log);host=psutil.Process(p.pid);start=time.monotonic();children=set();peak=0
            while p.poll() is None:
                if time.monotonic()-start>180:
                    for c in reversed(host.children(recursive=True)):c.terminate()
                    host.terminate();p.wait(10);raise TimeoutError('QA app timeout; owned process tree stopped')
                try:
                    tree=host.children(recursive=True);children.update((c.pid,c.create_time()) for c in tree);peak=max(peak,sum(c.memory_info().rss for c in [host]+tree))
                except psutil.Error:pass
                time.sleep(.05)
        r=json.loads((out/('ui-'+phase+'.json')).read_text(encoding='utf-8'));left=[]
        for pid,created in children:
            try:
                c=psutil.Process(pid)
                if c.create_time()==created and c.is_running():left.append(pid)
            except psutil.Error:pass
        r.update(exit_code=p.returncode,remaining_children=left,development=not bool(a.exe),exe_sha256=hashlib.sha256(a.exe.read_bytes()).hexdigest() if a.exe else None,elapsed_seconds=time.monotonic()-start,peak_tree_mib=peak/2**20,sample_ms=50,synthetic=True);runs.append(r)
        (out/('ui-'+phase+'.json')).write_bytes(json.dumps(r,ensure_ascii=False,indent=2).encode())
        if not r['passed'] or p.returncode or left:raise AssertionError(r)
    (out/'report.json').write_bytes(json.dumps({'passed':True,'runs':runs},ensure_ascii=False,indent=2).encode());print('PASS actual EXE' if a.exe else 'PASS source UI')
if __name__=='__main__':main()
