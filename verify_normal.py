"""0.7 normal-order/new-product acceptance, isolated synthetic fixtures and exact checks."""
import argparse,copy,hashlib,json,sqlite3,subprocess,sys
from pathlib import Path
ROOT=Path(__file__).resolve().parent;sys.path[:0]=[str(ROOT/'src'),str(ROOT/'tools')]
import new_product_focus as focus
import engine,batch_service,warehouse_catalog as wc,priority_list,work_repository as wr,normal_orders,new_products as np,normal_exports
from version_info import VERSION,MODEL_VERSION
from verify_batch import write_fixture
from parallel_qa import diff


def fixtures(out):
    kinds=['moving','new_must','new_plain','new_blank_box','new_disabled','new_stopped','new_local','new_no_price','new_anomaly','missing_n','recovery']
    for wh in wc.WAREHOUSES:
        rows=[{'商品编码':k,'商品名称':'合成-'+k,'分类':'饮品','基本售价':None if k=='new_no_price' else 2,'出库箱规':None if k=='new_blank_box' else 6 if k=='new_must' else 1,'商品状态':'禁用' if k=='new_disabled' else '启用','可用数':0 if wh=='河北仓' and k=='new_plain' else 100} for k in kinds if k!='new_local']
        write_fixture(out/(wh+'.xlsx'),list(rows[0]),rows)
    rows=[]
    for store in ['QA_A','QA_B']:
        for k in kinds:
            rows.append({'仓库/门店编码':store,'仓库/门店名称':'合成店'+store,'商品编码':k,'商品名称':'=合成-'+k,'日销':2 if k=='moving' else 1 if k=='new_anomaly' else 0,'周销':14 if k=='moving' else 0,'月销':60 if k=='moving' else 0,'可用库存':0,'在途库存':0,'缺货天数（天）':None if k=='missing_n' else 10 if k=='recovery' else 0,'补货状态':'停采' if k=='new_stopped' else '正常','基础单位':'瓶'})
    write_fixture(out/'sales.xlsx',list(rows[0]),rows)
    profile=priority_list.make_profile([{'sku':'new_must','name':'必卖新品','category':'饮品'}],'合成验收必卖')
    (out/'must.json').write_text(json.dumps(profile,ensure_ascii=False,indent=2),encoding='utf-8')
    write_fixture(out/'新品明细.xlsx',['商品编码','商品名称','关注'],[{'商品编码':'new_must','商品名称':'必卖新品','关注':'是'},{'商品编码':'new_plain','商品名称':'普通新品','关注':0},{'商品编码':'outside','关注':True},{'商品编码':'1991707366294683677','商品名称':'仅用于真实样本验收的关注选择','关注':'关注'}])
    return profile


def run(out):
    out.mkdir(parents=True,exist_ok=False);checks=[];profile=fixtures(out)
    catalogs={w:wc.import_catalog(out/(w+'.xlsx'),w) for w in wc.WAREHOUSES};loaded=batch_service.load([out/'sales.xlsx'],catalogs)
    cfg={'warehouses':engine.DEFAULT_WAREHOUSES,'exception':True,'budget':12,'supply_route':{'primary':'河北仓','secondary':'山东仓'}}
    base=batch_service.calculate_store(loaded,'QA_A',cfg,profile);other=batch_service.calculate_store(loaded,'QA_B',cfg,profile);unlimited=batch_service.calculate_store(loaded,'QA_A',{**cfg,'budget':None},profile)
    original=copy.deepcopy(base);other_before=copy.deepcopy(other);rows={d['E']:d for d in base['rows']}
    assert np.is_new(rows['new_must']) and not np.blocked_reason(rows['new_must'])
    assert np.is_new(rows['missing_n']) and not np.blocked_reason(rows['missing_n']) and not np.is_new(rows['recovery'])
    for sku in ['new_disabled','new_stopped','new_local','new_no_price','new_anomaly']:assert np.blocked_reason(rows[sku]),sku
    for value in [None,'']:
        d=copy.deepcopy(rows['new_must']);d['N']=value;assert np.is_new(d)
    for value in [False,-1,'待核']:
        d=copy.deepcopy(rows['new_must']);d['N']=value;assert not np.is_new(d),value
    for key in ['J','K','L']:
        for value in [None,'',False]:
            d=copy.deepcopy(rows['new_must']);d[key]=value;assert not np.is_new(d)
    assert rows['new_blank_box']['BD']==1 and rows['new_plain']['AU']=='山东仓'
    blank=np.apply(base,{'missing_n':1});blankrow=next(d for d in blank['rows'] if d['E']=='missing_n');assert blankrow['N'] is None and '按用户确认口径' in blankrow['_new_stockout_basis']
    checks.append('user-confirmed 2026-10-30: blank N recognized as zero with raw value/audit retained; J/K/L missing/empty/bool and N bool/negative/pending not zero; disabled/stopped/local/unpriced/sales-anomaly blocked; blank carton=1; secondary routing preserved')
    assert not diff(base,np.apply(base,{}))
    chosen=np.apply(base,{'new_must':6});newrows={d['E']:d for d in chosen['rows']}
    assert engine.order_quantity(newrows['new_must'])==6 and chosen['summary']['amount']==12
    assert engine.order_quantity(rows['moving'])>0 and engine.order_quantity(newrows['moving'])==0
    assert newrows['moving']['_holding_gap']
    assert newrows['new_must']['AS']==0 and newrows['new_must']['BV']==0
    assert not diff(other,other_before)
    # Explicit full-field invariant: only named candidate/execution allocation fields may differ.
    assert not diff(base,original)
    allowed=('_candidate_','_allocated_','_budget_','_holding_','_execution_','_new_')
    for d in chosen['rows']:
        before=rows[d['E']]
        assert not diff({k:v for k,v in before.items() if not k.startswith(allowed)},{k:v for k,v in d.items() if not k.startswith(allowed)})
    full=np.apply(unlimited,{'new_must':6,'new_blank_box':1});assert full['summary']['amount']==unlimited['summary']['amount']+14
    assert engine.order_quantity(next(d for d in full['rows'] if d['E']=='new_must'))==6
    zero=batch_service.calculate_store(loaded,'QA_A',{**cfg,'budget':0},profile);zeronew=np.apply(zero,{'new_must':6});assert zeronew['summary']['amount']==0 and next(d for d in zeronew['rows'] if d['E']=='new_must')['_new_requested']==6
    for bad in [{'new_must':1},{'new_disabled':1},{'new_stopped':1},{'new_no_price':1},{'missing':1},{'new_must':True},{'new_must':-6}]:
        try:np.apply(base,bad)
        except ValueError:pass
        else:raise AssertionError(bad)
    checks.append('empty selection exact field/type/order identity; mandatory new stock receives budget ahead of ordinary moving goods; minimum holding gap explicit; unlimited and zero budgets; AS/BV/intermediates invariant; invalid selection rejected')
    focusinfo=focus.inspect(out/'新品明细.xlsx');focused=focus.import_list(focusinfo,next(iter(focusinfo['sheets'])),'商品编码','关注',wr.today()[:7]);assert focus.members(focused)=={'new_must','outside','1991707366294683677'}
    work=wr.Repository(out/'settings');repo=normal_orders.Repository(work)
    payload={'schema':1,'day':wr.today(),'model_version':MODEL_VERSION,'app_version':VERSION,'states':{s:{'name':r['store_name'],'state':'已完成','error':''} for s,r in [('QA_A',base),('QA_B',other)]},'results':{'QA_A':base,'QA_B':other},'input_dirty':False,'new_focus':focused}
    issue=repo.confirm(payload,['QA_A','QA_B'],None,{'QA_A':None,'QA_B':None});assert len(repo.list())==2
    first=repo.get(wr.today(),'QA_A');job=work.get_job(wr.today());ids=first['payload']['item_ids'];assert first['revision']==1
    same=repo.confirm(payload,['QA_A'],job['revision'],{'QA_A':1})[0];assert same['revision']==1
    changed=copy.deepcopy(payload);changed['results']['QA_A']=chosen
    assert repo.get(wr.today(),'QA_A')['revision']==1  # unconfirmed draft does not touch normal order
    revised=repo.confirm(changed,['QA_A'],job['revision'],{'QA_A':1})[0]
    assert revised['number']==first['number'] and revised['revision']==2 and revised['payload']['item_ids']==ids
    assert repo.get(wr.today(),'QA_A',1)['payload']['result']==first['payload']['result'];assert repo.get(wr.today(),'QA_B')['revision']==1
    assert len(repo.list())==2
    try:repo.confirm(payload,['QA_A'],job['revision'],{'QA_A':1})
    except wr.Conflict:pass
    else:raise AssertionError('stale writes accepted')
    checks.append('first issue saves draft atomically; same result idempotent; same-day stable number with immutable V001/V002 and stable item IDs; other store unchanged; compare-and-swap rejects stale revision')
    beforejob=work.get_job(wr.today());beforeorder=repo.get(wr.today(),'QA_A')
    with work.connect() as db:db.execute("CREATE TRIGGER reject_revision BEFORE INSERT ON normal_revisions BEGIN SELECT RAISE(ABORT,'injected failure'); END")
    try:repo.confirm(payload,['QA_A'],beforejob['revision'],{'QA_A':2})
    except sqlite3.IntegrityError:pass
    else:raise AssertionError('failure injection failed')
    assert work.get_job(wr.today())==beforejob and repo.get(wr.today(),'QA_A')==beforeorder
    with work.connect() as db:db.execute('DROP TRIGGER reject_revision')
    for mutate in ['input','state','day']:
        bad=copy.deepcopy(changed)
        if mutate=='input':bad['input_dirty']=True
        elif mutate=='state':bad['states']['QA_A']['state']='待重算'
        else:bad['day']='2020-01-01'
        try:repo.confirm(bad,['QA_A'],beforejob['revision'],{'QA_A':2})
        except ValueError:pass
        else:raise AssertionError(mutate)
    assert wr.Repository(out/'settings').get_job(wr.today())==beforejob
    assert list((out/'settings').glob('workbench.before-normal-orders.*.sqlite3'))
    checks.append('database failure after draft update rolls entire transaction back; stale/dirty/historical results rejected; schema1 old repository remains readable; SQLite consistent pre-extension backup created')
    path=out/'确认单.xlsx';normal_exports.export(revised,path)
    import openpyxl
    book=openpyxl.load_workbook(path)
    assert book.sheetnames==normal_exports.SHEETS
    sheet=book['本月新品'];assert sheet.max_row==2;assert sheet['E2'].value=='new_must';assert sheet['D2'].value=='=合成-new_must' and sheet['D2'].data_type=='s';assert sheet['G2'].value=='瓶';assert sheet['Q2'].value==6
    assert sheet['U2'].value==revised['number'] and sheet['V2'].value=='必卖首铺' and sheet.max_column==22
    exported=[]
    for name in normal_exports.SHEETS:
        ws=book[name];assert [c.value for c in ws[1][:20]]==normal_exports.HEADERS;assert ws.freeze_panes is None and not ws.auto_filter.ref
        for row in ws.iter_rows(min_row=2,values_only=True):assert row[1]=='QA_A';exported.append(row[4])
        for row in ws:
            for c in row[:20]:assert c.font.name=='宋体' and c.font.sz==11 and not c.font.bold and c.number_format==('0.00' if c.row>1 and c.column==16 else '@' if c.row>1 and c.column==18 else 'General') and c.fill.patternType is None
        assert ws.max_column==22 and not any(d.hidden for d in ws.column_dimensions.values())
    assert len(exported)==len(set(exported))
    expected=[d['E'] for d in chosen['rows'] if d['_source']=='资料匹配' and not engine.exclusion_reason(d) and (np.is_new(d) or engine.orderable(d))];assert set(exported)==set(expected)
    oldnew={r[4]:r for r in book['往期新品'].iter_rows(min_row=2,values_only=True)}
    # 0.7.3 explicit user change: disabled/stopped no longer appear in any working export.
    assert 'new_disabled' not in oldnew and 'new_stopped' not in oldnew
    for code in ['new_plain','new_no_price','missing_n']:assert oldnew[code][16]==0
    assert 'new_local' not in oldnew
    assert oldnew['missing_n'][13] is None and oldnew['new_no_price'][15] is None
    assert normal_exports.filename(revised)==wr.today()+'合成店QA_A（精简版）.xlsx'
    book.close();oldhash=hashlib.sha256(path.read_bytes()).hexdigest()
    try:normal_exports.export(revised,path)
    except FileExistsError:pass
    else:raise AssertionError('export overwrote an existing file')
    assert hashlib.sha256(path.read_bytes()).hexdigest()==oldhash and repo.get(wr.today(),'QA_A')['revision']==2
    repo.log(revised,'带号导出',{'path':str(path),'sha256':oldhash});assert len(repo.events(revised['oid']))==3
    checks.append('merchant workbook exact sheet/row identity, inventory unit, all central-warehouse new candidates exported and unarranged zero; first 20 headers preserved; P numeric price/R status format; U order number/V concise reason; removed audit columns stay in DB (authorized 2026-10-30 change); raw blank N retained, formula-like product name stored as text, no cross-store rows; no-clobber failed export preserves archive; audit log saved')
    # Focus changes are revision-worthy even if quantities are exactly unchanged.
    emptied=copy.deepcopy(changed);emptied['new_focus']=focus.empty(wr.today()[:7]);job=work.get_job(wr.today())
    third=repo.confirm(emptied,['QA_A'],job['revision'],{'QA_A':2})[0];assert third['revision']==3 and third['payload']['result']==revised['payload']['result'];assert repo.get(wr.today(),'QA_A',2)['payload']['new_focus']==focused
    normal_exports.export(repo.get(wr.today(),'QA_A',2),out/'冻结关注修订.xlsx')
    focusrepo=focus.Repository(out/'focus-settings');focusrepo.save(focused);assert focusrepo.load()==focused
    for broken in [None,focus.empty('2020-01')]:
        bad=copy.deepcopy(changed);bad['new_focus']=broken
        try:repo.confirm(bad,['QA_A'],work.get_job(wr.today())['revision'],{'QA_A':3})
        except ValueError:pass
        else:raise AssertionError('missing/stale focus accepted')
    for values in [['是','否'],['是','??']]:
        write_fixture(out/'关注错误.xlsx',['商品编码','关注'],[{'商品编码':'001','关注':v} for v in values]);info=focus.inspect(out/'关注错误.xlsx')
        try:focus.import_list(info,next(iter(info['sheets'])),'商品编码','关注',wr.today()[:7])
        except ValueError:pass
        else:raise AssertionError('conflicting/unknown focus accepted')
    write_fixture(out/'关注编码.xlsx',['商品编码','关注'],[{'商品编码':'001','关注':1},{'商品编码':'001','关注':'关注'},{'商品编码':'1991707366294683677','关注':None}]);info=focus.inspect(out/'关注编码.xlsx');parsed=focus.import_list(info,next(iter(info['sheets'])),'商品编码','关注',wr.today()[:7]);assert [x['sku'] for x in parsed['items']]==['001','1991707366294683677'] and parsed['items'][0]['source_rows']==[2,3]
    info['sha256']='changed'
    try:focus.import_list(info,next(iter(info['sheets'])),'商品编码','关注',wr.today()[:7])
    except ValueError:pass
    else:raise AssertionError('changed source accepted')
    checks.append('explicit focus import/reuse; exact long and leading-zero codes; duplicate same status traced, conflict/unknown rejected; source hash change rejected; stale/missing month blocked; focus-only revision preserves quantities and old grouping; outside SKU never added')
    (out/'payload.json').write_text(wr.dumps(changed),encoding='utf-8')
    report={'passed':True,'version':VERSION,'source_commit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),'checks':checks,'fixture_note':'22 synthetic store-product rows, two synthetic warehouse feeds; not real observed demand','budget_example':{'before':base['summary'],'after':chosen['summary'],'changes':np.changes(base,chosen)},'protected_algorithm_files':['model.py','one_box_exception.py','allocation.py','default_config.json','default_must_sell.json']}
    (out/'normal_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=='__main__':
    import multiprocessing
    multiprocessing.freeze_support();p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);a=p.parse_args();run(a.output.resolve())
