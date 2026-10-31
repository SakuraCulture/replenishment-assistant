"""Validate authored cells and publish a synthetic-only ZIP/manifest, not sessions."""
import argparse,hashlib,json,zipfile
from pathlib import Path
import openpyxl

def main(package,authoring,destination):
    cases=json.loads((authoring/'cases.json').read_text(encoding='utf8'))
    lines=['# 38条业务用例','','先完整解压本目录ZIP，打开“从这里开始.html”。本页供云端查阅；实际结果填写包内验收记录。','']
    for c in cases:
        lines.extend([f'## {c["id"]} {c["scene"]} · {c["title"]}','',f'**前置：** {c["setup"]}',f'\n**材料：** {c["files"] or "按步骤操作，无额外文件"}',f'\n**操作：** {c["steps"]}',f'\n**通过标准：** {c["expected"]}',''])
    (destination/'用例索引.md').write_text('\n'.join(lines),encoding='utf8',newline='\n')
    checks=[]
    for specname in ['inputs-spec.json','lifecycle-spec.json','record-spec.json']:
        for spec in json.loads((authoring/specname).read_text(encoding='utf8')):
            book=openpyxl.load_workbook(package/spec['file'],data_only=False)
            if book.sheetnames!=[s['name'] for s in spec['sheets']]:raise ValueError('sheet mismatch')
            for tab in spec['sheets']:
                sheet=book[tab['name']]
                for ri,row in enumerate([tab['headers'],*tab['rows']],1):
                    for ci,value in enumerate(row,1):
                        cell=sheet.cell(ri,ci)
                        if not (cell.value==value or value=='' and cell.value is None):raise ValueError((spec['file'],ri,ci,cell.value,value))
                        if isinstance(value,str) and value and (value.startswith('0') or value.startswith('BH-')) and cell.data_type!='s':raise ValueError('identifier not text')
                        if cell.data_type=='e':raise ValueError('Excel error')
            checks.append({'file':spec['file'],'sheets':len(book.sheetnames),'exact_values':True});book.close()
    # Explicit top-level allowlist: never publish a tester's locally generated sessions.
    allowed={'01_补货输入','02_原签发参考','03_生命周期单据','04_异常练习','05_预期结果','seed','用例数据.json','业务验收记录.xlsx','从这里开始.html','launch-practice.ps1','开始补货练习.cmd','开始生命周期复查练习.cmd','新开生命周期复查练习.cmd','新开补货练习.cmd'}
    files=[p for p in package.rglob('*') if p.is_file() and p.relative_to(package).parts[0] in allowed and not p.name.endswith('.inspect.ndjson')]
    manifest={'synthetic':True,'files':{p.relative_to(package).as_posix():{'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(files)},'workbook_validation':checks,'manual_cases':len(cases)}
    archive=destination/'业务验收练习包.zip'
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        for p in files:z.write(p,'业务验收练习包/'+p.relative_to(package).as_posix())
    with zipfile.ZipFile(archive) as z:
        if z.testzip() is not None:raise ValueError('ZIP verification failed')
    manifest['archive']={'file':archive.name,'bytes':archive.stat().st_size,'sha256':hashlib.sha256(archive.read_bytes()).hexdigest()}
    (destination/'package-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf8',newline='\n')
    print(json.dumps({'package_files':len(files),'xlsx':sum(p.suffix=='.xlsx' for p in files),'archive':manifest['archive'],'artifact_workbooks_verified':len(checks)},ensure_ascii=False))

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--package',type=Path,required=True);p.add_argument('--authoring',type=Path,required=True);p.add_argument('--destination',type=Path,required=True);a=p.parse_args();main(a.package,a.authoring,a.destination)
