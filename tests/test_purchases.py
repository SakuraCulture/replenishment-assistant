"""Portable generated inputs. No user database, external Excel or Qt required."""
import copy,hashlib,json,sqlite3,sys,tempfile,threading,time,unittest
from contextlib import closing
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
import openpyxl
import purchase_io as io,purchases as p,followup,work_repository as wr,normal_orders as no,merchant_feedback as mf
from runtime_support import Cancelled
from test_archive_orders import seed

class PurchaseTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.work,self.repo,self.payload=seed(self.root);self.path=self.work.path;self.heads=p.targets(self.path);self.h=self.heads[0]
    def tearDown(self):self.tmp.cleanup()
    def scan(self,rows=None,filename='purchase.xlsx'):
        path=self.root/filename;book=openpyxl.Workbook();ws=book.active;ws.title='采购明细';ws.append(io.FIELDS)
        for r in rows or [{}]:
            data={'来源系统':'ERP','采购订单号':'P001','采购明细号':'L001','下单日期':wr.today(),'门店编码':'T','供货仓':'河北仓','商品编码':'0001','商品名称':'测试商品','本版实际下单数量':5,'数量单位':'库存单位','补货单号':self.h['number'],'签发修订':1,**r};ws.append([data.get(f) for f in io.FIELDS])
        book.save(path);book.close();info=io.inspect([path])[0];s=info['sheets'][0];spec={**info,**s,'mapping':{f:i for i,f in enumerate(io.FIELDS)}}
        return p.auto_links(io.scan([spec]),self.heads)
    def save(self,scan=None):
        scan=scan or self.scan();plan=p.prepare(self.path,scan);receipt=p.commit(self.path,plan);return scan,plan,receipt
    def original(self):
        with closing(sqlite3.connect(self.path)) as db:return {name:db.execute('SELECT * FROM '+name).fetchall() for name in ['normal_orders','normal_revisions','normal_work_revisions','source_files']}
    def home(self,**kw):return followup.read(self.path,self.root/'followup.db',kw)
    def reply(self,values,version=1):
        with closing(sqlite3.connect(self.path)) as db,db:
            db.execute('CREATE TABLE IF NOT EXISTS merchant_feedback(oid TEXT,revision INTEGER,version INTEGER,saved TEXT,payload TEXT,PRIMARY KEY(oid,revision,version))')
            db.execute('INSERT INTO merchant_feedback VALUES(?,?,?,?,?)',('oid',1,version,wr.now(),wr.dumps({'values':values,'warehouses':{'河北仓':{'received':len(values)}}})))
    def second_revision(self):
        h=self.repo.get(wr.today(),'T');result=copy.deepcopy(h['payload']);result['result']['marker']='revision2';digest=no.fingerprint(result['result'])
        with closing(sqlite3.connect(self.path)) as db,db:
            db.execute('INSERT INTO normal_revisions VALUES(?,?,?,?,?)',('oid',2,digest,wr.now(),no.encode(result)));db.execute('UPDATE normal_orders SET current=2')

    def test_first_import_versions_units_and_signed_immutable(self):
        before=self.original();scan,plan,receipts=self.save();self.assertEqual(self.original(),before);self.assertEqual(receipts[0]['version'],1);self.assertEqual(scan['rows'][0]['sku'],'0001');self.assertEqual(p.comparison(self.path,self.h)['rows'][0][4],5);self.assertEqual(p.comparison(self.path,self.h)['rows'][0][3],None)
        source=scan['files'][0];dest=self.root/'download.xlsx';p.attachment(self.path,source['sha'],dest);self.assertEqual(dest.read_bytes(),Path(source['path']).read_bytes())
    def test_repeated_import_does_not_accumulate(self):
        scan,plan,r=self.save();again=p.commit(self.path,p.prepare(self.path,scan));self.assertEqual(again[0]['version'],1);self.assertFalse(again[0]['changed']);self.assertEqual(p.comparison(self.path,self.h)['rows'][0][4],5)
    def test_full_new_version_replaces_keeps_history(self):
        self.save();scan=self.scan([{'本版实际下单数量':3}],'v2.xlsx');_,_,r=self.save(scan);self.assertEqual(r[0]['version'],2);self.assertEqual(p.comparison(self.path,self.h)['rows'][0][4],3)
        with closing(p.connect(self.path)) as db:self.assertEqual(p.current(db,p.docid(scan['rows'][0]),1)['payload']['rows'][0]['base_quantity'],'5')
    def test_missing_quantity_not_zero_zero_explicit(self):
        with self.assertRaisesRegex(ValueError,'空白'):self.scan([{'本版实际下单数量':None}])
        self.save(self.scan([{'本版实际下单数量':0}]));self.assertEqual(p.comparison(self.path,self.h)['rows'][0][4],0)
    def test_unknown_units_then_explicit_conversion(self):
        scan=self.scan([{'数量单位':'箱'}]);self.assertIsNone(scan['rows'][0]['base_quantity']);self.assertEqual(scan['rows'][0]['links'],[]);self.save(scan);self.assertIsNone(p.comparison(self.path,self.h)['rows'][0][4])
        scan=self.scan([{'数量单位':'箱','折合库存单位系数':24,'换算依据':'订单明确一箱24包'}],'converted.xlsx');self.save(scan);self.assertEqual(p.comparison(self.path,self.h)['rows'][0][4],120)
    def test_two_lines_same_sku_and_two_orders_add_only_links(self):
        scan=self.scan([{}, {'采购明细号':'L002','本版实际下单数量':2},{'采购订单号':'P002','本版实际下单数量':3}]);self.save(scan);self.assertEqual(p.comparison(self.path,self.h)['rows'][0][4],10);self.assertEqual(len(p.documents(self.path)),2)
    def test_duplicate_id_same_content_dedup_conflict_reject(self):
        scan=self.scan([{},{}]);plan=p.prepare(self.path,scan);self.assertEqual(plan['plans'][0]['duplicates'],1);p.commit(self.path,plan);self.assertEqual(p.comparison(self.path,self.h)['rows'][0][4],5)
        with self.assertRaisesRegex(ValueError,'冲突'):p.prepare(self.path,self.scan([{}, {'本版实际下单数量':2}]))
    def test_cross_warehouse_and_extra_requires_explicit(self):
        scan=self.scan();scan['rows'][0]['links']=[p.link(self.heads[1],5)]
        with self.assertRaisesRegex(ValueError,'跨门店/仓库'):p.prepare(self.path,scan)
        scan=self.scan([{'商品编码':'EXTRA'}]);scan['rows'][0]['links']=[p.link(self.h,5)]
        with self.assertRaisesRegex(ValueError,'额外采购'):p.prepare(self.path,scan)
        scan['rows'][0]['links']=[p.link(self.h,5,True)];self.save(scan);rows=p.comparison(self.path,self.h)['rows'];self.assertEqual(rows[-1][0],'EXTRA');self.assertIsNone(rows[-1][2]);self.assertEqual(rows[-1][7],'额外采购')
    def test_split_allocations_and_overallocation(self):
        self.second_revision();heads=p.targets(self.path);old=next(h for h in heads if h['warehouse']=='河北仓' and h['revision']==1);new=next(h for h in heads if h['warehouse']=='河北仓' and h['revision']==2);scan=self.scan();scan['rows'][0]['links']=[p.link(old,2),p.link(new,3)];self.save(scan);self.assertEqual(p.comparison(self.path,old)['rows'][0][4],2);self.assertEqual(p.comparison(self.path,new)['rows'][0][4],3)
        scan['rows'][0]['links'][0]['quantity']='3'
        with self.assertRaisesRegex(ValueError,'超过'):p.prepare(self.path,scan)
    def test_missing_id_update_needs_identity_review(self):
        scan=self.scan([{'采购明细号':None}]);self.save(scan);updated=self.scan([{'采购明细号':None,'本版实际下单数量':3}],'different.xlsx')
        with self.assertRaisesRegex(ValueError,'逐行确认'):p.prepare(self.path,updated)
        updated['rows'][0].update(line_id=scan['rows'][0]['line_id'],identity_reviewed=True);self.save(updated);self.assertEqual(p.comparison(self.path,self.h)['rows'][0][4],3)
    def test_old_revision_requires_explicit_binding_no_migration(self):
        self.save();self.second_revision();heads=p.targets(self.path);new=next(h for h in heads if h['warehouse']=='河北仓' and h['revision']==2);self.assertIsNone(p.comparison(self.path,new)['rows'][0][4]);self.assertEqual(p.comparison(self.path,self.h)['rows'][0][4],5);self.assertTrue(next(r for r in self.home()['records'] if r['warehouse']=='河北仓')['old_revision'])
    def test_preview_stale_signing_and_purchase_conflicts_atomic(self):
        scan=self.scan();plan=p.prepare(self.path,scan);self.second_revision()
        with self.assertRaises(wr.Conflict):p.commit(self.path,plan)
        self.assertEqual(p.documents(self.path),[])
        scan['rows'][0]['links']=[];a=p.prepare(self.path,scan);b=p.prepare(self.path,scan);p.commit(self.path,a)
        with self.assertRaises(wr.Conflict):p.commit(self.path,b)
    def test_reply_changes_invalidate_preview(self):
        plan=p.prepare(self.path,self.scan());self.reply({'0001':5})
        with self.assertRaisesRegex(wr.Conflict,'回传'):p.commit(self.path,plan)
        self.assertEqual(p.documents(self.path),[])
    def test_cancel_preparation_before_transaction_and_boundary(self):
        scan=self.scan();event=threading.Event();event.set()
        with self.assertRaises(Cancelled):p.prepare(self.path,scan,cancel=event)
        event.clear();plan=p.prepare(self.path,scan)
        def progress(s):
            if '即将提交' in str(s):event.set()
        before=self.original()
        with self.assertRaises(Cancelled):p.commit(self.path,plan,progress,event)
        self.assertEqual(self.original(),before);self.assertEqual(p.documents(self.path),[]);event.clear();self.assertEqual(p.commit(self.path,plan,cancel=event)[0]['version'],1)
    def test_cancel_waiting_for_lock_rolls_back(self):
        plan=p.prepare(self.path,self.scan());event=threading.Event();errors=[]
        # Backup first so the intentional write-lock test measures transaction waiting.
        import storage_chunks
        storage_chunks.backup_before_write(self.path,lambda _:None)
        with closing(sqlite3.connect(self.path)) as lock:
            lock.execute('BEGIN IMMEDIATE')
            def save():
                try:p.commit(self.path,plan,cancel=event)
                except Exception as exc:errors.append(exc)
            with patch.object(storage_chunks,'backup_before_write',return_value=None):
                t=threading.Thread(target=save);t.start();time.sleep(.12);event.set();lock.rollback();t.join(5)
        self.assertFalse(t.is_alive());self.assertIsInstance(errors[0],Cancelled);self.assertEqual(p.documents(self.path),[])
    def test_source_changes_rejected_and_failed_batch_atomic(self):
        scan=self.scan([{}, {'采购订单号':'P2'}]);plan=p.prepare(self.path,scan);Path(scan['files'][0]['path']).write_bytes(b'changed')
        with self.assertRaisesRegex(ValueError,'变化'):p.commit(self.path,plan)
        self.assertEqual(p.documents(self.path),[])
    def test_cancelled_order_clears_active_allocations_keeps_old(self):
        self.save();scan=self.scan([{'订单状态':'已取消'}],'cancelled.xlsx');self.assertEqual(scan['rows'][0]['links'],[]);self.save(scan);self.assertIsNone(p.comparison(self.path,self.h)['rows'][0][4]);self.assertEqual(p.documents(self.path)[0]['version'],2)
    def test_todo_does_not_write_user_db_missing_is_pending(self):
        before=hashlib.sha256(self.path.read_bytes()).hexdigest();home=self.home();self.assertEqual(len(home['records']),2);self.assertEqual(hashlib.sha256(self.path.read_bytes()).hexdigest(),before);self.assertEqual(home['records'][0]['purchase_status'],'尚未关联下单资料')
    def test_ack_bound_to_observation_defer_restore_and_conflict(self):
        self.save();self.reply({'0001':3});h=next(r for r in self.home()['records'] if r['warehouse']=='河北仓');self.assertIn('差异待核',h['issues']);followup.note(self.path,h,'已核实保留差异','店主追加两件');h2=next(r for r in self.home()['records'] if r['warehouse']=='河北仓');self.assertNotIn('差异待核',h2['issues']);self.assertTrue(h2['acknowledged'])
        with self.assertRaises(wr.Conflict):followup.note(self.path,h,'暂缓跟进','旧窗口')
        self.reply({'0001':2},2);h3=next(r for r in self.home()['records'] if r['warehouse']=='河北仓');self.assertIn('差异待核',h3['issues']);followup.note(self.path,h3,'暂缓跟进','等待核对');h4=next(r for r in self.home(status='暂缓跟进')['records'] if r['warehouse']=='河北仓');self.assertTrue(h4['deferred']);followup.note(self.path,h4,'恢复跟进','已回复');self.assertFalse(next(r for r in self.home()['records'] if r['warehouse']=='河北仓')['deferred'])
    def test_exports_scopes_order_values_types_and_originals_unchanged(self):
        self.save();before=self.original();dest=self.root/'exports';dest.mkdir();files=p.export_comparisons(self.path,self.heads,dest,'separate')['files'];combined=p.export_comparisons(self.path,list(reversed(self.heads)),dest,'combined')['files'];self.assertEqual(len(files),2);b=openpyxl.load_workbook(combined[0]);ws=b.active;rows=list(ws.values);self.assertEqual([r[2] for r in rows[1:]],['山东仓','河北仓']);self.assertEqual(rows[2][10],5);self.assertIsNone(rows[1][10]);self.assertEqual(ws.cell(3,11).data_type,'n');b.close();self.assertEqual(before,self.original())
    def test_mapping_profile_template_duplicate_headers_and_numeric_codes(self):
        template=self.root/'template.xlsx';io.template(template);book=openpyxl.load_workbook(template);self.assertEqual(list(book.worksheets[0].values)[0],tuple(io.FIELDS));self.assertIsNone(book.worksheets[0].cell(2,1).value);book.close()
        p.mappings(self.path,'ERP',{'mapping':{'商品编码':0},'headers':['编码']});self.assertEqual(p.mappings(self.path)['ERP']['mapping']['商品编码'],0)
        self.assertEqual(io.code(12345),'12345');self.assertEqual(io.code('0001'),'0001')
        with self.assertRaises(ValueError):io.code(True)
        with self.assertRaises(ValueError):io.number(float('nan'))
    def test_empty_and_formula_without_cache(self):
        with self.assertRaisesRegex(ValueError,'公式没有缓存值'):self.scan([{'本版实际下单数量':'=1+1'}])
        file=self.root/'empty.xlsx';io.template(file);info=io.inspect([file])[0];s=info['sheets'][0]
        with self.assertRaisesRegex(ValueError,'没有采购明细'):io.scan([{**info,**s,'mapping':{f:i for i,f in enumerate(io.FIELDS)}}])

    def test_replay_retains_explicit_conversion_and_manual_links(self):
        scan=self.scan([{'数量单位':'箱','补货单号':None,'签发修订':None}]);reviewed=copy.deepcopy(scan);r=reviewed['rows'][0];r.update(factor='24',conversion_basis='已核对一箱24包',base_quantity='120',links=[p.link(self.h,120)]);self.save(reviewed)
        replay=p.retain_links(scan,p.documents(self.path));self.assertEqual(replay['rows'][0]['base_quantity'],'120');receipt=p.commit(self.path,p.prepare(self.path,replay));self.assertFalse(receipt[0]['changed']);self.assertEqual(p.comparison(self.path,self.h)['rows'][0][4],120)
    def test_revising_links_is_versioned_and_noop_idempotent(self):
        self.save();doc=p.documents(self.path)[0];rows=copy.deepcopy(doc['payload']['rows']);plan=p.revise(self.path,doc['docid'],1,rows);self.assertFalse(plan['plans'][0]['changed']);self.assertEqual(p.commit(self.path,plan)[0]['version'],1)
        rows[0]['links']=[];plan=p.revise(self.path,doc['docid'],1,rows);self.assertEqual(p.commit(self.path,plan)[0]['version'],2);self.assertIsNone(p.comparison(self.path,self.h)['rows'][0][4])
    def test_partial_export_cancel_reports_files_and_no_db_change(self):
        self.save();before=self.original();event=threading.Event();folder=self.root/'partial';folder.mkdir();import xlsx_exports
        original=xlsx_exports.write_table;calls=[]
        def write(*args,**kwargs):
            calls.append(1)
            if len(calls)==2:event.set()
            return original(*args,**kwargs)
        with patch.object(xlsx_exports,'write_table',write):receipt=p.export_comparisons(self.path,self.heads,folder,'separate',cancel=event)
        self.assertTrue(receipt['cancelled']);self.assertEqual(len(receipt['files']),1);self.assertEqual(len(list(folder.glob('*.xlsx'))),1);self.assertEqual(self.original(),before)
    def test_corrupt_cache_rebuild_and_payload_tamper_rejected(self):
        self.home()
        with closing(sqlite3.connect(self.root/'followup.db')) as db,db:db.execute("UPDATE followup_projection SET payload='corrupt'")
        self.assertEqual(len(self.home()['records']),2);self.save()
        with closing(p.connect(self.path)) as db,db:
            row=db.execute('SELECT docid,payload FROM purchase_versions').fetchone();payload=wr.loads(row[1]);payload['rows'][0]['quantity']='999';db.execute('UPDATE purchase_versions SET payload=?',(wr.dumps(payload),))
        with self.assertRaisesRegex(ValueError,'校验失败'):p.documents(self.path)
    def test_trash_cannot_link_or_export_but_original_stays(self):
        import archive_orders as ao
        before=self.original();h=self.home()['records'][0];ao.change(self.path,h,True)
        self.assertEqual(len(self.home()['records']),1);self.assertFalse(any(x['warehouse']=='河北仓' for x in p.targets(self.path)))
        with self.assertRaises(wr.Conflict):p.prepare(self.path,self.scan())
        with self.assertRaises(wr.Conflict):p.export_comparisons(self.path,[self.h],self.root,'separate')
        self.assertEqual(self.original(),before)
    def test_old_feedback_without_warehouse_is_review_not_guessed(self):
        self.reply({'0001':2})
        with closing(sqlite3.connect(self.path)) as db,db:db.execute('UPDATE merchant_feedback SET payload=?',(wr.dumps({'values':{'0001':2}}),))
        self.second_revision();self.assertTrue(all(h['old_revision'] for h in self.home()['records']))
    def test_legacy_unlabelled_warehouse_does_not_block_other_orders(self):
        payload=copy.deepcopy(self.payload);payload['result']['rows'][0]['AU']='旧混合仓';digest=no.fingerprint(payload['result'])
        with closing(sqlite3.connect(self.path)) as db,db:db.execute('UPDATE normal_revisions SET payload=?,digest=?',(no.encode(payload),digest))
        before=self.original();home=self.home()['records'];legacy=next(h for h in home if h['warehouse']=='旧版未标仓');self.assertIn('旧单仓归属待核',legacy['issues']);self.assertEqual(len(p.targets(self.path)),1);self.assertEqual(p.comparison(self.path,legacy)['rows'][0][0],'0001');self.assertEqual(self.original(),before)
    def test_retained_raw_cells_are_integrity_checked(self):
        self.save()
        with closing(p.connect(self.path)) as db,db:
            raw=db.execute('SELECT payload FROM purchase_versions').fetchone()[0];payload=wr.loads(raw);payload['rows'][0]['raw']['商品名称']='损坏的原始值';db.execute('UPDATE purchase_versions SET payload=?',(wr.dumps(payload),))
        with self.assertRaisesRegex(ValueError,'校验失败'):p.documents(self.path)

if __name__=='__main__':unittest.main()
