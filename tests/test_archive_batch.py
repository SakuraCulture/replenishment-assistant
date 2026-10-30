"""Portable batch export tests: isolated original files, dates, stores and revisions."""
import copy,hashlib,sys,tempfile,threading,unittest
from pathlib import Path
from contextlib import closing
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
import openpyxl
from openpyxl.styles import Font
import archive_batch as ab,archive_exports as ae,history_read as hr
import normal_orders as no,work_repository as wr,source_archive,archive_orders as ao
from test_archive_orders import seed

def fixture(root):
    work,repo,p=seed(root)
    with closing(work.connect()) as db:
        catalogs=no.decode(db.execute('SELECT payload FROM normal_work_revisions').fetchone()[0])['catalogs']
    for wh,c in catalogs.items():
        path=Path(c['source_path']);b=openpyxl.load_workbook(path);b.active['A2']='0001' if wh=='河北仓' else '0002';b.active['B2']='=1+2';b.active['C2'].font=Font(bold=True,color='FF0000');b.create_sheet('原件其他页')['A1']='不能丢失';b.save(path);b.close()
        c['source_sha256']=hashlib.sha256(path.read_bytes()).hexdigest();source_archive.capture(work.path,path,c['source_sha256'])
        sku='0001' if wh=='河北仓' else '0002';c['items']={sku:{'商品编码':sku,'_row':2}}
    with closing(work.connect()) as db,db:
        db.execute('UPDATE normal_work_revisions SET payload=?',(no.encode({'catalogs':catalogs}),))
        p2=copy.deepcopy(p);p2['result']['store']='T2';p2['result']['store_name']='另一店';p2['result']['rows'][0]['J']=188
        p2['result']['rows'][0]['_extra_audit']='历史新增字段';p2['result']['config']['parameters']['alpha']=.8
        day='2026-10-28';digest=no.fingerprint(p2['result'])
        db.execute('INSERT INTO normal_orders VALUES(?,?,?,?,?,?)',(day,'T2','normal_model','oid2','BH-T2-ZC',1))
        db.execute('INSERT INTO normal_revisions VALUES(?,?,?,?,?)',('oid2',1,digest,wr.now(),no.encode(p2)))
    heads=hr.execute(work.path,root/'index.db',{'kind':'list'},threading.Event(),lambda _:None)['records']
    return work,repo,heads,catalogs

def cells(path):
    b=openpyxl.load_workbook(path);result=[[(c.value,c.data_type,c.number_format) for c in row] for row in b.active];names=b.sheetnames;b.close();return names,result

class BatchTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.work,self.repo,self.heads,self.catalogs=fixture(self.root);self.cancel=threading.Event()
        self.before=self.work.path.read_bytes()
    def tearDown(self):self.tmp.cleanup()
    def plan(self,kind='calculation',mode='separate',heads=None):return ab.prepare(self.work.path,heads or self.heads,kind,mode)
    def execute(self,plan,**kw):return ab.execute(self.work.path,plan,self.root,**kw)
    def test_separate_and_combined_values_types_order_and_baseline(self):
        for kind,sheet in [('calculation',0),('signed',2)]:
            separate=self.execute(self.plan(kind));combined=self.execute(self.plan(kind,'combined'))
            self.assertEqual(len(separate['files']),4);self.assertEqual(len(combined['files']),2)
            baseline={}
            for f in separate['files']:
                h=f['heads'][0];name,data=cells(f['path']);self.assertEqual(len(name),1)
                target=self.root/(kind+h['oid']+h['warehouse']+'.xlsx');ae.export(self.work.path,h,target)
                b=openpyxl.load_workbook(target);old=[[(c.value,c.data_type,c.number_format) for c in r] for r in b.worksheets[sheet]];b.close()
                self.assertEqual([r[6:] for r in data],old)
                baseline[ab.key(h)]=data
            for f in combined['files']:
                name,actual=cells(f['path']);headers=[c[0] for c in actual[0]];expected=[]
                for h in f['heads']:
                    data=baseline[ab.key(h)];old=[c[0] for c in data[0]]
                    for row in data[1:]:
                        expected.append([row[old.index(c)] if c in old else (None,'n','General') for c in headers])
                self.assertEqual(actual[1:],expected)
        self.assertEqual(self.work.path.read_bytes(),self.before)
    def test_original_byte_identity_and_dedup_associations(self):
        plan=self.plan('original','combined');result=self.execute(plan);self.assertEqual(len(result['files']),2)
        for f in result['files']:
            wh=f['heads'][0]['warehouse'];self.assertEqual(len(f['heads']),2)
            self.assertEqual(Path(f['path']).read_bytes(),Path(self.catalogs[wh]['source_path']).read_bytes())
            b=openpyxl.load_workbook(f['path']);self.assertEqual(len(b.sheetnames),2);self.assertEqual(b.active['B2'].data_type,'f');b.close()
        self.assertEqual(len(self.execute(self.plan('original'))['files']),4)
    def test_reference_is_labelled_derived_and_retains_duplicate_columns(self):
        r=self.execute(self.plan('reference','combined'));self.assertEqual(len(r['files']),2)
        for f in r['files']:
            names,data=cells(f['path']);self.assertIn('加工表',names[0]);headers=[c[0] for c in data[0]]
            self.assertIn('原表第3列 / 重复',headers);self.assertIn('原表第4列 / 重复',headers)
            self.assertEqual(len(data),3);i=headers.index('原表第2列 / 备注')
            self.assertEqual(data[1][i][:2],('=1+2','s'))
    def test_missing_corrupt_source_and_explicit_partial(self):
        with closing(self.work.connect()) as db,db:db.execute('DELETE FROM source_files WHERE sha=?',(self.catalogs['河北仓']['source_sha256'],))
        plan=self.plan('original','combined');self.assertEqual(sum(bool(i['error']) for i in plan['items']),2)
        with self.assertRaises(ValueError):self.execute(plan)
        result=self.execute(plan,allow_partial=True);self.assertEqual(len(result['files']),1);self.assertEqual(len(result['skipped']),2)
        with closing(self.work.connect()) as db,db:db.execute("UPDATE source_files SET content=X'00'")
        self.assertTrue(all(i['error'] for i in self.plan('original')['items']))
    def test_preview_bound_revision_current_update_does_not_replace(self):
        plan=self.plan('signed');h=self.heads[0]
        with closing(self.work.connect()) as db,db:db.execute('UPDATE normal_orders SET current=2 WHERE oid=?',(h['oid'],))
        self.assertEqual(len(self.execute(plan)['files']),4)
    def test_combined_fails_atomic_per_warehouse_other_warehouse_completes(self):
        plan=self.plan('signed','combined');h=next(h for h in self.heads if h['warehouse']=='河北仓');ao.change(self.work.path,h,True)
        result=self.execute(plan);self.assertEqual(len(result['files']),1);self.assertEqual(result['files'][0]['heads'][0]['warehouse'],'山东仓');self.assertEqual(len(result['failed']),2)
        self.assertFalse(list(Path(result['directory']).glob('.archive-*')))
    def test_cancel_before_and_during_then_retry(self):
        plan=self.plan('calculation','combined');self.cancel.set();r=self.execute(plan,cancel=self.cancel);self.assertFalse(r['files']);self.assertEqual(len(r['cancelled']),4)
        self.cancel.clear()
        def progress(message):
            if '压缩保存' in str(message):self.cancel.set()
        r=self.execute(plan,cancel=self.cancel,progress=progress);self.assertFalse(r['files']);self.assertFalse(list(Path(r['directory']).glob('*.xlsx')))
        self.cancel.clear();self.assertEqual(len(self.execute(plan,cancel=self.cancel)['files']),2)
    def test_partial_completion_and_no_overwrite(self):
        plan=self.plan('signed');self.cancel.clear()
        def progress(message):
            if str(message).startswith('文件 2/'):self.cancel.set()
        r=self.execute(plan,cancel=self.cancel,progress=progress);self.assertEqual(len(r['files']),1);self.assertEqual(len(r['cancelled']),3)
        again=self.execute(plan);self.assertNotEqual(r['directory'],again['directory']);self.assertTrue(Path(r['files'][0]['path']).exists())
    def test_capacity_no_truncation_and_duplicate_revision_rejected(self):
        with self.assertRaises(ValueError):self.plan(heads=[self.heads[0],self.heads[0]])
        with patch.object(ab,'MAX_ROWS',2):
            r=self.execute(self.plan('calculation','combined'));self.assertFalse(r['files']);self.assertEqual(len(r['failed']),4)
    def test_reverse_selection_preserved(self):
        heads=list(reversed(self.heads));r=self.execute(self.plan('signed','combined',heads))
        for f in r['files']:self.assertEqual(f['heads'],[h for h in heads if h['warehouse']==f['heads'][0]['warehouse']])
    def test_date_range_and_source_status_and_legacy_form_guard(self):
        listing=hr.execute(self.work.path,self.root/'date-index.db',{'kind':'list','from':'2026-10-28','to':'2026-10-28'},self.cancel,lambda _:None)
        self.assertEqual(sum(listing['matches']),2)
        self.assertTrue(all('已留存' in h['source_status'] for h in listing['records']))
        with self.assertRaises(ValueError):hr.execute(self.work.path,self.root/'date-index.db',{'kind':'list','from':'2026-11-01','to':'2026-10-28'},self.cancel,lambda _:None)
        with self.assertRaisesRegex(ValueError,'资料说明不是原表'):
            hr.execute(self.work.path,self.root/'date-index.db',{'kind':'form_export','head':self.heads[0],'scope':'当时商品资料快照','path':str(self.root/'wrong.xlsx')},self.cancel,lambda _:None)
        self.assertFalse((self.root/'wrong.xlsx').exists())

if __name__=='__main__':unittest.main()
