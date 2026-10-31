"""Self-contained archive tests. Generates its own Excel/SQLite; no Qt or real data."""
import copy,hashlib,sys,tempfile,threading,unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from contextlib import closing
import openpyxl
import archive_orders as ao,archive_exports as ae,normal_orders as no,work_repository as wr
import history_read as hr,history_forms as hf,source_archive,normal_exports as ne
import result_semantics as sem,storage_chunks as chunks
from test_rules import row,META
import model,engine


def seed(folder,chunked=False):
    work=wr.Repository(Path(folder)/'settings');repo=no.Repository(work)
    rows=[]
    for sku,wh,qty in [('0001','河北仓',24),('0002','山东仓',0)]:
        d=model.calculate(row(E=sku,L=0));d.update(AU=wh,_source='资料匹配',_meta=dict(META,export_row=len(rows)+2),_candidate_qty=d['AS'],_candidate_amount=d['BV'],_must_sell=False)
        rows.append(d)
    r={'store':'T','store_name':'测试店','rows':rows,'summary':sem.summarize(rows),'config':engine.config(),'warehouses':engine.DEFAULT_WAREHOUSES,'enabled':True,'budget':{'enabled':False},'sources':{}}
    p={'schema':1,'result':r,'work_revision':'frozen','new_focus':None};h=no.fingerprint(r)
    catalogs={}
    for wh in ('河北仓','山东仓'):
        path=Path(folder)/(wh+'.xlsx');b=openpyxl.Workbook();s=b.active;s.title='原资料';s.append(['商品编码','备注','重复','重复','日期']);s.append(['0001','=保留原文字',None,0,None]);s.append(['原表未入选商品','保留全部行',None,False,None]);b.save(path);b.close()
        sha=hashlib.sha256(path.read_bytes()).hexdigest();source_archive.capture(work.path,path,sha)
        catalogs[wh]={'sheet':'原资料','revision':'frozen-'+wh,'source_path':str(path),'source_sha256':sha,'items':{}}
    with closing(work.connect()) as db,db:
        db.execute('INSERT INTO normal_orders VALUES(?,?,?,?,?,?)',(wr.today(),'T','normal_model','oid','BH-T-ZC',1))
        blob=chunks.put(db,chunks.prepare(p,binding=h)) if chunked else no.encode(p)
        db.execute('INSERT INTO normal_revisions VALUES(?,?,?,?,?)',('oid',1,h,wr.now(),blob))
        db.execute('INSERT INTO normal_work_revisions VALUES(?,?,?)',('frozen',wr.today(),no.encode({'catalogs':catalogs})))
    return work,repo,p


class ArchiveTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.work,self.repo,self.p=seed(self.root);self.cancel=threading.Event()
    def tearDown(self):self.tmp.cleanup()
    def listing(self,**kw):return hr.execute(self.work.path,self.root/'index.db',{'kind':'list',**kw},self.cancel,lambda _:None)
    def test_partition_order_amount_search_and_detail(self):
        records=self.listing()['records'];self.assertEqual([r['warehouse'] for r in records],['河北仓','山东仓']);self.assertEqual(sum(r['summary']['amount'] for r in records),self.p['result']['summary']['amount'])
        self.assertEqual(self.listing(search='0002')['matches'],[False,True])
        v=hr.execute(self.work.path,self.root/'index.db',{'kind':'detail','head':records[1]},self.cancel,lambda _:None)
        self.assertEqual([r[0] for r in v['rows']],['0002']);self.assertEqual(v['head']['number'],'BH-T-ZC-SD')
    def test_delete_restore_preserves_payload_sources_and_other_warehouse(self):
        before=self.repo.get(wr.today(),'T')
        with closing(self.work.connect()) as db:original_sources=db.execute('SELECT * FROM source_files ORDER BY sha').fetchall()
        head=self.listing()['records'][0];ao.change(self.work.path,head,True)
        self.assertEqual(self.repo.get(wr.today(),'T'),before);self.assertEqual([r['warehouse'] for r in self.listing()['records']],['山东仓'])
        trash=self.listing(trash=True)['records'];self.assertEqual(len(trash),1)
        with closing(self.work.connect()) as db,db:
            with self.assertRaises(wr.Conflict):ao.require_active(db,'oid','河北仓')
            ao.require_active(db,'oid','山东仓');self.assertEqual(db.execute('SELECT * FROM source_files ORDER BY sha').fetchall(),original_sources)
        with self.assertRaises(wr.Conflict):ao.change(self.work.path,head,False)
        ao.change(self.work.path,trash[0],False);self.assertEqual(len(self.listing()['records']),2);self.assertEqual(self.repo.get(wr.today(),'T'),before)
    def test_stale_revision_and_cancel_roll_back(self):
        h=self.listing()['records'][0];self.cancel.set()
        with self.assertRaises(Exception):ao.change(self.work.path,h,True,cancel=self.cancel)
        with closing(self.work.connect()) as db,db:self.assertEqual(ao.states(db),{});db.execute('UPDATE normal_orders SET current=2')
        with self.assertRaises(wr.Conflict):ao.change(self.work.path,h,True)
        with closing(self.work.connect()) as db,db:self.assertEqual(ao.states(db),{})
    def test_three_sheets_raw_source_and_signed_rows(self):
        before=copy.deepcopy(self.repo.get(wr.today(),'T'));h=self.listing()['records'][0];path=self.root/'archive.xlsx';ae.export(self.work.path,h,path)
        b=openpyxl.load_workbook(path);self.assertEqual(b.sheetnames,ae.SHEETS)
        calc=list(b.worksheets[0].values);self.assertEqual(len(calc),2);self.assertIn('最终签发执行数',calc[0]);self.assertEqual(calc[1][-3],sem.order_quantity(self.p['result']['rows'][0]))
        raw=list(b.worksheets[1].values);self.assertEqual(raw[0][:4],('商品编码','备注','重复','重复'));self.assertEqual(raw[1][:4],('0001','=保留原文字',None,0));self.assertEqual(raw[2][0],'原表未入选商品')
        self.assertEqual(b.worksheets[1]['B2'].data_type,'s');signed=list(b.worksheets[2].values);self.assertEqual(signed[1][5],'0001');self.assertEqual(signed[1][-2],'BH-T-ZC-HB');b.close()
        self.assertEqual(self.repo.get(wr.today(),'T'),before)
        with self.assertRaises(FileExistsError):ae.export(self.work.path,h,path)
    def test_missing_source_explicit_and_cancel_no_file(self):
        h=self.listing()['records'][0]
        with closing(self.work.connect()) as db,db:db.execute('DELETE FROM source_files')
        path=self.root/'missing.xlsx';ae.export(self.work.path,h,path);b=openpyxl.load_workbook(path);self.assertIn('未留存',b.worksheets[1]['A2'].value);b.close()
        self.cancel.set()
        with self.assertRaises(Exception):ae.export(self.work.path,h,self.root/'cancel.xlsx',cancel=self.cancel)
        self.assertFalse((self.root/'cancel.xlsx').exists())
    def test_deleted_export_blocked_and_event_per_warehouse(self):
        h=self.listing()['records'][0];ao.change(self.work.path,h,True)
        with self.assertRaises(wr.Conflict):ae.export(self.work.path,h,self.root/'blocked.xlsx')
        events=hf.events(self.work.path,'oid','河北仓');self.assertEqual(len(events['rows']),1);self.assertEqual(events['rows'][0][1],'BH-T-ZC-HB')
        self.assertFalse(hf.events(self.work.path,'oid','山东仓')['records'])
    def test_chunk_storage_identical_view(self):
        dest=self.root/'chunk';dest.mkdir();work,repo,p=seed(dest,True)
        self.assertEqual(wr.dumps(repo.get(wr.today(),'T')['payload']['result']),wr.dumps(self.p['result']))
        data=hr.execute(work.path,dest/'index.db',{'kind':'list'},self.cancel,lambda _:None);self.assertEqual(len(data['records']),2)
    def test_exception_isolation_survives_python_optimization(self):
        d=row(L=0);d.update(_source='资料匹配',AU='河北仓');loaded={'normalized':[(d,dict(META,export_row=2))],'config':engine.config(),'stores':{'T':'测试店'},'hashes':{}}
        def wrong(b,m,enabled=True):out=copy.deepcopy(b);out['J']+=1;return out,{'applied':False}
        with patch.object(engine,'apply_one_box_exception',wrong):
            with self.assertRaisesRegex(RuntimeError,'隔离检查失败'):engine.run_calculation(loaded,'T',engine.DEFAULT_WAREHOUSES,True)

    def test_deleted_confirm_and_feedback_commit_reject_without_mutation(self):
        import merchant_feedback as mf
        from version_info import MODEL_VERSION
        before=self.repo.get(wr.today(),'T');file=self.root/'return.xlsx';ne.export(before,file,warehouse='河北仓')
        plan=mf.prepare(self.work.path,mf.scan([file]),{})
        head=self.listing()['records'][0];ao.change(self.work.path,head,True)
        payload={'schema':1,'model_version':MODEL_VERSION,'day':wr.today(),'input_dirty':False,'states':{'T':{'name':'测试店','state':'已完成'}},'results':{'T':self.p['result']},'new_focus':None}
        with self.assertRaises(wr.Conflict):self.repo.confirm(payload,['T'],None,{'T':1})
        with self.assertRaises(wr.Conflict):mf.commit(self.work.path,plan)
        self.assertIsNone(self.work.get_job(wr.today()));self.assertEqual(self.repo.get(wr.today(),'T'),before)
        trash=self.listing(trash=True)['records'][0];ao.change(self.work.path,trash,False)
        mf.commit(self.work.path,plan);self.assertEqual(self.repo.get(wr.today(),'T'),before)

    def test_source_hash_tamper_and_no_partial_export(self):
        h=self.listing()['records'][0]
        with closing(self.work.connect()) as db,db:db.execute("UPDATE source_files SET content=X'00'")
        with self.assertRaisesRegex(ValueError,'校验失败'):ae.export(self.work.path,h,self.root/'tampered.xlsx')
        self.assertFalse((self.root/'tampered.xlsx').exists())

    def test_empty_fields_and_formula_text_are_not_dropped(self):
        self.assertEqual(list(ae.flatten({'x':{},'y':[],'z':None,'n':0},())),[(('x',),'空记录'),(('y',),'空清单'),(('z',),None),(('n',),0)])
        with self.assertRaisesRegex(ValueError,'上限'):list(ae.checked_rows([['x'*32768]]))

if __name__=='__main__':unittest.main()
