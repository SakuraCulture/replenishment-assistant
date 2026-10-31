"""Portable full-document lifecycle fixtures. No Qt, real Excel or user settings."""
import copy,hashlib,sqlite3,sys,tempfile,threading,time,unittest
from pathlib import Path
from contextlib import closing
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
import openpyxl
import lifecycle_io as io,order_lifecycle as lc,lifecycle_home as home
import purchases,work_repository as wr,normal_orders as no,archive_orders as ao
from runtime_support import Cancelled
from test_archive_orders import seed

class LifecycleTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.work,self.repo,self.payload=seed(self.root);self.path=self.work.path;self.heads=purchases.targets(self.path);self.h=self.heads[0];self.count=0
    def tearDown(self):self.tmp.cleanup()
    def scan(self,stage='return',items=None):
        self.count+=1;p=self.root/f'{stage}-{self.count}.xlsx';book=openpyxl.Workbook();ws=book.active;ws.title='业务明细';ws.append(io.fields(stage))
        for item in ([{}] if items is None else items):
            r={'补货单号':self.h['number'],'签发修订':1,'门店编码':'T','供货仓':'河北仓','商品编码':'0001','商品名称':'测试商品','数量单位':'库存单位',io.QUANTITIES[stage]:5,io.DOCNUM[stage]:'D001',io.DOCDATE[stage]:wr.today(),'实际金额':12,'币种':'CNY',**item};ws.append([r.get(f) for f in io.fields(stage)])
        book.save(p);book.close();f=io.inspect([p])[0];s=f['sheets'][0];return io.scan(stage,[{**f,**s,'mapping':{h:i for i,h in enumerate(s['headers'])}}])
    def plan(self,scan=None):
        scan=scan or self.scan();heads=lc.available(self.path,scan);bindings={lc.binding_key(r):lc.matching(r,heads)[0] for r in scan['rows']};return lc.prepare(self.path,scan,bindings)
    def save(self,stage='return',items=None):
        plan=self.plan(self.scan(stage,items));return plan,lc.commit(self.path,plan)
    def old_tables(self):
        with closing(sqlite3.connect(self.path)) as db:
            # Existing backup mechanism may create these two storage metadata tables.
            # They are not signed/job facts. All original business tables stay exact.
            names=[x[0] for x in db.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'lifecycle_%' AND name NOT IN ('snapshot_chunks','snapshot_storage_meta')")]
            return {n:db.execute('SELECT * FROM "'+n+'"').fetchall() for n in names}
    def read(self,**kw):return home.read(self.path,self.root/'home.db',kw)
    def view(self,h=None):return lc.snapshot(self.path,h or self.h)
    def test_four_stages_and_original_immutable(self):
        before=self.old_tables();self.assertEqual(self.view()['status'],'待回传')
        for stage,status in [('return','待付款'),('payment','待收货'),('receipt','已收齐')]:
            self.save(stage);self.assertEqual(self.view()['status'],status)
        self.assertEqual(self.view()['rows'][0][3:6],[5.,5.,5.]);self.assertEqual(self.old_tables(),before)
        self.assertEqual(next(r for r in self.read()['records'] if r['warehouse']=='山东仓')['status'],'待回传')
    def test_out_of_order_documents_next_missing_not_highest_stage(self):
        self.save('receipt');self.assertEqual(self.view()['next'],'return');self.assertTrue(self.view()['present']['receipt']);self.save('return');self.assertEqual(self.view()['next'],'payment');self.save('payment');self.assertEqual(self.view()['next'],'receipt')
        lc.commit(self.path,lc.recheck(self.path,self.h,'receipt'));self.assertIsNone(self.view()['next'])
    def test_quantity_differences_do_not_make_partial_stage(self):
        self.save('return',[{'门店确认数量':0}]);self.save('payment',[{'付款单商品数量':8}]);self.save('receipt',[{'实际收货数量':3}]);v=self.view();self.assertEqual(v['status'],'已收齐');self.assertEqual(v['rows'][0][3:6],[0,8,3]);self.assertEqual(v['rows'][0][7:9],[8,-5])
    def test_missing_sku_not_zero_and_extra_kept(self):
        self.save('payment',[{'商品编码':'EXTRA','付款单商品数量':9}]);v=self.view();self.assertIsNone(v['rows'][0][4]);self.assertEqual(v['rows'][-1][0],'EXTRA');self.assertEqual(v['rows'][-1][4],9);self.assertTrue(v['present']['payment']);self.assertEqual(v['status'],'待回传')
    def test_blank_quantity_zero_effect_preserves_raw(self):
        for stage in io.STAGES:
            with self.subTest(stage=stage):
                plan=self.plan(self.scan(stage,[{io.QUANTITIES[stage]:None}]))
                self.assertIsNone(plan['plans'][0]['payload']['raw_rows'][0]['quantity'])
                self.assertEqual(plan['plans'][0]['payload']['rows'][0]['base_quantity'],'0')
                lc.commit(self.path,plan)
        self.assertEqual(self.view()['rows'][0][3:6],[0,0,0])
    def test_missing_paid_amount_rejected_currency_not_invented(self):
        with self.assertRaisesRegex(ValueError,'空白'):self.scan('payment',[{'实际金额':None}])
        self.assertEqual(self.scan('payment',[{'币种':None}])['rows'][0]['currency'],'')
    def test_payment_uses_source_amount_not_model_price(self):
        self.save('payment',[{'付款单商品数量':7,'实际金额':0.99,'币种':'USD'}]);v=self.view();self.assertEqual(v['rows'][0][9:11],['0.99','USD'])
    def test_exact_codes_and_units(self):
        self.save('receipt',[{'实际收货数量':2,'数量单位':'箱','折合库存单位系数':24,'换算依据':'源单标明24包/箱'}]);self.assertEqual(self.view()['rows'][0][5],48);self.assertEqual(self.view()['rows'][0][0],'0001')
        with self.assertRaisesRegex(ValueError,'换算系数'):self.scan('payment',[{'数量单位':'箱'}])
        with self.assertRaisesRegex(ValueError,'只能为1'):self.scan('return',[{'折合库存单位系数':24,'换算依据':'错误单位'}])
    def test_repeat_no_accumulation_no_revision(self):
        plan,first=self.save();scan={'stage':'return','rows':[],'files':plan['files']}
        with closing(purchases.connect(self.path)) as db:old=lc.current(db,self.h,'return')
        again=self.plan(io.scan('return',[{**io.inspect([plan['files'][0]['path']])[0],**io.inspect([plan['files'][0]['path']])[0]['sheets'][0],'mapping':{h:i for i,h in enumerate(io.fields('return'))}}]));receipt=lc.commit(self.path,again);self.assertFalse(receipt[0]['changed']);self.assertEqual(receipt[0]['version'],1);self.assertEqual(len(lc.history(self.path,self.h,'return')),1)
    def test_new_complete_version_replaces_old_not_accumulates(self):
        self.save();self.save('return',[{'门店确认数量':2}]);self.assertEqual(self.view()['rows'][0][3],2);history=lc.history(self.path,self.h,'return');self.assertEqual([d['version'] for d in history],[2,1]);self.assertEqual(history[1]['payload']['rows'][0]['base_quantity'],'5')
    def test_duplicate_same_row_dedup_conflict_reject(self):
        plan=self.plan(self.scan(items=[{},{}]));self.assertEqual(plan['plans'][0]['duplicates'],1);lc.commit(self.path,plan);self.assertEqual(self.view()['rows'][0][3],5)
        with self.assertRaisesRegex(ValueError,'冲突'):self.plan(self.scan(items=[{},{'门店确认数量':6}]))
    def test_multiple_external_numbers_not_silently_merged(self):
        with self.assertRaisesRegex(ValueError,'不同单号'):self.plan(self.scan('payment',[{}, {'付款单号':'D002','商品编码':'0003'}]))
    def test_cross_warehouse_and_wrong_revision_reject(self):
        scan=self.scan();bindings={lc.binding_key(scan['rows'][0]):self.heads[1]}
        with self.assertRaisesRegex(ValueError,'不一致'):lc.prepare(self.path,scan,bindings)
        scan['rows'][0]['revision']=2
        with self.assertRaisesRegex(ValueError,'不一致'):lc.prepare(self.path,scan,{lc.binding_key(scan['rows'][0]):self.h})
    def test_missing_revision_requires_binding(self):
        scan=self.scan(items=[{'签发修订':None}])
        with self.assertRaisesRegex(ValueError,'未明确'):lc.prepare(self.path,scan,{})
        lc.commit(self.path,lc.prepare(self.path,scan,{lc.binding_key(scan['rows'][0]):self.h}));self.assertEqual(self.view()['status'],'待付款')
    def test_batch_order_preserved_and_atomic(self):
        h=self.heads[1];scan=self.scan('payment',[{'补货单号':h['number'],'供货仓':h['warehouse'],'商品编码':'0002'},{}]);plan=self.plan(scan);self.assertEqual([p['head']['warehouse'] for p in plan['plans']],['山东仓','河北仓']);lc.commit(self.path,plan);self.assertEqual([r['warehouse'] for r in self.read()['records']],['河北仓','山东仓'])
    def test_preview_stale_rejected_whole_batch(self):
        h=self.heads[1];scan=self.scan(items=[{}, {'补货单号':h['number'],'供货仓':h['warehouse'],'商品编码':'0002'}]);plan=self.plan(scan);self.save()
        with self.assertRaises(wr.Conflict):lc.commit(self.path,plan)
        self.assertEqual(lc.history(self.path,h,'return'),[])
    def second_revision(self):
        payload=copy.deepcopy(self.payload);payload['result']['revision_marker']='second';digest=no.fingerprint(payload['result'])
        with closing(sqlite3.connect(self.path)) as db,db:db.execute('INSERT INTO normal_revisions VALUES(?,?,?,?,?)',('oid',2,digest,wr.now(),no.encode(payload)));db.execute('UPDATE normal_orders SET current=2')
    def test_new_signing_no_auto_transfer_old_documents(self):
        self.save('payment');self.second_revision();heads=purchases.targets(self.path);h=next(h for h in heads if h['warehouse']=='河北仓' and h['revision']==2);self.assertFalse(self.view(h)['present']['payment']);self.assertTrue(self.view()['present']['payment']);self.assertTrue(next(r for r in self.read()['records'] if r['warehouse']=='河北仓')['old_revision'])
    def test_signing_changes_after_preview_reject(self):
        plan=self.plan();self.second_revision()
        with self.assertRaises(wr.Conflict):lc.commit(self.path,plan)
    def test_trash_restore_token_invalidates_preview(self):
        plan=self.plan();ao.change(self.path,self.h,True)
        with self.assertRaises(wr.Conflict):lc.commit(self.path,plan)
    def test_cancel_before_commit_and_retry(self):
        event=threading.Event();event.set();before=self.old_tables()
        with self.assertRaises(Cancelled):lc.prepare(self.path,self.scan(),{},cancel=event)
        event.clear();plan=self.plan();event.set()
        with self.assertRaises(Cancelled):lc.commit(self.path,plan,cancel=event)
        self.assertEqual(lc.history(self.path,self.h,'return'),[]);self.assertEqual(self.old_tables(),before);event.clear();self.assertEqual(lc.commit(self.path,plan,cancel=event)[0]['version'],1)
    def test_cancel_waiting_write_lock(self):
        plan=self.plan();event=threading.Event();errors=[];db=sqlite3.connect(self.path);db.execute('BEGIN IMMEDIATE')
        def writer():
            try:lc.commit(self.path,plan,cancel=event)
            except Exception as exc:errors.append(exc)
        with patch('storage_chunks.backup_before_write'):
            task=threading.Thread(target=writer);task.start();time.sleep(.2);event.set();task.join(3)
        db.rollback();db.close();self.assertFalse(task.is_alive());self.assertIsInstance(errors[0],Cancelled);self.assertEqual(lc.history(self.path,self.h,'return'),[])
    def test_commit_boundary_finishes_atomically(self):
        event=threading.Event();plan=self.plan()
        def progress(s):
            if s==lc.COMMIT_BOUNDARY:event.set()
        receipt=lc.commit(self.path,plan,progress,event);self.assertEqual(receipt[0]['version'],1)
    def test_sql_failure_rolls_back_entire_batch(self):
        h=self.heads[1];plan=self.plan(self.scan(items=[{}, {'补货单号':h['number'],'供货仓':h['warehouse'],'商品编码':'0002'}]))
        with closing(sqlite3.connect(self.path)) as db,db:
            lc.schema(db);db.execute("CREATE TRIGGER fail_second BEFORE INSERT ON lifecycle_versions WHEN NEW.warehouse='山东仓' BEGIN SELECT RAISE(ABORT,'injected'); END")
        with self.assertRaisesRegex(sqlite3.IntegrityError,'injected'):lc.commit(self.path,plan)
        self.assertEqual(lc.history(self.path,self.h,'return'),[])
        with closing(sqlite3.connect(self.path)) as db:self.assertEqual(db.execute('SELECT COUNT(*) FROM lifecycle_files').fetchone()[0],0)
    def test_source_change_before_commit_reject(self):
        plan=self.plan();p=Path(plan['files'][0]['path']);p.write_bytes(p.read_bytes()+b'changed')
        with self.assertRaisesRegex(ValueError,'变化'):lc.commit(self.path,plan)
    def test_tampered_payload_reject_even_optimized_python(self):
        plan=self.plan();plan['plans'][0]['payload']['rows'][0]['base_quantity']='100'
        with self.assertRaisesRegex(ValueError,'预览内容'):lc.commit(self.path,plan)
        self.save()
        with closing(sqlite3.connect(self.path)) as db,db:db.execute("UPDATE lifecycle_versions SET digest='wrong'")
        with self.assertRaisesRegex(ValueError,'校验'):self.view()
    def test_original_download_exact_no_overwrite(self):
        plan,receipt=self.save('payment');f=plan['files'][0];out=self.root/'copy.xlsx';lc.attachment(self.path,self.h,'payment',1,f['sha'],out);self.assertEqual(out.read_bytes(),Path(f['path']).read_bytes())
        with self.assertRaises(FileExistsError):lc.attachment(self.path,self.h,'payment',1,f['sha'],out)
    def test_old_notes_and_purchase_not_payment(self):
        with closing(sqlite3.connect(self.path)) as db,db:
            purchases.schema(db);db.execute('INSERT INTO followup_notes VALUES(?,?,?,?,?)',('oid',1,'河北仓','t',wr.dumps({'action':'暂缓跟进','reason':'旧记录'})));db.execute('INSERT INTO purchase_link_index VALUES(?,?,?,?,?,?,?)',('doc','line','oid',1,'河北仓','0001','5'))
        before=self.old_tables();self.assertEqual(self.view()['status'],'待回传');self.assertFalse(self.view()['present']['payment']);r=self.read();self.assertEqual(r['counts']['待回传'],2);self.assertEqual(self.old_tables(),before)
    def test_old_complete_return_readonly_compatible(self):
        with closing(sqlite3.connect(self.path)) as db,db:
            db.execute('CREATE TABLE merchant_feedback(oid TEXT,revision INTEGER,version INTEGER,saved TEXT,payload TEXT)');db.execute('INSERT INTO merchant_feedback VALUES(?,?,?,?,?)',('oid',1,1,wr.now(),wr.dumps({'values':{'0001':0},'warehouses':{}})))
        before=self.old_tables();self.assertEqual(self.view()['status'],'待付款');self.assertEqual(self.view()['rows'][0][3],0);self.assertEqual(self.old_tables(),before);self.save('return');self.assertEqual(self.view()['rows'][0][3],5);self.assertEqual(self.old_tables(),before)
    def test_unknown_missing_quantity_does_not_complete_legacy(self):
        with closing(sqlite3.connect(self.path)) as db,db:
            db.execute('CREATE TABLE merchant_feedback(oid TEXT,revision INTEGER,version INTEGER,saved TEXT,payload TEXT)');db.execute('INSERT INTO merchant_feedback VALUES(?,?,?,?,?)',('oid',1,1,wr.now(),wr.dumps({'values':{}})))
        self.assertEqual(self.view()['status'],'待回传');self.assertTrue(self.read()['records'][0]['legacy_incomplete'])
    def test_template_formula_and_empty_inputs(self):
        for stage in io.STAGES:
            p=self.root/(stage+'template.xlsx');io.template(stage,p);book=openpyxl.load_workbook(p);self.assertEqual([io.field_key(h) for h in next(book.active.values)],io.PAYMENT_HEADERS if stage=='payment' else io.input_fields(stage));book.close()
        with self.assertRaisesRegex(ValueError,'缓存'):self.scan('return',[{'门店确认数量':'=1+1'}])
        with self.assertRaisesRegex(ValueError,'没有商品'):self.scan(items=[])
    def test_signed_confirmation_export_can_be_read(self):
        import normal_exports
        order=self.repo.get(wr.today(),'T');path=self.root/'signed.xlsx';normal_exports.export(order,path,warehouse='河北仓')
        info=io.inspect([path])[0];specs=[]
        for s in info['sheets']:
            if '商品编码' not in s['headers']:continue
            mapping={f:(s['headers'].index(f) if f in s['headers'] else None) for f in io.fields('return')};mapping['门店确认数量']=s['headers'].index('补货数量');mapping['供货仓']=s['headers'].index('所属仓');specs.append({**info,**s,'mapping':mapping,'defaults':{'数量单位':'库存单位'}})
        scan=io.scan('return',specs);plan=self.plan(scan);lc.commit(self.path,plan);self.assertEqual(self.view()['rows'][0][3],self.view()['rows'][0][2])
    def test_export_scope_order_types_and_reimport(self):
        h=self.heads[1];self.save('payment',[{'补货单号':h['number'],'供货仓':h['warehouse'],'商品编码':'0002'},{}]);before=self.old_tables()
        result=lc.export(self.path,[h,self.h],self.root,'payment','combined');self.assertIsNone(result['error']);self.assertEqual(len(result['files']),2)
        b=openpyxl.load_workbook(result['files'][0]);values=list(b.active.values);self.assertEqual(values[1][2],'山东仓');self.assertEqual(values[1][7],'0002');self.assertIsInstance(values[1][9],(int,float));self.assertNotIn('付款日期',values[0]);b.close();self.assertEqual(self.old_tables(),before)
        missing=lc.export(self.path,[self.h],self.root,'receipt');self.assertFalse(missing['files']);self.assertIn('尚未导入',missing['error'])
    def test_mapping_profiles_separate_stages(self):
        lc.mappings(self.path,'payment','ERP',{'headers':['SKU'],'mapping':{'商品编码':0},'defaults':{}});self.assertIn('ERP',lc.mappings(self.path,'payment'));self.assertEqual(lc.mappings(self.path,'receipt'),{})
    def test_count_partition_filter_dates_and_order(self):
        self.save('return');data=self.read();self.assertEqual(sum(data['counts'].values()),2);self.assertEqual(data['counts']['待付款'],1);self.assertEqual(self.read(status='待付款')['records'][0]['warehouse'],'河北仓');self.assertEqual(self.read(search='山东 T')['count'],1)
        with self.assertRaises(ValueError):self.read(**{'from':'2026-10-31','to':'2026-01-01'})

if __name__=='__main__':unittest.main()
