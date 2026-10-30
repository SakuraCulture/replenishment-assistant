"""Minimum accepted headers; examples never mix with imported business records."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
import os,tempfile

TEMPLATES={
 'sales':('补货建议','商品补货建议',['仓库/门店编码','商品编码','日销','周销','月销','可用库存','在途库存','缺货天数（天）'],['仓库/门店名称','商品名称','单位类型','补货状态']),
 'catalog':('分仓商品资料','Sheet1',['商品编码','分类','基本售价','出库箱规','商品状态','可用数'],['商品名称','商品等级','补货模型分类']),
 'legacy':('旧版混合商品资料','Sheet1',['商品编码','分类','基本售价','出库箱规','商品状态'],['商品名称','仓归属']),
 'priority':('必卖清单','必卖清单',['商品编码'],['商品名称','分类']),
 'focus':('新品关注清单','新品明细',['商品编码','关注'],['商品名称']),
 'feedback':('门店确认回传','门店确认回传',['补货单号','商品编码','门店确认数量'],['商品名称','备注'])}

def write(kind,path,progress=lambda _:None,cancel=None):
    title,sheet,required,optional=TEMPLATES[kind];target=Path(path)
    if target.exists():raise FileExistsError('模板目标已存在，请另存为新文件')
    wb=Workbook();ws=wb.active;ws.title=sheet;ws.append(required+optional);ws.freeze_panes='A2'
    for i,c in enumerate(ws[1]):
        c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='126C63' if i<len(required) else '657B83');ws.column_dimensions[c.column_letter].width=22
        for row in range(2,202):ws.cell(row,i+1).number_format='@' if ('编码' in c.value or c.value=='补货单号') else 'General'
    note=wb.create_sheet('填写说明');note.append(['事项','要求'])
    for item in [('适用内容',title),('必需表头','、'.join(required)),('可选表头','、'.join(optional)),('编码','按文本保存；保留前导零，禁止科学计数法和模糊匹配'),('数量单位','库存单位；回传数量为非负数字，0表示明确不补，空白表示未确认'),('缺失值','保留空白，不能为了通过校验编造0；缺货天数空白按现有规则处理'),('分仓','在导入入口选择河北仓或山东仓；不要以文件名猜测仓库'),('关注','填写关注/未关注、是/否、1/0'),('回传关联','无修订信息的文件，导入时必须明确选择签发修订；也可直接改加盟商确认单Q列回传'),('数据区','第1张表从第2行填写；模板不附可误导入的示例商品')]:note.append(item)
    note.column_dimensions['A'].width=22;note.column_dimensions['B'].width=100
    fd,tmp=tempfile.mkstemp(dir=target.parent,suffix='.xlsx');os.close(fd)
    try:wb.save(tmp);os.link(tmp,target)
    finally:wb.close();Path(tmp).unlink(missing_ok=True)
    return str(target)
