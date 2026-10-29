"""读取门店销售库存导出和商品资料，输出可追溯的 Python 补货结果。"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from dataclasses import asdict
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import sys
import unicodedata
import warnings

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from model import VERSION, Rules, calculate, number, blank, turnover

HERE = Path(__file__).resolve().parent


def digest(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        for block in iter(lambda:f.read(1024*1024),b''):h.update(block)
    return h.hexdigest()


def numeric(value):
    """Only parse designated numerical columns. Slash/missing/invalid stay distinct."""
    if number(value):return value
    if value is None:return None
    if isinstance(value,str):
        value=value.strip()
        if not value:return None
        if re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)',value):
            n=float(value)
            return n if number(n) else value
    return value


def identifier(value):
    """Exact identifiers; never numeric joins, fuzzy names or wildcard matching."""
    if isinstance(value,str):return (value,None) if value.strip() else ('','编码缺失')
    if number(value) and float(value).is_integer() and abs(value)<10**15:
        return str(int(value)),None
    if blank(value):return '', '编码缺失'
    return str(value),'编码以数值保存且超过15位或格式无效，需原始文本编码'


def read_table(path, sheet, required):
    warnings.filterwarnings('ignore',message='Workbook contains no default style')
    book=openpyxl.load_workbook(path,read_only=True,data_only=True,keep_links=False)
    try:
        if sheet is not None:
            if sheet not in book.sheetnames:raise ValueError(f'找不到工作表：{sheet}')
            ws=book[sheet]
        elif len(book.sheetnames)==1:
            ws=book.active
        else:
            raise ValueError('多工作表文件必须在配置中指定 sheet，不能默认选错表')
        # The supplied export says A1:A1 although it contains 15,967 rows.
        ws.reset_dimensions()
        it=ws.iter_rows(values_only=True)
        header=next(it,None)
        if not header:raise ValueError('文件为空')
        header=[str(x).strip() if x is not None else '' for x in header]
        positions=defaultdict(list)
        for i,h in enumerate(header):positions[h].append(i)
        missing=[h for h in required if h not in positions]
        repeated=[h for h in required if len(positions.get(h,[]))>1]
        if missing:raise ValueError('缺少必需表头：'+','.join(missing))
        if repeated:raise ValueError('必需表头重复：'+','.join(repeated))
        rows=[]
        for rn,values in enumerate(it,2):
            if not any(x is not None and x!='' for x in values):continue
            # Duplicate unused headers (three 毛利 columns) are not business inputs.
            row={h:values[idx[0]] if idx[0]<len(values) else None for h,idx in positions.items() if len(idx)==1}
            row['_row']=rn;rows.append(row)
        return rows,{'sheet':ws.title,'rows':len(rows),'headers':header,
                     'duplicate_unused_headers':[h for h,idx in positions.items() if len(idx)>1]}
    finally:book.close()


def load_config(path):
    config=json.loads(Path(path).read_text(encoding='utf-8-sig'))
    rules=Rules(**config['parameters'])
    if not rules.valid():raise ValueError('规则参数无效；请检查 config.json')
    if config['interface']['blank_outbound_unit'] not in ('one','pending'):
        raise ValueError('blank_outbound_unit 必须为 one 或 pending')
    return config,rules


def prepare(sales_path, products_path, config):
    sf,pf=config['sales_fields'],config['product_fields']
    sr,sinfo=read_table(sales_path,config['interface']['sales_sheet'],
                       [sf[c] for c in ['store_code','sku','daily14','sales7','sales30','stock','inbound','stockout_days']])
    pr,pinfo=read_table(products_path,config['interface']['product_sheet'],
                       [pf[c] for c in ['sku','category','outbound_unit','price','supply_status']])
    for mapping,info in [(sf,sinfo),(pf,pinfo)]:
        ambiguous=set(mapping.values()) & set(info['duplicate_unused_headers'])
        if ambiguous:raise ValueError('已映射字段表头重复：'+','.join(sorted(ambiguous)))
    index=defaultdict(list);master_issues=[]
    for r in pr:
        key,err=identifier(r.get(pf['sku']))
        if err:
            master_issues.append([r['_row'],key,err]);continue
        index[key].append(r)
    keys=[];key_errors={}
    for r in sr:
        store,e1=identifier(r.get(sf['store_code']));sku,e2=identifier(r.get(sf['sku']))
        key=(store,sku);keys.append(key)
        if e1 or e2:key_errors[r['_row']]='；'.join(e for e in (e1,e2) if e)
    keycounts=Counter(keys)
    normalized=[];counts=Counter()
    smap={'A':'store_name','B':'store_code','C':'supplier_code','D':'name','E':'sku','F':'barcode','G':'unit_type',
          'H':'daily14','I':'sales7','J':'sales30','K':'stock','L':'inbound','M':'inbound_days','N':'stockout_days','T':'tags'}
    for raw,key in zip(sr,keys):
        d={c:raw.get(sf[k]) for c,k in smap.items()}
        d['B'],d['E']=key
        for c in ['H','I','J','K','L','M','N']:d[c]=numeric(d[c])
        d['P']=turnover(*(d[c] for c in 'H I J K L'.split()))
        d.update(AT='未提供（不作为数量拦截）',AU='未提供',AV='未提供')
        notes=[];blocking=[];record=None
        if raw['_row'] in key_errors:blocking.append(key_errors[raw['_row']])
        if keycounts[key]>1:blocking.append('导出门店商品编码重复，未汇总或择一')
        matches=index.get(key[1],[]) if not blocking else []
        if len(matches)>1:
            blocking.append('商品资料编码重复，无法唯一匹配');match='资料重复'
        elif not matches:
            blocking.append('待补商品资料');match='资料未匹配'
        else:
            record=matches[0];match='商品编码精确匹配'
        if record:
            cat=record.get(pf['category'])
            profile_name=config['category_profiles'].get(str(cat))
            # An explicitly supplied profile column is an optional integration seam.
            override=pf.get('profile_override')
            if override and not blank(record.get(override)):profile_name=str(record[override]).strip()
            profile=config['weight_profiles'].get(profile_name)
            d['U']=profile_name or '待核分类'
            if profile:
                d.update(AB=profile['alpha'],AC=profile['w1'],AD=profile['w2'])
            else:
                blocking.append('未配置的商品分类：'+str(cat));d.update(AB=None,AC=None,AD=None)
            unit=numeric(record.get(pf['outbound_unit']))
            if blank(unit):
                if config['interface']['blank_outbound_unit']=='one':
                    unit='/';notes.append('资料出库箱规空白，按已确认口径视为无约束')
                else:notes.append('出库箱规空白，待核；未使用标准装箱数量')
            d['O']=unit
            d['V']=numeric(record.get(pf['price']))
            d['S']=record.get(pf['supply_status'])
            d['AV']=record.get(pf['grade']) if pf.get('grade') else '未提供'
            d['AU']='商品资料对应供货方'
            d['_category']=cat
        else:
            d.update(O=None,V=None,S=None,U=None,AB=None,AC=None,AD=None)
        for c,label in [('H','日销/14天日均'),('I','周销'),('J','月销'),('K','可用库存'),('L','在途库存')]:
            if not number(d[c]):blocking.append(label+'缺失或格式无效')
        if not(blank(d['N']) or number(d['N'])):blocking.append('缺货天数格式无效')
        if raw.get(sf.get('replenishment_status'))=='停采':notes.append('导出补货状态为停采；原模型未将此状态作为自动停补门槛，需业务核验')
        meta={'export_row':raw['_row'],'product_row':record['_row'] if record else None,
              'match':match,'blocking':list(dict.fromkeys(blocking)),'notes':notes,
              'category':record.get(pf['category']) if record else None,
              'export_status':raw.get(sf.get('replenishment_status')),
              'export_updated':raw.get(sf.get('updated'))}
        normalized.append((d,meta));counts[match]+=1
    return normalized,{'sales':sinfo,'products':pinfo,'matching':dict(counts),
                       'invalid_master_codes':master_issues,
                       'duplicate_master_keys':{k:[r['_row'] for r in rs] for k,rs in index.items() if len(rs)>1},
                       'duplicate_sales_keys':[list(k) for k,n in keycounts.items() if n>1]}


def execute(normalized,rules):
    results=[]
    for inp,meta in normalized:
        if meta['blocking']:
            # Report all SKU rows; the interface does not manufacture data to force zero.
            d=calculate(inp,rules)
            d.update(AS='待补商品资料' if '待补商品资料' in meta['blocking'] else '输入待核',
                     BM='；'.join(meta['blocking']),BV='数量未定',BW='；'.join(meta['blocking']))
        else:
            try:d=calculate(inp,rules)
            except (ArithmeticError,ValueError,OverflowError) as exc:
                # Unexpected arithmetic cannot become a healthy 0 or disappear.
                d=dict(inp);d.update(AS='计算异常',BM=str(exc),BV='数量未定',BW='计算异常',AW='计算异常')
                meta['blocking'].append('计算异常：'+str(exc))
        d['_meta']=meta
        results.append(d)
    return results


def business_checks(results,rules):
    checks=Counter();failures=[]
    for d in results:
        q=d.get('AS');step=d.get('BD');rn=d['_meta']['export_row']
        if not number(q):checks['数量未定行数']+=1;continue
        checks['数值结果行数']+=1
        if q<0:failures.append([rn,'负补货量',q])
        if number(step) and (abs(q/step-round(q/step))>1e-8):failures.append([rn,'未满足出库倍数',q,step])
        if number(d.get('BQ')) and q<d['BQ']:failures.append([rn,'最低恢复量未兑现',q,d['BQ']])
        if (all(number(d.get(c)) for c in ['H','I','J','K','L']) and number(step)
            and d['K']<=0 and max(d['L'],0)<min(rules.minimum_reference_units,step)
            and (max(d['H'],d['I'],d['J'])>0 or d['J']==0 and number(d['N']) and d['N']>0)):
            checks['应最低恢复行数']+=1
            if q<step:failures.append([rn,'有销或缺货零库存漏补',q,step])
        if number(d.get('BQ')) and number(step) and max(d['L'],0)>=min(rules.minimum_reference_units,step) and d['BQ']>0:
            failures.append([rn,'在途足够仍重复保底',d['BQ']])
        if d.get('AW')=='极端缺货自动拉回' and number(step) and q not in (0,step):failures.append([rn,'极端恢复非一个单位',q])
    return {'counts':dict(checks),'failures':failures}


def safe_cell(ws,value):
    c=WriteOnlyCell(ws,value=value)
    if isinstance(value,str):c.data_type='s'  # product strings beginning '=' are data
    elif number(value):c.number_format='0.##########'
    return c


def append_visible(ws,row,widths=None):
    """Plain, uncoloured output; long text wraps and stays expanded."""
    widths=widths or {}
    cells=[];lines=1
    for i,value in enumerate(row,1):
        cell=safe_cell(ws,value)
        if isinstance(value,str):
            cell.alignment=Alignment(wrap_text=True,vertical='top')
            visual=sum(2 if unicodedata.east_asian_width(ch) in 'WF' else 1 for ch in value)
            lines=max(lines,math_ceil(visual/max(5,widths.get(i,18)-1)))
        cells.append(cell)
    # XLSX row height has a 409-point limit. Text remains available in full.
    nextrow=getattr(ws,'_output_row',1)+1
    ws.row_dimensions[nextrow].height=min(409,max(22,lines*15+6))
    ws._output_row=nextrow
    ws.append(cells)


def math_ceil(x):
    return int(x) if x==int(x) else int(x)+1


def sheet_setup(ws,headers,widths=None):
    ws.freeze_panes='H2'
    ws.row_dimensions[1].height=42
    for i,h in enumerate(headers,1):
        width=(widths or {}).get(i,18)
        ws.column_dimensions[get_column_letter(i)].width=width
    cells=[]
    for h in headers:
        c=safe_cell(ws,h);c.alignment=Alignment(wrap_text=True,vertical='center');cells.append(c)
    ws.append(cells)


def write_workbook(path,results,report,config,all_headers):
    if Path(path).exists():raise FileExistsError(f'为避免覆盖，输出文件已存在：{path}')
    book=openpyxl.Workbook(write_only=True)
    ws=book.create_sheet('计算结果')
    columns=list(all_headers)
    extra=['CV（阶段差异）','销售状态诊断','三段日均极差率','导出源行','商品资料源行','资料匹配','资料分类','接口说明','导出补货状态','导出更新时间']
    headers=[all_headers[c] for c in columns]+extra
    main_widths={1:34,2:20,4:65,5:25,6:34,65:100,75:80,77:26,78:42,79:26,82:24,83:30,84:100,85:22,86:25}
    sheet_setup(ws,headers,main_widths)
    for d in results:
        m=d['_meta']
        row=[d.get(c) for c in columns]+[d.get('CV'),d.get('诊断销售状态'),d.get('三段日均极差率'),
              m['export_row'],m['product_row'],m['match'],m['category'],'；'.join(m['blocking']+m['notes']),m['export_status'],m['export_updated']]
        append_visible(ws,row,main_widths)
    ws.auto_filter.ref=f'A1:{get_column_letter(len(headers))}{len(results)+1}'
    summary=book.create_sheet('参数及运行说明')
    sheet_setup(summary,['项目','值','说明'],{1:38,2:70,3:110})
    summary.freeze_panes='A2'
    entries=[['模型版本',VERSION,'计算结果 AS 为最终数量；Python直接计算，工作簿内为本次结果快照'],
             ['输入销售库存文件',report['sales_path'],report['hashes_before']['sales']],
             ['输入商品资料文件',report['products_path'],report['hashes_before']['products']],
             ['运行时间',report['run_time'],'使用所提供文件快照，不表示已刷新到实时库存或实时价格'],
             ['价格口径',config['product_fields']['price'],'对应原模型V；不使用商品资料采购价或导出门店采购价替代'],
             ['空白出库箱规',config['interface']['blank_outbound_unit'],'one=按1件；pending=待核；标准装箱数量不参与出库倍数'],
             ['销量口径','日销=14天日均；周销=近7天累计；月销=近30天累计','沿用原模型：14天累计=ROUND(日销×14,0)，三个窗口同截止时点'],
             ['资料未匹配','明确待补资料，不视为非中心仓或无箱规','需要补齐商品资料，或另外提供明确的非中心仓资料记录'],
             ['统计期限制','连续缺货天数只描述最后一次缺货','不等于完整上架历史；未引入新品上架日期或虚构试销期']]
    for k,v in report['summary'].items():entries.append([k,v,''])
    labels=config['parameter_labels']
    for k,v in config['parameters'].items():entries.append([labels.get(k,k),v,'规则字段：'+k])
    for k,profile in config['weight_profiles'].items():
        for key,value in profile.items():entries.append([k+' '+key,value,'固定常数；不随销量自动调参'])
        entries.append([k+' w3',1-profile['w1']-profile['w2'],'自动计算 1-w1-w2；保留10位小数'])
    for k,v in config['category_profiles'].items():entries.append(['分类映射：'+k,v,'从商品资料分类映射；更细的品类权重尚未引入'])
    for section in ['sales_fields','product_fields']:
        for k,v in config[section].items():entries.append([section+'.'+k,v,'字段映射接口'])
    entries.extend([['原资格门槛','P≤12且P≠-1，或月销≤5且V<6','最低有货恢复仍可豁免原资格门槛'],
                    ['周转P','ROUND(MAX(K+0.7L,0)/(I/7或J/30),0)','周销优先，周销为0时用月销；沿用原P，0.7仅用于资格参考，AQ在途仍只扣一次'],
                    ['最低有货','缺货且无足够在途时，优先恢复一个实际出库单位','月销0且0<N<30、极端缺货也保留；原首铺数量不作为永久库存下限'],
                    ['供货/停采状态','保留说明，不新增自动停补门槛','数量是模型建议；核验供货可用性和导出停采状态后再执行'],
                    ['资金金额','仅汇总BV为数值的结果','未知价格与未定数量不当作0计入完整资金；另有计数提示'],
                    ['使用方式','每次更换输入后重新运行Python','在结果工作簿中手动改输入不会触发Python重算'],
                    ['检核失败数',len(report['checks']['failures']),'应为0；不代表供货、上架和真实销售表现已验证']])
    for row in entries:append_visible(summary,row,{1:38,2:70,3:110})
    audit=book.create_sheet('重点核验')
    ah=['导出源行','门店编码','商品编码','商品名称','最终AS','采购价V','出库单位BD','月销J','缺货天数N','校正状态AW','计算类型BB','业务状态BL','模型重点关注BW','接口与资料说明','商品资料源行']
    sheet_setup(audit,ah,{2:22,3:26,4:65,13:90,14:110})
    count=0
    for d in results:
        m=d['_meta'];notes='；'.join(m['blocking']+m['notes'])
        # Ordinary blank-to-one note stays in detail, not an error by itself.
        attention=bool(d.get('BW') or m['blocking'] or any('停采' in n for n in m['notes']))
        if not attention:continue
        row=[m['export_row'],d['B'],d['E'],d['D'],d.get('AS'),d.get('V'),d.get('BD'),d.get('J'),d.get('N'),d.get('AW'),d.get('BB'),d.get('BL'),d.get('BW'),notes,m['product_row']]
        append_visible(audit,row,{2:22,3:26,4:65,13:90,14:110});count+=1
    audit.auto_filter.ref=f'A1:O{count+1}'
    Path(path).parent.mkdir(parents=True,exist_ok=True)
    book.save(path)


def run(sales,products,output,config_path):
    config,rules=load_config(config_path)
    sales=Path(sales).resolve();products=Path(products).resolve();output=Path(output).resolve()
    if output in (sales,products) or output.exists():raise ValueError('输出必须是未存在的新文件，不能覆盖源文件')
    before={'sales':digest(sales),'products':digest(products)}
    data,ingestion=prepare(sales,products,config)
    result=execute(data,rules)
    checks=business_checks(result,rules)
    summary={'输入商品行数':len(result),'资料精确匹配行数':ingestion['matching'].get('商品编码精确匹配',0),
             '资料未匹配行数':ingestion['matching'].get('资料未匹配',0),
             '建议补货SKU数':sum(number(d.get('AS')) and d['AS']>0 for d in result),
             '建议0补货SKU数':sum(number(d.get('AS')) and d['AS']==0 for d in result),
             '最终数量未定SKU数':sum(not number(d.get('AS')) for d in result),
             '已知价格补货金额':round(sum(d['BV'] for d in result if number(d.get('BV'))),2),
             '有补货但金额未定SKU数':sum(number(d.get('AS')) and d['AS']>0 and not number(d.get('BV')) for d in result),
             '额外量待定SKU数':sum(number(d.get('BX')) and d['BX']>0 for d in result),
             '销量窗口异常SKU数':sum(d.get('BB')=='销量窗口异常' for d in result),
             '缺货销量冲突SKU数':sum(d.get('BB')=='缺货销量冲突' for d in result),
             '计算异常SKU数':sum(d.get('AW')=='计算异常' for d in result)}
    report={'model_version':VERSION,'run_time':datetime.now().astimezone().isoformat(),
            'sales_path':str(sales),'products_path':str(products),'output_path':str(output),
            'hashes_before':before,'config':config,'summary':summary,'ingestion':ingestion,'checks':checks}
    if checks['failures'] or summary['计算异常SKU数']:
        report['status']='验证失败，未生成可用结果'
        output.with_suffix('.failure.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
        raise ValueError('计算或业务检核失败，已生成 failure.json，未把异常当成0')
    headers=json.loads((HERE/'schema.json').read_text(encoding='utf8'))['headers']
    report['status']='部分SKU待补数据' if summary['最终数量未定SKU数'] else '全部数量已计算'
    write_workbook(output,result,report,config,headers)
    report['hashes_after']={'sales':digest(sales),'products':digest(products)}
    report['sources_unchanged']=before==report['hashes_after']
    report['output_sha256']=digest(output)
    output.with_suffix('.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
    if not report['sources_unchanged']:raise ValueError('运行期间输入文件发生变化，请使用固定快照重跑')
    print(json.dumps({'output':str(output),'status':report['status'],'summary':summary,'check_failures':len(checks['failures'])},ensure_ascii=False,indent=2))
    return report


def main():
    parser=argparse.ArgumentParser(description='最新补货模型；销售库存导出+商品资料；AS为最终量。')
    parser.add_argument('--sales',required=True,help='销售库存导出xlsx')
    parser.add_argument('--products',required=True,help='商品资料xlsx')
    parser.add_argument('--output',help='新结果xlsx；不指定则写到results目录')
    parser.add_argument('--config',default=str(HERE/'config.json'),help='字段映射、固定权重和规则参数')
    args=parser.parse_args()
    output=args.output or str(HERE/'results'/('补货结果_'+datetime.now().strftime('%Y%m%d_%H%M%S_%f')+'.xlsx'))
    try:run(args.sales,args.products,output,args.config)
    except (ValueError,FileNotFoundError,PermissionError,FileExistsError) as exc:
        print('未完成：'+str(exc),file=sys.stderr);return 2
    return 0


if __name__=='__main__':raise SystemExit(main())
