"""Opt-in local release self-test. Receives all fixture paths on command line."""
import json,sys,time,traceback
from pathlib import Path
from PySide6.QtCore import QTimer,Qt
from PySide6.QtTest import QTest
import engine

def start_smoke(app,win,args):
    sales,products,output=args[:3]
    out=Path(output);out.mkdir(parents=True,exist_ok=True)
    state={'stage':0,'started':time.monotonic(),'checks':[]}
    def finish(error=None):
        timer.stop()
        report={'passed':error is None,'error':error,'checks':state['checks'],'summary':win.result['summary'] if win.result else None}
        (out/'ui_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
        app.exit(1 if error else 0)
    def fail(message):finish(message)
    win.failed=fail
    sys.excepthook=lambda t,v,tb:finish(''.join(traceback.format_exception(t,v,tb)))
    def snap(name):
        win.repaint();assert win.grab().save(str(out/(name+'.png')))
    def tick():
        try:
            if time.monotonic()-state['started']>360:raise TimeoutError(str(state))
            stage=state['stage']
            if stage==0:
                snap('01_初始界面');win.sales.set_path(sales);win.products.set_path(products);state['stage']=1
            elif stage==1 and win.loaded and win.worker is None:
                assert len(win.model.rows)==14244
                assert engine.summarize(win.model.rows)['local']==5161
                assert win.calculate_btn.isEnabled() and not win.export_all.isEnabled()
                snap('02_导入检查');state['checks'].append('import preview and local classification')
                QTest.mouseClick(win.calculate_btn,Qt.MouseButton.LeftButton);state['stage']=2
            elif stage==2 and win.result and win.worker is None:
                assert win.result['summary']['amount']==12268.87
                assert win.result['summary']['exception']==8
                assert win.proxy.rowCount()==430 and win.export_all.isEnabled()
                assert all(not engine.exclusion_reason(d) for d in win.visible_rows())
                snap('03_计算结果');win.mode.setCurrentText('地采商品');state['stage']=3
            elif stage==3:
                assert win.proxy.rowCount()==5161
                win.show_detail(win.proxy.index(0,0));assert '地采商品' in win.details.toPlainText()
                snap('04_地采明细');win.detail_panel.hide();win.mode.setCurrentText('一箱例外');state['stage']=4
            elif stage==4:
                assert win.proxy.rowCount()==8
                win.show_detail(win.proxy.index(0,0));assert '放行依据' in win.details.toPlainText()
                snap('05_一箱例外');win.resize(1080,720);state['stage']=5
            elif stage==5:
                snap('06_小窗口');win.resize(1380,900)
                first=win.visible_rows()[0]['E'];win.search.setText(first);assert win.proxy.rowCount()==1
                win.search.setText('no_such_product_0');assert win.proxy.rowCount()==0
                win.search.clear();win.detail_panel.hide();win.mode.setCurrentText('全部商品');assert win.proxy.rowCount()==14244
                win.mode.setCurrentText('商品资料禁用');assert win.proxy.rowCount()==1973
                win.mode.setCurrentText('补货建议停采');assert win.proxy.rowCount()==1564
                win.mode.setCurrentText('已排除');assert win.proxy.rowCount()==3269
                assert all(engine.order_quantity(d)==0 and engine.order_amount(d)==0 for d in win.visible_rows())
                win.show_detail(win.proxy.index(0,0));assert '已排除' in win.details.toPlainText()
                state['checks'].append('disabled / stopped / union filters; all excluded execution values zero')
                state['stage']=51
            elif stage==51:
                snap('07_已排除商品');win.detail_panel.hide()
                state['checks'].append('local / exception / search filters and explanation drawer')
                win.exception.setChecked(False);assert win.dirty and not win.export_all.isEnabled()
                win.start_calculation();state['stage']=6
            elif stage==6 and win.worker is None:
                assert win.result['summary']['amount']==11890.07 and win.result['summary']['exception']==0
                state['checks'].append('toggle off restores exact baseline; stale export disabled')
                win.exception.setChecked(True);win.start_calculation();state['stage']=7
            elif stage==7 and win.worker is None:
                assert win.result['summary']['amount']==12268.87
                win.mode.setCurrentText('一箱例外')
                def exported(path):state['export']=path
                win.job(engine.export_xlsx,(win.result,str(out/'界面导出_一箱例外.xlsx'),win.visible_rows(),'当前筛选：一箱例外'),exported)
                state['stage']=8
            elif stage==8 and win.worker is None:
                assert Path(state['export']).is_file()
                import openpyxl
                book=openpyxl.load_workbook(state['export'],read_only=True)
                assert sum(1 for _ in book['筛选明细'].iter_rows())==9;book.close()
                state['checks'].append('background filtered export reopened with eight rows')
                spin=win.warehouses.cellWidget(0,1);spin.setValue(spin.value()+.5)
                assert win.dirty and not win.export_all.isEnabled()
                state['checks'].append('warehouse parameter change invalidates prior result')
                spin.setValue(spin.value()-.5)
                win.budget_enabled.setChecked(True);win.budget_amount.setValue(3000)
                assert win.dirty and not win.export_all.isEnabled()
                win.start_calculation();state['stage']=9
            elif stage==9 and win.worker is None:
                b=win.result['budget'];assert b['enabled'] and b['limit']==3000 and b['spent']<=3000
                assert len(win.profile['items'])==2142
                win.mode.setCurrentText('预算未满足');assert win.proxy.rowCount()==b['deferred']>0
                win.show_detail(win.proxy.index(0,0));assert '资金分配' in win.details.toPlainText()
                state['stage']=91
            elif stage==91:
                snap('08_预算缺口');win.detail_panel.hide();win.mode.setCurrentText('最低持有缺口')
                assert win.proxy.rowCount()==win.result['budget']['holding_gaps']
                win.budget_amount.setValue(0);assert win.dirty and not win.export_all.isEnabled()
                win.start_calculation();state['stage']=10
            elif stage==10 and win.worker is None:
                assert win.result['summary']['amount']==0 and win.proxy.rowCount()==0
                twin=win.__class__();twin.loaded_ok(win.loaded)
                assert twin.budget_enabled.isChecked() and twin.budget_amount.value()==0
                assert twin.profile['revision']==win.profile['revision'];twin.close()
                state['checks'].append('zero budget means zero orders; fresh window restores list and store budget')
                win.budget_amount.setValue(3000);win.start_calculation();state['stage']=11
            elif stage==11 and win.worker is None:
                snap('09_预算分配结果')
                def exported_budget(path):state['budget_export']=path
                win.job(engine.export_xlsx,(win.result,str(out/'界面导出_预算3000.xlsx')),exported_budget);state['stage']=12
            elif stage==12 and win.worker is None:
                import openpyxl
                book=openpyxl.load_workbook(state['budget_export'],read_only=True)
                assert sum(1 for _ in book['补货清单'].iter_rows())-1==win.result['summary']['positive']
                assert sum(1 for _ in book['全店必卖匹配'].iter_rows())-1==2142;book.close()
                state['checks'].append('budget output sheets and must-list audit reopened')
                import priority_list
                saved=out/'清单备份.json';priority_list.export_profile(win.profile,saved)
                win.job(priority_list.import_profile,(str(saved),),win.priority_imported);state['stage']=13
            elif stage==13 and win.worker is None:
                assert win.dirty and not win.export_all.isEnabled()
                state['checks'].append('imported list persists and invalidates previous result')
                finish()
        except Exception:finish(traceback.format_exc())
    timer=QTimer(win);timer.setInterval(300);timer.timeout.connect(tick);timer.start()
