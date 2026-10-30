"""Opt-in GUI acceptance extension, called only by warehouse-delivery-smoke."""
from pathlib import Path
from PySide6.QtWidgets import QFileDialog,QMessageBox
import openpyxl
import work_repository as wr


def exercise(win,out,checks,shot,wait,phase):
    from history_ui import ArchiveDialog,AuditDialog
    from contextlib import closing
    import archive_orders as ao
    win.show_section(2);wait(lambda:win.history_reader.active is None and bool(win.archive_all))
    if phase=='reopen':
        win.archive_trash.setCurrentIndex(1);win.refresh_archive();wait(lambda:win.history_reader.active is None)
        if len(win.archive_orders)!=1:raise AssertionError('trash not persisted across restart')
        win.archive_table.selectRow(0);old=QMessageBox.question;QMessageBox.question=lambda *a,**k:QMessageBox.StandardButton.Yes
        try:win.change_archive_visibility();wait(lambda:win.history_reader.active is None)
        finally:QMessageBox.question=old
        if win.archive_orders:raise AssertionError('restore failed')
        win.archive_trash.setCurrentIndex(0);win.refresh_archive();wait(lambda:win.history_reader.active is None)
        if len(win.archive_orders)!=4:raise AssertionError('other warehouse affected')
        checks.append('EXE restart preserves recycle bin; restore returns original four warehouse orders');shot(win,'archive-restored.png');return
    if len(win.archive_orders)!=4:raise AssertionError('expected two stores x two warehouses')
    h=next(h for h in win.archive_orders if h['warehouse']=='河北仓');index=win.archive_orders.index(h);win.archive_table.selectRow(index);shot(win,'archive-warehouse-list.png')
    order=win.orders.get(h['day'],h['store']);before=wr.dumps(order['payload'])
    dialog=ArchiveDialog(win,h);dialog.show();wait(lambda:dialog.selected is not None)
    expected=[d['E'] for d in order['payload']['result']['rows'] if ao.warehouse(d)=='河北仓']
    if [r[0] for r in dialog.model.rows]!=expected:raise AssertionError('detail scope leaked')
    audit=AuditDialog(dialog,win,dialog.selected,products=dialog.model.rows,sku=expected[0]);audit.show();wait(lambda:audit.loaded_key is not None)
    for scope in ('当时参数','商品建议与销量库存','当时商品资料快照','商品完整计算依据'):
        audit.scope.setCurrentText(scope);wait(lambda:audit.loaded_key is not None)
        if not audit.model.rows:raise AssertionError('empty business form '+scope)
    shot(audit,'archive-calculation-form.png');audit.reject();audit.deleteLater()
    target=out/'archive-three.xlsx';old=QFileDialog.getSaveFileName;QFileDialog.getSaveFileName=lambda *a,**k:(str(target),'')
    try:dialog.export_archive()
    finally:QFileDialog.getSaveFileName=old
    b=openpyxl.load_workbook(target)
    if b.sheetnames!=['完整计算建议表','制单时商品资料原表','模型建议签发表']:raise AssertionError(b.sheetnames)
    if b.worksheets[1].max_row<2:raise AssertionError('raw source absent')
    b.close()
    merchant=out/'archive-history-hb.xlsx';QFileDialog.getSaveFileName=lambda *a,**k:(str(merchant),'')
    try:dialog.export();wait(lambda:merchant.exists() and win.worker is None)
    finally:QFileDialog.getSaveFileName=old
    book=openpyxl.load_workbook(merchant)
    if {p.name:p.value for p in book.custom_doc_props}['供货仓']!='河北仓':raise AssertionError('history selected wrong warehouse')
    book.close();dialog.reject();dialog.deleteLater()
    from history_ui import EventsDialog
    from history_forms import events
    data=events(win.work_repository.path,h['oid'],'河北仓');events_dialog=EventsDialog(win,data['records'],data['rows']);events_dialog.show();events_dialog.table.selectRow(0)
    event_export=out/'event-three.xlsx';QFileDialog.getSaveFileName=lambda *a,**k:(str(event_export),'')
    try:events_dialog.export()
    finally:QFileDialog.getSaveFileName=old
    book=openpyxl.load_workbook(event_export)
    if book.sheetnames!=['完整计算建议表','制单时商品资料原表','模型建议签发表']:raise AssertionError('event export mismatch')
    book.close();events_dialog.reject();events_dialog.deleteLater()
    checks.append('actual per-warehouse detail/four forms/merchant scope; archive and operation record both export the exact three sheets')
    old=QMessageBox.question;QMessageBox.question=lambda *a,**k:QMessageBox.StandardButton.Cancel
    try:win.change_archive_visibility()
    finally:QMessageBox.question=old
    with closing(win.work_repository.connect()) as db:
        if ao.states(db):raise AssertionError('cancel changed trash')
    QMessageBox.question=lambda *a,**k:QMessageBox.StandardButton.Yes
    try:win.change_archive_visibility();wait(lambda:win.history_reader.active is None)
    finally:QMessageBox.question=old
    if len(win.archive_orders)!=3:raise AssertionError('wrong child hidden')
    win.archive_trash.setCurrentIndex(1);win.refresh_archive();wait(lambda:win.history_reader.active is None)
    if len(win.archive_orders)!=1:raise AssertionError('trash view')
    trash=ArchiveDialog(win,win.archive_orders[0]);trash.show();wait(lambda:trash.selected is not None)
    if trash.archive_export_btn.isEnabled() or trash.export_btn.isEnabled():raise AssertionError('trash export enabled')
    shot(trash,'archive-recycle-readonly.png');trash.reject();trash.deleteLater()
    if wr.dumps(win.orders.get(h['day'],h['store'])['payload'])!=before:raise AssertionError('signed payload changed')
    checks.append('delete defaults to cancel; confirmed one-warehouse recycle, immutable signed data, export disabled until restored')
    win.show_section(0);wait(lambda:win.history_reader.active is None)
    if len(win.archive_all)!=3 or any(r.get('deleted') for r in win.archive_all):raise AssertionError('home inherited recycle-bin filter')
    checks.append('home always counts active warehouse orders even after recycle-bin browsing')
