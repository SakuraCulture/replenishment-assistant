"""Opt-in actual Qt/EXE acceptance; private fixtures/config supplied by launcher."""
import copy,json,sys,time,traceback
from pathlib import Path
from PySide6.QtCore import Qt,QTimer
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication,QMessageBox,QFileDialog,QLineEdit
import warehouse_catalog as wc,work_repository as wr,normal_orders
from new_product_ui import NewProductDialog
from new_product_focus_ui import FocusMapping
import new_product_focus as focus
from parameter_ui import NumberInput
from version_info import diagnostic
import new_products,priority_list


def start_smoke(app,win,args):
    fixture,out=[Path(p).resolve() for p in args[:2]];out.mkdir(parents=True,exist_ok=True);state={'step':-3,'start':time.monotonic(),'checks':[]};timer=QTimer(win)
    app.setAttribute(Qt.ApplicationAttribute.AA_DontUseNativeDialogs,True)
    def snap(name):win.repaint();assert win.grab().save(str(out/(name+'.png')))
    def finish(error=None):
        timer.stop();meta=diagnostic();r={'passed':error is None,'error':error,'version':meta['version'],'source_commit':meta['source_commit'],'checks':state['checks']};(out/'normal_ui_report.json').write_text(json.dumps(r,ensure_ascii=False,indent=2),encoding='utf-8')
        if error:app.exit(1);return
        # Exercise the same accepted-close path as an operator, not a forced
        # event-loop exit while modal/worker callbacks are still unwinding.
        def close():
            def discard():
                box=QApplication.activeModalWidget()
                if isinstance(box,QMessageBox):QTest.mouseClick(next(b for b in box.buttons() if b.text()=='放弃本次变更'),Qt.MouseButton.LeftButton)
            QTimer.singleShot(150,discard);win.close()
        QTimer.singleShot(0,close)
    win.failed=lambda message:finish(message);sys.excepthook=lambda t,v,tb:finish(''.join(traceback.format_exception(t,v,tb)))
    def later(fn):
        def guarded():
            try:fn()
            except Exception:finish(traceback.format_exc())
        QTimer.singleShot(150,guarded)
    def modal_click(name,text,attempt=0):
        box=QApplication.activeModalWidget()
        if not (isinstance(box,QMessageBox) and box.objectName()==name) and attempt<300:
            later(lambda:modal_click(name,text,attempt+1));return
        assert isinstance(box,QMessageBox) and box.objectName()==name;assert box.defaultButton().text()=='取消';assert box.grab().save(str(out/(name+'_'+text+'.png')));QTest.mouseClick(next(b for b in box.buttons() if b.text()==text),Qt.MouseButton.LeftButton)
    def select_new(amount=6,accept=True):
        dialog=QApplication.activeModalWidget();assert isinstance(dialog,NewProductDialog)
        i=next(i for i,d in enumerate(dialog.rows) if d['E']=='new_must');check,qty=dialog.widgets[i]
        dialog.table.edit(dialog.model.index(dialog.model.positions[i],8));editor=dialog.table.findChild(NumberInput);assert editor is not None;editor.setValue(amount);QTest.keyClick(editor,Qt.Key.Key_Return);QApplication.processEvents();assert qty.value()==amount;check.setChecked(True)
        dialog.scope.setCurrentText('本月新品');assert i in dialog.model.positions;dialog.scope.setCurrentText('往期新品');assert i not in dialog.model.positions;dialog.scope.setCurrentText('可选');dialog.search.setText('new_plain');dialog.refresh();assert check.isChecked() and i not in dialog.model.positions;assert '筛选外已选 1' in dialog.note.text();dialog.search.clear();dialog.refresh();assert dialog.grab().save(str(out/'新品选择.png'))
        later(lambda:modal_click('newBudgetPreview','应用到本店草稿' if accept else '取消'))
        QTest.mouseClick(dialog.preview_btn,Qt.MouseButton.LeftButton)
    def focus_file(accept):
        from import_ui import ImportDialog
        dialog=QApplication.activeModalWidget();assert isinstance(dialog,ImportDialog);dialog.set_paths([str(fixture/'新品明细.xlsx')]);dialog.accept()
        def mapping(attempt=0):
            dialog=QApplication.activeModalWidget()
            if not isinstance(dialog,FocusMapping) and attempt<200:later(lambda:mapping(attempt+1));return
            assert isinstance(dialog,FocusMapping);assert dialog.sku.currentText()=='商品编码' and dialog.focus.currentText()=='关注';assert dialog.grab().save(str(out/'00_新品明细映射.png'))
            later(lambda:modal_click('focusConfirm','使用并保存清单' if accept else '取消'));dialog.accept()
        later(mapping)
    def tick():
        try:
            if time.monotonic()-state['start']>240:raise TimeoutError(str(state))
            if QApplication.activeModalWidget() is not None:return
            step=state['step']
            if step==-3:
                win.show_section(1);state['step']=-2;later(lambda:focus_file(False));win.import_focus()
            elif step==-2 and win.worker is None:
                assert win.new_focus is None;state['step']=-1;later(lambda:focus_file(True));win.import_focus()
            elif step==-1 and win.worker is None:
                assert 'new_must' in focus.members(win.new_focus);assert win.focus_repository.load()==win.new_focus;state['checks'].append('actual independent focus file dialog, sheet/column mapping, preview cancel/confirm, reusable JSON persistence');win.show_section(0);state['step']=0
            elif step==0:
                assert win.sections.currentIndex()==0;snap('01_首页');QTest.mouseClick(win.top_nav[1],Qt.MouseButton.LeftButton);assert win.sections.currentIndex()==1
                for w in wc.WAREHOUSES:win.catalog_imported(wc.import_catalog(fixture/(w+'.xlsx'),w))
                win.profile=json.loads((fixture/'must.json').read_text(encoding='utf-8'));win.repository.save(win.profile);win.refresh_profile_label();win.sales.add_paths([str(fixture/'sales.xlsx')]);win.start_load();state['step']=1
            elif step==1 and win.worker is None:
                assert len(win.loaded['normalized'])==22
                for s,w in win.store_widgets.items():
                    w['mode'].setCurrentIndex(2);w['amount'].setText('12');w['primary'].setCurrentIndex(w['primary'].findData('河北仓'));w['secondary'].setCurrentIndex(w['secondary'].findData('山东仓'))
                win.start_calculation();state['step']=2
            elif step==2 and win.worker is None:
                assert all(x['state']=='已完成' for x in win.states.values()),win.states
                state['original']=copy.deepcopy(win.results);win.current_store='QA_A';win.navigate(3);state['step']=20
                later(lambda:select_new(6,False));win.choose_new_products();state['step']=3
            elif step==3 and win.worker is None:
                assert win.results==state['original'];state['checks'].append('real file check, two-store spawn calculation; filtered selection retained; budget preview cancel preserves both stores')
                state['step']=30;later(lambda:select_new(6,True));win.choose_new_products();state['step']=4
            elif step==4 and win.worker is None:
                assert win.results['QA_A']['new_products']['selection']=={'new_must':6.0};assert win.results['QA_B']==state['original']['QA_B'];assert win.results['QA_A']['summary']['must_amount']==12
                win.mode.setCurrentText('新品推荐');assert win.proxy.rowCount()==1;snap('02_新品预算执行');state['step']=40
                later(lambda:modal_click('normalConfirm','取消'));win.confirm_normal();assert not win.orders.list()
                later(lambda:modal_click('normalConfirm','确认并归档'));win.confirm_normal();assert len(win.orders.list())==2
                state['order']=win.orders.get(wr.today(),'QA_A');assert state['order']['revision']==1 and not win.job_dirty
                later(lambda:modal_click('normalConfirm','确认并归档'));win.confirm_normal();assert win.orders.get(wr.today(),'QA_A')['revision']==1
                state['checks'].append('mandatory new item takes shared budget ahead of ordinary sales; other store unchanged; new filter; first confirmation cancel/accept; auto daily save; duplicate confirm idempotent')
                # Prepare a changed first-fill candidate; execution may stay capped by budget.
                later(lambda:select_new(12,True));win.choose_new_products();state['step']=5
            elif step==5 and win.worker is None:
                assert win.orders.get(wr.today(),'QA_A')['revision']==1;state['step']=50
                later(lambda:modal_click('normalConfirm','取消'));win.confirm_normal();assert win.orders.get(wr.today(),'QA_A')['revision']==1
                later(lambda:modal_click('normalConfirm','确认并归档'));win.confirm_normal();order=win.orders.get(wr.today(),'QA_A');assert order['revision']==2 and order['number']==state['order']['number'];assert win.orders.get(wr.today(),'QA_B')['revision']==1
                assert win.orders.get(wr.today(),'QA_A',1)['payload']['result']['new_products']['selection']['new_must']==6
                state['checks'].append('changed candidate leaves archived V001 until explicit overwrite; default-cancel; same number V002; unchanged other store remains V001; old revision readable')
                def choose_file():
                    dialog=QApplication.activeModalWidget()
                    if not isinstance(dialog,QFileDialog):later(choose_file);return
                    dialog.setDirectory(str(out));dialog.findChild(QLineEdit,'fileNameEdit').setText('EXE确认单.xlsx');assert Path(dialog.selectedFiles()[0])==out/'EXE确认单.xlsx';dialog.accept()
                later(choose_file);win.export_current([win.current_store]);state['step']=6
            elif step==6 and win.worker is None:
                assert (out/'EXE确认单.xlsx').is_file();
                import openpyxl
                with_book=openpyxl.load_workbook(out/'EXE确认单.xlsx');assert with_book.sheetnames==['往期新品','本月新品','中心仓动销补货清单'];assert with_book['本月新品']['E2'].value=='new_must' and with_book['本月新品']['Q2'].value==6;assert with_book['往期新品'].max_row==6;assert not {'new_disabled','new_stopped'} & {r[4] for r in with_book['往期新品'].iter_rows(min_row=2,values_only=True)};with_book.close();assert any(e['action']=='带号导出' for e in win.orders.events())
                record=win.work_repository.get_job(wr.today());state['record']=record;win.reset_job();assert not win.new_bases;win.restore_job(record);assert 'QA_A' in win.new_bases and win.results['QA_A']['new_products']['selection']['new_must']==12
                state['checks'].append('actual file-save dialog and frozen numbered Excel export; queryable event; saved new choices and original base restored after resetting work area')
                win.show_section(2);win.archive_search.setText('new_must');state['step']=61
            elif step==61 and not win.archive_timer.isActive() and win.history_reader.active is None:
                assert win.archive_orders and all(r['store'] in ('QA_A','QA_B') for r in win.archive_orders);snap('03_归档查询');state['step']=60
                def inspect_archive():
                    dialog=QApplication.activeModalWidget();assert dialog is not None
                    if not getattr(dialog,'selected',None):later(inspect_archive);return
                    assert dialog.grab().save(str(out/'04_修订明细.png'));dialog.accept()
                later(inspect_archive);win.open_archive(win.archive_orders[0])
                win.show_section(1);win.store_widgets['QA_A']['amount'].setText('0');assert win.states['QA_A']['state']=='待重算';assert win.orders.get(wr.today(),'QA_A')['revision']==2
                win.current_store='QA_A';assert not win.orders.get(wr.today(),'QA_A')['payload']['result']['budget']['limit']==0
                # Recalculation from the saved immutable input snapshot retains standard behavior.
                later(lambda:modal_click('newSelectionReset','重新计算并重选新品'));win.start_calculation(['QA_A']);state['step']=7
            elif step==7 and win.worker is None:
                assert win.results['QA_A']['summary']['amount']==0 and 'QA_A' not in win.new_bases and not win.results['QA_A'].get('new_products')
                assert win.orders.get(wr.today(),'QA_A')['revision']==2
                state['checks'].append('history query/detail independent of work area; parameter changes invalidate only selected store; rerun resets selections and leaves issued order untouched')
                win.show_section(0);snap('05_首页归档状态');state['step']=70
                def cancel_close():
                    box=QApplication.activeModalWidget();assert isinstance(box,QMessageBox);assert box.defaultButton().text()=='取消';QTest.mouseClick(next(b for b in box.buttons() if b.text()=='取消'),Qt.MouseButton.LeftButton)
                later(cancel_close);win.close();assert win.isVisible();state['checks'].append('home navigation retains unsaved work; close cancel keeps application alive')
                if len(args)<3:finish();return
                real=Path(args[2]);win.reset_job();win.show_section(1);win.master_legacy.setChecked(True);win.products.path=str(real/'混合仓商品资料.xlsx');win.products.filename.setText(win.products.path);win.sales.add_paths([str(real/'补货建议导出服务.xlsx')]);win.profile=priority_list.validate_profile(json.loads((Path(__file__).parent/'default_must_sell.json').read_text(encoding='utf-8')));win.refresh_profile_label();win.start_load();state['step']=8
            elif step==8 and win.worker is None:
                assert len(win.loaded['normalized'])==14244
                for s,w in win.store_widgets.items():w['mode'].setCurrentIndex(1)
                win.start_calculation();state['step']=9
            elif step==9 and win.worker is None:
                store=next(iter(win.results));win.current_store=store;result=win.results[store];assert result['summary']['positive']==430 and result['summary']['amount']==12268.87
                state['real_before']=copy.deepcopy(result);state['step']=90
                def select_real():
                    dialog=QApplication.activeModalWidget();assert isinstance(dialog,NewProductDialog);assert sum(new_products.is_new(d) for d in state['real_before']['rows'])==9023;assert dialog.rows==[d for d in state['real_before']['rows'] if new_products.is_new(d) and not new_products.blocked_reason(d).startswith(('商品资料禁用','补货建议停采'))]
                    valid=[i for i,d in enumerate(dialog.rows) if not new_products.blocked_reason(d)];assert len(valid)==3455
                    def allow_bulk():
                        box=QApplication.activeModalWidget();assert isinstance(box,QMessageBox);QTest.mouseClick(box.button(QMessageBox.StandardButton.Ok),Qt.MouseButton.LeftButton)
                    later(allow_bulk);dialog.bulk(True);assert len(dialog.selections())==3455
                    later(allow_bulk);dialog.clear_all();assert not dialog.selections()
                    i=valid[0];d=dialog.rows[i];assert d['N'] is None;state['real_sku']=d['E'];state['real_step']=d['BD'];state['real_price']=d['V'];dialog.widgets[i][0].setChecked(True);dialog.search.setText(d['E']);assert dialog.grab().save(str(out/'06_真实新品候选.png'))
                    later(lambda:modal_click('newBudgetPreview','应用到本店草稿'));QTest.mouseClick(dialog.preview_btn,Qt.MouseButton.LeftButton)
                later(select_real);win.choose_new_products();state['step']=10
            elif step==10 and win.worker is None:
                result=win.results[win.current_store];d=next(d for d in result['rows'] if d['E']==state['real_sku']);assert d['AS']==0 and d['N'] is None and d['_execution_qty']==state['real_step'];assert result['summary']['positive']==431;assert result['summary']['amount']==round(12268.87+state['real_step']*state['real_price'],2)
                old={d['E']:d for d in state['real_before']['rows']}
                assert all(d==old[d['E']] for d in result['rows'] if d['E']!=state['real_sku'])
                state['checks'].append('real 14244 rows: 9023 raw blank-stockout new candidates; disabled/stopped removed from selector / 3455 eligible; actual selector searches and applies one; original N/AS retained; other 14243 rows exactly unchanged; unlimited amount reconciles')
                state['step']=100;later(lambda:modal_click('normalConfirm','确认并归档'));win.confirm_normal();assert win.orders.get(wr.today(),win.current_store)
                def real_file():
                    dialog=QApplication.activeModalWidget()
                    if not isinstance(dialog,QFileDialog):later(real_file);return
                    dialog.setDirectory(str(out));dialog.findChild(QLineEdit,'fileNameEdit').setText('EXE真实新品确认单.xlsx');dialog.accept()
                later(real_file);win.export_current([win.current_store]);state['step']=11
            elif step==11 and win.worker is None:
                assert (out/'EXE真实新品确认单.xlsx').is_file();
                import openpyxl
                book=openpyxl.load_workbook(out/'EXE真实新品确认单.xlsx',read_only=True);assert book['本月新品'].max_row==2 and book['本月新品']['E2'].value==state['real_sku'];assert book['往期新品'].max_row==1+sum(new_products.is_new(d) and d['_source']=='资料匹配' and d['E']!=state['real_sku'] for d in win.results[win.current_store]['rows']);assert book['中心仓动销补货清单'].max_row==431;assert sum(row[16]==0 for row in book['往期新品'].iter_rows(min_row=2,values_only=True))==book['往期新品'].max_row-1;book.close();snap('07_真实归档结果');state['checks'].append('real-data normal order signed and numbered merchant workbook exported from frozen revision');finish()
        except Exception:finish(traceback.format_exc())
    timer.timeout.connect(tick);timer.start(250)
