"""Opt-in actual Qt acceptance for one-action delivery and measurable export progress."""
import sys,json,time,traceback,copy
from pathlib import Path
from PySide6.QtCore import QTimer,Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication,QMessageBox,QFileDialog,QLineEdit
import warehouse_catalog as wc,new_product_focus as focus,work_repository as wr,xlsx_exports as x
from version_info import diagnostic


def start_smoke(app,win,args):
    fixture,out=[Path(p).resolve() for p in args[:2]];out.mkdir(parents=True,exist_ok=True);state={'step':0,'start':time.monotonic(),'checks':[]};timer=QTimer(win);beats=[0]
    app.setAttribute(Qt.ApplicationAttribute.AA_DontUseNativeDialogs,True)
    pulse=QTimer(win);pulse.timeout.connect(lambda:beats.__setitem__(0,beats[0]+1));pulse.start(20)
    def finish(error=None):
        timer.stop();pulse.stop();d=diagnostic();r={'passed':error is None,'error':error,'version':d['version'],'source_commit':d['source_commit'],'checks':state['checks'],'event_loop_ticks':beats[0]};(out/'export_ui_report.json').write_text(json.dumps(r,ensure_ascii=False,indent=2),encoding='utf-8');win.job_dirty=False;QTimer.singleShot(0,win.close if error is None else lambda:app.exit(1))
    win.failed=lambda msg:finish(msg);sys.excepthook=lambda t,v,tb:finish(''.join(traceback.format_exception(t,v,tb)))
    def soon(fn):
        def safe():
            try:fn()
            except Exception:finish(traceback.format_exc())
        QTimer.singleShot(60,safe)
    def confirmation(accept,filename,tries=0):
        box=QApplication.activeModalWidget()
        if not isinstance(box,QMessageBox) or box.objectName()!='deliveryConfirm':
            assert tries<500;soon(lambda:confirmation(accept,filename,tries+1));return
        assert box.defaultButton().text()=='取消';assert box.grab().save(str(out/('confirmation-'+str(state['step'])+'.png')))
        QTest.mouseClick(next(b for b in box.buttons() if b.text()==('确认并导出' if accept else '取消')),Qt.MouseButton.LeftButton)
        if accept:soon(lambda:savefile(filename))
    def savefile(name,tries=0):
        dialog=QApplication.activeModalWidget()
        if not isinstance(dialog,QFileDialog):assert tries<500;soon(lambda:savefile(name,tries+1));return
        dialog.setDirectory(str(out));dialog.findChild(QLineEdit,'fileNameEdit').setText(name);dialog.accept()
    def batch_confirm(tries=0):
        box=QApplication.activeModalWidget()
        if not isinstance(box,QMessageBox) or box.objectName()!='deliveryConfirm':assert tries<500;soon(lambda:batch_confirm(tries+1));return
        assert '已选 4 家' in box.text();QTest.mouseClick(next(b for b in box.buttons() if b.text()=='确认并导出'),Qt.MouseButton.LeftButton);soon(batch_folder)
    def batch_folder(tries=0):
        dialog=QApplication.activeModalWidget()
        if not isinstance(dialog,QFileDialog):assert tries<500;soon(lambda:batch_folder(tries+1));return
        folder=out/'batch';folder.mkdir();dialog.setDirectory(str(folder));dialog.accept();state['batch_confirmed']=True
    def tick():
        try:
            if time.monotonic()-state['start']>240:raise TimeoutError(str(state))
            if QApplication.activeModalWidget() is not None:return
            step=state['step']
            if step==0:
                win.show_section(1)
                for w in wc.WAREHOUSES:win.catalog_imported(wc.import_catalog(fixture/(w+'.xlsx'),w))
                win.new_focus=focus.empty(wr.today()[:7]);win.refresh_focus();win.sales.add_paths([str(fixture/'five.xlsx')]);win.concurrency.setValue(4);win.start_load();state['step']=1
            elif step==1 and win.worker is None:
                assert len(win.store_widgets)==5
                for w in win.store_widgets.values():w['mode'].setCurrentIndex(1);w['primary'].setCurrentIndex(w['primary'].findData('河北仓'));w['secondary'].setCurrentIndex(w['secondary'].findData('山东仓'))
                win.start_calculation();state['step']=2
            elif step==2 and win.worker is None:
                assert len(win.results)==5 and all(s['state']=='已完成' for s in win.states.values());assert not win.orders.list();assert win.concurrency.value()==4
                assert '自动补位' in win.run_note.text();state['checks'].append('actual EXE configured four workers, all five stores completed before any archive/export, queue guidance visible')
                win.current_store='QA_0';state['step']=3;soon(lambda:confirmation(False,'cancel.xlsx'));win.export_current(['QA_0'])
            elif step==3 and win.worker is None:
                assert not win.orders.list() and not (out/'cancel.xlsx').exists();state['step']=4;soon(lambda:confirmation(True,'first.xlsx'));win.export_current(['QA_0'])
            elif step==4 and win.worker is None and (out/'first.xlsx').exists():
                order=win.orders.get(wr.today(),'QA_0');assert order['revision']==1 and not win.job_dirty;state['checks'].append('one-action unarchived delivery: default-cancel leaves DB/files unchanged; accept signs, saves draft, and writes numbered template without separate archive action')
                win.results['QA_0']['evidence_note']='second revision';win.touch();state['step']=5;soon(lambda:confirmation(False,'cancel2.xlsx'));win.export_current(['QA_0'])
            elif step==5 and win.worker is None:
                assert win.orders.get(wr.today(),'QA_0')['revision']==1;state['step']=6;soon(lambda:confirmation(True,'second.xlsx'));win.export_current(['QA_0'])
            elif step==6 and win.worker is None and (out/'second.xlsx').exists():
                assert win.orders.get(wr.today(),'QA_0')['revision']==2;assert win.orders.get(wr.today(),'QA_0',1);state['checks'].append('changed same-day order requires explicit cancel/confirm; same number V2 with V1 readable')
                import pickle,zlib
                real=pickle.loads(zlib.decompress((fixture/'real.pklz').read_bytes()));state['real']=real;state['rows']=[];state['step']=7
                win.job(x.export_xlsx,(real,str(out/'cancel-internal.xlsx')),lambda _:None)
                def progress(current,total):
                    state['rows'].append((current,total))
                    if current>=500 and total>500 and not state.get('cancelled'):
                        state['cancelled']=True;assert win.progress.maximum()==total;assert win.progress.value()==current;assert win.grab().save(str(out/'actual-row-progress.png'));win.cancel_work()
                win.worker.row_progress.connect(progress)
            elif step==7 and win.worker is None:
                assert state.get('cancelled') and not (out/'cancel-internal.xlsx').exists();state['checks'].append('actual measured row progress updates bar; cancel during row writing leaves no final file and reclaims workbook XML')
                state['step']=8;win.job(x.export_xlsx,(state['real'],str(out/'retry-internal.xlsx')),lambda _:None)
            elif step==8 and win.worker is None:
                assert (out/'retry-internal.xlsx').exists();assert beats[0]>20;state['checks'].append('retry after cancellation succeeds; Qt event loop remains active during background sign/export');state['step']=9
                def pick():
                    from merchant_selection import MerchantSelection
                    dialog=QApplication.activeModalWidget();assert isinstance(dialog,MerchantSelection);assert len(dialog.codes())==5
                    dialog.select_all(False);assert not dialog.ok.isEnabled();dialog.select_all(True);assert len(dialog.codes())==5
                    dialog.items[4][1].setCheckState(Qt.CheckState.Unchecked);assert dialog.codes()==['QA_0','QA_1','QA_2','QA_3'];dialog.grab().save(str(out/'store-picker.png'));dialog.accept();soon(batch_confirm)
                soon(pick);win.export_current()
            elif step==9 and win.worker is None and state.get('batch_confirmed'):
                files=list((out/'batch').glob('*.xlsx'));assert len(files)==4;assert win.orders.get(wr.today(),'QA_4') is None
                import openpyxl
                for file in files:
                    book=openpyxl.load_workbook(file);assert book.sheetnames==['往期新品','本月新品','中心仓动销补货清单']
                    for sheet in book:
                        assert sheet.max_column==22 and sheet['V1'].value=='决策理由';assert not any(row[18]=='地采' for row in sheet.iter_rows(min_row=2,values_only=True))
                    book.close()
                state['checks'].append('visible store picker defaults all five; deselect-all disables Next; choose four; one confirmation and folder selection writes four separate central-only 22-column workbooks; unselected store remains unissued; normal close');finish()
        except Exception:finish(traceback.format_exc())
    timer.timeout.connect(tick);timer.start(150)
