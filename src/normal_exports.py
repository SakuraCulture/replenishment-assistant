"""Merchant template: central-warehouse rows; audit detail stays in the archive."""
import os,re,tempfile
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,Color
from openpyxl.comments import Comment
from openpyxl.packaging.custom import StringProperty
import result_semantics as sem
import new_products
import new_product_focus as focus
import export_progress as ep
from io_stages import stage

HEADERS=['仓库/门店名称','仓库/门店编码','供应商编码','商品名称','商品编码','商品条码','基础单位','日销','周销','月销','可用库存','在途库存','在途库存天数','缺货天数（天）','箱规','中心仓价格','补货数量','黑名单状态','所属仓','商品等级']
AUDIT_HEADERS=['补货单号','决策理由']
SHEETS=['往期新品','本月新品','中心仓动销补货清单']
TEMPLATE_SHA256='1cc7a8e9f99fec764d17286c6b0a028a7ea79416add7178a64e91512f311f7c6'


def filename(order,warehouse=None):
    name=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',order['payload']['result']['store_name']).rstrip(' .')[:100] or order['store']
    if warehouse:
        from warehouse_orders import child_number
        child_number('',warehouse)
    return order['day']+name+'（精简版'+('·'+warehouse if warehouse else '')+'）.xlsx'


def groups(result,profile):
    """New pool has precedence, including anomalous sales rows: no duplicate SKU."""
    focused=focus.members(profile);out={name:[] for name in SHEETS}
    for d in result['rows']:
        if d.get('_source')!='资料匹配' or not sem.work_eligible(d):continue
        if new_products.is_new(d):out['本月新品' if d['E'] in focused else '往期新品'].append(d)
        elif sem.orderable(d):out['中心仓动销补货清单'].append(d)
    return out


def decision(d):
    """Short explanation of existing decisions only; never changes quantity."""
    excluded=sem.exclusion_reason(d)
    if excluded:return '禁用不补' if '禁用' in excluded else '停采不补'
    if sem.issue_text(d):return '资料待核'
    if d.get('_budget_overrun'):return '预算补足'
    if d.get('_budget_gap',0)>0:return '预算限量' if sem.orderable(d) else '预算暂缓'
    if new_products.is_new(d):
        if d.get('_new_requested',0)>0:return '必卖首铺' if d.get('_must_sell') else '新品首铺'
        return '新品未选'
    if d.get('_audit',{}).get('applied'):return '一箱恢复'
    if d.get('AW')=='极端缺货自动拉回':return '极端缺货恢复'
    # Counterfactual quantity without the holding floor, from persisted model
    # intermediates. Do not re-run prediction or infer from AS == BQ alone.
    keys=('AR','BI','BP','BQ','BR','BU','AS')
    if all(sem.numeric(d.get(k)) for k in keys):
        ordinary=d['AR'] if d['BI']==1 else 0
        without_floor=max(0,min(ordinary,d['BU']))
        raised=d['AS']>without_floor and (d['BQ']>without_floor or d['BP']>ordinary)
        if raised:
            if d.get('J')==0 and sem.numeric(d.get('N')) and d['N']>0:return '零销恢复'
            return '最低持有' if d['BQ']>without_floor else '断货保底'
    elif sem.numeric(d.get('BQ')) and d['BQ']>0:return '补货建议'
    return '必卖补货' if d.get('_must_sell') else '动销补货'


def export(order,path,cancel=None,progress=lambda s:None,*,warehouse=None):
    from engine import checkpoint
    path=Path(path)
    if path.exists():raise FileExistsError('文件已存在，请选择新文件名；不会覆盖已发出的单据')
    checkpoint(cancel);result=order['payload']['result'];profile=order['payload'].get('new_focus');focus.require_for(profile,order['day'],[result])
    parent_number=order['number'];grouped=groups(result,profile)
    if warehouse:
        from warehouse_orders import partitions,child_number
        grouped=partitions(result,profile).get(warehouse)
        if grouped is None:raise ValueError('该修订没有可导出的'+warehouse+'商品')
        order={**order,'number':child_number(parent_number,warehouse)}
    wb=Workbook();wb.remove(wb.active)
    rev=f"V{order['revision']:03d}";historical=order['revision']!=order['current_revision'];font=Font(name='宋体',sz=11,charset=134,scheme='minor',color=Color(theme=1))
    wb.properties.title=order['number']+' '+rev;wb.properties.subject='内部签发建议，非实际采购订单';wb.properties.description='仅中心仓商品；Q列可供门店调整；完整签发建议与来源保存在应用历史归档。数量为库存单位，价格为基本售价。'
    for key,value in [('补货单号',order['number']),('修订',rev),('模板SHA256',TEMPLATE_SHA256),('版本状态','历史修订' if historical else '当前修订')]:wb.custom_doc_props.append(StringProperty(name=key,value=value))
    if warehouse:
        for key,value in [('门店正常单号',parent_number),('供货仓',warehouse)]:wb.custom_doc_props.append(StringProperty(name=key,value=value))
    for name,rows in grouped.items():
        ep.rows(progress,name,0,len(rows));ws=wb.create_sheet(name);ws.sheet_format.defaultColWidth=9;ws.sheet_format.defaultRowHeight=13.5;ws.append(HEADERS+AUDIT_HEADERS)
        ws.oddHeader.center.text=order['number']+' '+rev+(' 历史修订' if historical else '')
        ws['A1'].comment=Comment(order['number']+' '+rev+'；业务日期 '+order['day']+'；'+('历史修订' if historical else '当前修订'),'补货助手')
        ws['Q1'].comment=Comment('签发执行数量（库存单位）；未选择或预算未安排的新品为0。门店可修改本列，原签发建议保存在应用历史归档。箱规品请填写箱规倍数。','补货助手')
        ws['P1'].comment=Comment('取商品资料基本售价；缺价保持空白或待核，不补造0。','补货助手')
        for rn,d in enumerate(rows,2):
            checkpoint(cancel);qty=sem.order_quantity(d) if sem.orderable(d) else 0
            reason=sem.exclusion_reason(d);status=d.get('AT')
            if status=='未提供（不作为数量拦截）':status='未提供'
            visible=[d.get(k) for k in 'A B C D E F G H I J K L M N'.split()]+[d.get('BD'),d.get('V'),qty,status,d.get('AU'),d.get('AV')]
            explanation=sem.explanation(d)
            audit=[order['number'],decision(d)]
            ws.append(visible+audit)
            ws.cell(rn,16).number_format='0.00'
            ws.cell(rn,18).number_format='@'
            if reason or qty==0:ws.cell(rn,17).comment=Comment(explanation,'补货助手')
            if (rn-1)%500==0:ep.rows(progress,name,rn-1,len(rows))
        ep.rows(progress,name,len(rows),len(rows));progress(name+'：整理格式…')
        for row in ws:
            checkpoint(cancel)
            for cell in row:
                cell.font=font
                if isinstance(cell.value,str):cell.data_type='s'
        # Original A:T layout retained, with explicit price/status presentation.
        ws.column_dimensions['P'].width=12;ws.column_dimensions['R'].width=14
        ws.column_dimensions['U'].width=38;ws.column_dimensions['V'].width=14
        ws.print_options.horizontalCentered=False
    progress('数据行写入完成，正在压缩保存确认单；此阶段暂不能中断…');path.parent.mkdir(parents=True,exist_ok=True)
    fd,temp=tempfile.mkstemp(prefix='.normal-',suffix='.xlsx',dir=path.parent);os.close(fd)
    try:
        with stage('压缩保存 Excel（完成后响应取消）',progress):wb.save(temp)
        checkpoint(cancel);os.rename(temp,path)
    finally:
        wb.close()
        if os.path.exists(temp):os.unlink(temp)
    return str(path)
