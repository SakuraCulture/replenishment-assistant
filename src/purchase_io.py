"""Purchase snapshot input. No model calls, exact codes, immutable source rows."""
import hashlib,io,math,uuid
from pathlib import Path
from datetime import date,datetime
from decimal import Decimal,InvalidOperation,localcontext
import openpyxl
import input_adapter as adapter
from runtime_support import checkpoint
from export_progress import RowProgress

FIELDS=['来源系统','采购订单号','采购明细号','下单日期','门店编码','供货仓','商品编码','商品名称','本版实际下单数量','数量单位','折合库存单位系数','换算依据','补货单号','签发修订','实际含税单价','实际金额','币种','订单状态','备注']
REQUIRED=['来源系统','采购订单号','下单日期','门店编码','供货仓','商品编码','本版实际下单数量','数量单位']

def number(value,optional=False):
    if value in (None,'') and optional:return None
    if isinstance(value,bool):raise ValueError('数量/金额不能是布尔值')
    try:d=Decimal(str(value))
    except (InvalidOperation,ValueError):raise ValueError('需要明确数字，不能将空白当0') from None
    if not d.is_finite() or d<0:raise ValueError('需要有限非负数字；撤单/退货不能用负数猜测')
    if not math.isfinite(float(d)):raise ValueError('数值超出可显示和导出的范围')
    with localcontext() as ctx:
        ctx.prec=max(28,len(d.as_tuple().digits))
        return str(d.normalize()) if d else '0'

def code(value,optional=False):
    if value in (None,'') and optional:return ''
    value,error=adapter.identifier(value)
    if error:raise ValueError(error)
    return value

def digest(data):return hashlib.sha256(data).hexdigest()

def inspect(paths,progress=lambda _:None,cancel=None):
    out=[]
    for index,p in enumerate(dict.fromkeys(str(Path(p).resolve()) for p in paths)):
        checkpoint(cancel);progress(RowProgress(f'识别采购文件 {index+1}/{len(paths)}：{Path(p).name}',index,len(paths)))
        raw=Path(p).read_bytes();sha=digest(raw);book=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
        sheets=[]
        try:
            for ws in book.worksheets:
                ws.reset_dimensions();it=ws.values;headers=list(next(it,()));sample=[]
                for _,row in zip(range(3),it):sample.append(list(row))
                sheets.append({'sheet':ws.title,'headers':headers,'sample':sample})
        finally:book.close()
        if adapter.digest(p)!=sha:raise ValueError('读取期间文件变化：'+p)
        out.append({'path':p,'sha':sha,'name':Path(p).name,'sheets':sheets})
        progress(RowProgress(f'已识别 {index+1}/{len(paths)} 个文件',index+1,len(paths)))
    return out

def scan(specs,progress=lambda _:None,cancel=None):
    """spec mapping values are zero-based column positions, never duplicate names."""
    rows=[];files=[]
    for fi,spec in enumerate(specs):
        checkpoint(cancel);p=Path(spec['path']);raw=p.read_bytes()
        if digest(raw)!=spec['sha']:raise ValueError('映射后文件变化，请重新读取：'+p.name)
        files.append({'path':str(p),'name':p.name,'sha':spec['sha']})
        mapping=spec['mapping'];defaults=spec.get('defaults',{})
        used=[i for i in mapping.values() if i is not None]
        if len(used)!=len(set(used)):raise ValueError('同一源列不能映射多个业务字段')
        for f in REQUIRED:
            if mapping.get(f) is None and not defaults.get(f):raise ValueError('缺少字段映射：'+f)
        book=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False);cached=None
        try:
            ws=book[spec['sheet']];ws.reset_dimensions();it=ws.values;headers=list(next(it,()))
            if headers!=spec['headers']:raise ValueError('表头已变化，请重新映射')
            for n,values in enumerate(it,2):
                checkpoint(cancel)
                if not any(v not in (None,'') for v in values):continue
                try:
                    d={f:(values[i] if i is not None and i<len(values) else defaults.get(f)) for f in FIELDS for i in [mapping.get(f)]}
                    if any(isinstance(v,str) and v.startswith('=') for v in d.values()):
                        if cached is None:cached=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
                        cv=next(cached[spec['sheet']].iter_rows(min_row=n,max_row=n,values_only=True))
                        for f,v in list(d.items()):
                            if isinstance(v,str) and v.startswith('='):
                                i=mapping[f];d[f]=cv[i] if i<len(cv) else None
                                if d[f] is None:raise ValueError(f+'公式没有缓存值，请在Excel计算保存后重试')
                    day=d['下单日期']
                    if isinstance(day,datetime):day=day.date().isoformat()
                    elif isinstance(day,date):day=day.isoformat()
                    else:day=str(day).strip();date.fromisoformat(day)
                    wh=str(d['供货仓'] or '').strip()
                    if wh not in ('河北仓','山东仓'):raise ValueError('供货仓须明确为河北仓或山东仓')
                    qty=number(d['本版实际下单数量']);unit=str(d['数量单位'] or '').strip()
                    if not unit:raise ValueError('数量单位不能为空')
                    factor=number(d['折合库存单位系数'],True);basis=str(d['换算依据'] or '').strip()
                    if factor is not None and (Decimal(factor)<=0 or not basis):raise ValueError('换算系数须大于0且填写换算依据')
                    if unit=='库存单位' and factor is not None and Decimal(factor)!=1:raise ValueError('已按库存单位填写，换算系数不能同时不是1')
                    with localcontext() as ctx:
                        ctx.prec=max(28,len(Decimal(qty).as_tuple().digits)+len(Decimal(factor or 1).as_tuple().digits))
                        base=qty if unit=='库存单位' else number(Decimal(qty)*Decimal(factor)) if factor else None
                    revision=d['签发修订'];rev=None
                    if revision not in (None,''):
                        rv=str(revision).strip().lstrip('Vv');dec=Decimal(rv)
                        if not dec.is_finite() or dec<=0 or dec!=dec.to_integral_value():raise ValueError('签发修订须为正整数或V001形式')
                        rev=int(dec)
                    external=code(d['采购明细号'],True)
                    source={'sha':spec['sha'],'file':p.name,'sheet':spec['sheet'],'row':n}
                    lid=external or 'local:'+digest((spec['sha']+'\0'+spec['sheet']+'\0'+str(n)).encode())[:24]
                    status=str(d['订单状态'] or '有效').strip()
                    if status not in ('有效','已取消'):raise ValueError('订单状态仅接受有效/已取消；不推断其它系统状态')
                    rows.append({'system':code(d['来源系统']),'number':code(d['采购订单号']),'store':code(d['门店编码']),'warehouse':wh,'day':day,'line_id':lid,'external_line_id':bool(external),'sku':code(d['商品编码']),'name':str(d['商品名称'] or ''),'quantity':qty,'unit':unit,'factor':factor,'conversion_basis':basis,'base_quantity':base,'signed_number':code(d['补货单号'],True),'signed_revision':rev,'price':number(d['实际含税单价'],True),'amount':number(d['实际金额'],True),'currency':str(d['币种'] or ''),'status':status,'note':str(d['备注'] or ''),'source':source,'raw':d,'links':[]})
                except (ValueError,InvalidOperation) as exc:raise ValueError(f'{p.name} / {ws.title} 第{n}行：{exc}') from exc
                if n%200==0:progress(RowProgress(f'已完成 {fi}/{len(specs)} 个文件；当前文件已读取 {n-1} 行',fi,len(specs)))
        finally:
            book.close()
            if cached:cached.close()
        if adapter.digest(p)!=spec['sha']:raise ValueError('读取期间文件变化：'+p.name)
        progress(RowProgress(f'读取校验已完成 {fi+1}/{len(specs)} 个文件',fi+1,len(specs)))
    if not rows:raise ValueError('没有采购明细')
    return {'rows':rows,'files':files}

def template(path):
    target=Path(path)
    if target.exists():raise FileExistsError('文件已存在，请换新文件名')
    book=openpyxl.Workbook();ws=book.active;ws.title='实际下单明细';ws.append(FIELDS);ws.freeze_panes='A2'
    from openpyxl.styles import Font
    for c in ws[1]:c.font=Font(bold=True);ws.column_dimensions[c.column_letter].width=22
    for n in range(2,202):
        for i,f in enumerate(FIELDS,1):
            if '编码' in f or '单号' in f or f in ('采购明细号','签发修订'):ws.cell(n,i).number_format='@'
    note=book.create_sheet('填写说明')
    for row in [('必需列','、'.join(REQUIRED)),('口径','每行是本版订单总量，不是本次增量；同采购单更版完整替换，不累计。'),('单位','数量单位填写库存单位；其它单位暂存待核，或提供明确换算系数及依据。'),('空值','实际下单数量不能空；0表示明确为0；缺金额仍为空。'),('身份','编码按文本保留前导零；采购明细号强烈建议填写。'),('修订','补货单号按仓子单号，签发修订填写V001；缺失时人工选择，不能自动用最新版。'),('更新','无稳定明细号的更新必须逐行确认映射或明确作为新增行；不按行序猜。'),('取消','订单状态填写有效或已取消；不把负数、缺行视为取消。'),('样本','模板不含示例业务行；源文件保持不变。')]:note.append(row)
    note.column_dimensions['A'].width=20;note.column_dimensions['B'].width=100
    import os,tempfile
    fd,tmp=tempfile.mkstemp(dir=target.parent,suffix='.xlsx');os.close(fd)
    try:book.save(tmp);os.link(tmp,target)
    finally:book.close();Path(tmp).unlink(missing_ok=True)
