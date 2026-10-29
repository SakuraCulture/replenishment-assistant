"""Reusable must-sell membership; sales, stock and price are never read here."""
import csv,hashlib,json,os,uuid
from pathlib import Path
from datetime import datetime
import openpyxl
from input_adapter import identifier

def atomic_json(path,data):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    tmp=path.with_name('.'+path.name+'.'+uuid.uuid4().hex)
    try:
        tmp.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
        os.replace(tmp,path)
    finally:
        if tmp.exists():tmp.unlink()

def make_profile(items,name,source_sha='',ignored=0,duplicates=0):
    if not items:raise ValueError('没有找到必卖商品，请检查商品编码及清单归属。')
    normalized=sorted(items,key=lambda x:x['sku'])
    revision=hashlib.sha256(json.dumps(normalized,ensure_ascii=False,sort_keys=True).encode()).hexdigest()[:12]
    return {'schema':1,'name':name,'revision':revision,'saved_at':datetime.now().isoformat(timespec='seconds'),
        'source_sha256':source_sha,'ignored_rows':ignored,'duplicate_rows':duplicates,'items':normalized}

def validate_profile(data):
    if not isinstance(data,dict) or data.get('schema')!=1 or not isinstance(data.get('items'),list):raise ValueError('必卖清单格式不正确')
    if not isinstance(data.get('name'),str) or not isinstance(data.get('revision'),str):raise ValueError('必卖清单缺少名称或版本')
    seen=set()
    for item in data['items']:
        if not isinstance(item,dict):raise ValueError('必卖清单条目格式错误')
        if not isinstance(item.get('name',''),str) or not isinstance(item.get('category',''),str):raise ValueError('清单商品名称或分类格式错误')
        sku,err=identifier(item.get('sku'))
        if err or not isinstance(item.get('sku'),str) or sku!=item['sku']:raise ValueError('必卖清单存在不可靠商品编码')
        if sku in seen:raise ValueError('必卖清单存在重复商品编码')
        seen.add(sku)
    if not seen:raise ValueError('必卖清单为空')
    actual=hashlib.sha256(json.dumps(sorted(data['items'],key=lambda x:x['sku']),ensure_ascii=False,sort_keys=True).encode()).hexdigest()[:12]
    if actual!=data['revision']:raise ValueError('清单内容与版本不一致，请重新导入原Excel清单')
    return data

def import_profile(path,progress=lambda s:None,cancel=None):
    p=Path(path);progress('正在读取必卖清单…')
    if cancel and cancel.is_set():
        from engine import Cancelled
        raise Cancelled()
    if p.suffix.lower()=='.json':return validate_profile(json.loads(p.read_text(encoding='utf-8-sig')))
    sha=hashlib.sha256(p.read_bytes()).hexdigest();book=None
    if p.suffix.lower()=='.csv':
        with p.open(encoding='utf-8-sig',newline='') as f:raw=list(csv.reader(f))
    elif p.suffix.lower()=='.xlsx':
        book=openpyxl.load_workbook(p,read_only=True,data_only=True,keep_links=False)
        try:
            candidates=[]
            for ws in book.worksheets:
                ws.reset_dimensions();heads=next(ws.iter_rows(values_only=True),())
                names=[str(x or '').strip() for x in heads]
                if '商品编码' in names and not any(x in names for x in ['门店编码','仓库/门店编码']):candidates.append(ws)
            if len(candidates)!=1:raise ValueError('请提供一张含“商品编码”的商品清单表，不使用门店-商品重复明细。')
            raw=list(candidates[0].iter_rows(values_only=True))
        finally:book.close()
    else:raise ValueError('清单支持.xlsx、.csv或本软件导出的.json')
    if not raw:raise ValueError('清单为空')
    heads=[str(x or '').strip() for x in raw[0]]
    if heads.count('商品编码')!=1:raise ValueError('需要唯一“商品编码”列；长编码请以文本格式保存')
    selected={};ignored=duplicates=0
    for rowno,row in enumerate(raw[1:],2):
        if cancel and cancel.is_set():
            from engine import Cancelled
            raise Cancelled()
        if not any(x is not None and str(x).strip() for x in row):continue
        d=dict(zip(heads,row))
        if '清单归属' in heads and str(d.get('清单归属') or '').strip() not in {'必卖','必卖+冬季','冬季+必卖','必卖清单'}:
            ignored+=1;continue
        sku,err=identifier(d.get('商品编码'))
        if err:raise ValueError(f'第{rowno}行商品编码待核：{err}；请使用文本编码')
        item={'sku':sku,'name':str(d.get('商品名称') or '').strip(),'category':str(d.get('一级分类') or d.get('分类') or '').strip()}
        if sku in selected:
            if selected[sku]['name']!=item['name']:raise ValueError(f'商品编码{sku}出现不同名称，请先核对重复行')
            duplicates+=1
        else:selected[sku]=item
    if hashlib.sha256(p.read_bytes()).hexdigest()!=sha:raise ValueError('读取期间清单发生变化，请重新导入')
    return make_profile(list(selected.values()),p.stem,sha,ignored,duplicates)

class Repository:
    def __init__(self,folder):self.folder=Path(folder)
    def profile(self):
        p=self.folder/'must_sell.json'
        if p.exists():return validate_profile(json.loads(p.read_text(encoding='utf-8')))
        data=validate_profile(json.loads((Path(__file__).parent/'default_must_sell.json').read_text(encoding='utf-8')))
        self.save(data);return data
    def save(self,profile):atomic_json(self.folder/'must_sell.json',validate_profile(profile))
    def budgets(self):
        p=self.folder/'store_budgets.json'
        data=json.loads(p.read_text(encoding='utf-8')) if p.exists() else {}
        if not isinstance(data,dict):raise ValueError('门店预算配置格式错误')
        for key,value in data.items():
            if not isinstance(value,dict) or not isinstance(value.get('enabled'),bool) or isinstance(value.get('amount'),bool) or not isinstance(value.get('amount'),(int,float)) or not 0<=value['amount']<=1000000000:raise ValueError('门店预算配置值无效')
        return data
    def save_budget(self,store,enabled,amount):
        data=self.budgets();data[store]={'enabled':bool(enabled),'amount':float(amount)}
        atomic_json(self.folder/'store_budgets.json',data)

def export_profile(profile,path):
    p=Path(path)
    if p.exists():raise ValueError('文件已存在，请选择新文件名')
    validate_profile(profile)
    if p.suffix.lower()=='.json':
        with p.open('x',encoding='utf-8') as f:json.dump(profile,f,ensure_ascii=False,indent=2)
    else:raise ValueError('请选择.json文件；该格式保留长编码，适合备份和迁移')
    return str(p)
