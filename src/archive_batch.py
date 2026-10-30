"""Read-only, revision-bound archive exports. No model execution or source mutation."""
import hashlib, io, os, re, tempfile, uuid
from pathlib import Path
from contextlib import closing
from collections import Counter
import openpyxl
import archive_exports as legacy
import archive_orders as ao
import normal_exports as ne
import result_semantics as sem
import history_forms as forms
import source_archive
from history_read import repository
from runtime_support import checkpoint, Cancelled
from xlsx_exports import stream_book, write_table
from export_progress import RowProgress

TITLES={'calculation':'完整计算表','signed':'签发建议','original':'商品资料原文件','reference':'商品资料核对表（加工表）'}
IDENTITY=['订单 / 业务日期','订单 / 门店编码','订单 / 门店名称','订单 / 仓库','订单 / 补货单号','订单 / 修订']
MAX_ROWS=1048576
MAX_COLUMNS=16384

def key(h):
    return (h['oid'],h['warehouse'],h['revision'],h['digest'])

def safe(value):
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',str(value)).rstrip(' .')[:100]

def load(repo,head,cancel=None):
    checkpoint(cancel)
    with closing(repo.work.connect()) as db:ao.require_active(db,head['oid'],head['warehouse'])
    order=repo.get(head['day'],head['store'],head['revision'])
    if order is None or order['digest']!=head['digest'] or order['oid']!=head['oid']:
        raise ValueError('归档修订或指纹不一致，请重新查询')
    if head['warehouse'] not in ao.scopes(order['payload']['result']['rows']):raise ValueError('此修订没有该仓数据')
    return order

def source(repo,order,wh):
    parts=forms.work_parts(repo,order,[('catalogs',)])
    catalog=parts.get(('catalogs',),{}).get(wh)
    if not catalog:raise ValueError('历史未留存可识别的本仓原文件；不能以当前资料替代')
    with closing(repo.work.connect()) as db:stored=source_archive.read(db,catalog.get('source_sha256'))
    if not stored:raise ValueError('历史原文件未留存；不能以字段说明或当前资料替代')
    return catalog,stored

def signed(order,wh):
    rows=legacy.signed_rows(order,wh)
    return ['推荐类型']+ne.HEADERS+ne.AUDIT_HEADERS,iter(rows),len(rows)


def reference(repo,order,wh,cancel):
    catalog,stored=source(repo,order,wh)
    needed=[d for d in order['payload']['result']['rows'] if ao.warehouse(d)==wh]
    positions={}
    for d in needed:
        item=catalog.get('items',{}).get(d['E'])
        meta=d.get('_meta',{})
        if not item or not isinstance(item.get('_row'),int):raise ValueError('缺少商品原行绑定：'+d['E'])
        if meta.get('product_row') is not None and meta['product_row']!=item['_row']:raise ValueError('商品原行绑定冲突：'+d['E'])
        positions[item['_row']]=None
    book=openpyxl.load_workbook(io.BytesIO(stored[1]),read_only=True,data_only=False,keep_links=False)
    try:
        ws=book[catalog['sheet']];ws.reset_dimensions();it=ws.values;headers=list(next(it,()))
        for line,values in enumerate(it,2):
            checkpoint(cancel)
            if line in positions:positions[line]=list(values)
        if any(v is None for v in positions.values()):raise ValueError('原文件缺少绑定的商品行')
    finally:book.close()
    width=max([len(headers)]+[len(v) for v in positions.values()]);headers+=['']*(width-len(headers))
    hs=['商品编码','商品名称','资料版本','原工作表','原行号']+[f'原表第{i+1}列 / {h or "空表头"}' for i,h in enumerate(headers)]
    def values():
        for d in needed:
            line=catalog['items'][d['E']]['_row'];v=positions[line]
            yield [d['E'],d.get('D'),catalog.get('revision'),catalog['sheet'],line]+v+[None]*(width-len(v))
    return hs,values(),len(needed)

def table(repo,order,wh,kind,cancel):
    if kind=='calculation':return legacy.calculation(order,wh)
    if kind=='signed':return signed(order,wh)
    if kind=='reference':return reference(repo,order,wh,cancel)
    raise ValueError('未知导出内容')

def columns(headers):
    # Preserve repeated labels positionally; never collapse them via dict(header).
    seen=Counter();result=[]
    for h in headers:seen[h]+=1;result.append((h,seen[h]))
    return result

def prepare(path,heads,kind,mode,progress=lambda _:None,cancel=None):
    if kind not in TITLES or mode not in ('separate','combined'):raise ValueError('未知导出方式')
    if not heads:raise ValueError('请先勾选仓单')
    unique=set();orders=set();items=[];repo=repository(path)
    for i,head in enumerate(heads):
        checkpoint(cancel)
        if key(head) in unique:raise ValueError('重复选择同一仓单')
        if (head['oid'],head['warehouse']) in orders:raise ValueError('同一仓单不能同时选择多个修订')
        unique.add(key(head));orders.add((head['oid'],head['warehouse']))
        progress(RowProgress(f'核对仓单 {i+1}/{len(heads)} · {head["store"]} · {head["warehouse"]}',i,len(heads)))
        item={'head':dict(head),'error':None}
        try:
            order=load(repo,head,cancel)
            if kind=='original':
                c,stored=source(repo,order,head['warehouse'])
                item.update(sha=c['source_sha256'],name=stored[0],count=1,columns=[])
            else:
                hs,rs,n=table(repo,order,head['warehouse'],kind,cancel)
                item.update(columns=columns(hs),count=n)
                if kind=='reference':list(legacy.checked_rows(rs))
        except Cancelled:raise
        except Exception as exc:item['error']=str(exc)
        items.append(item)
    progress(RowProgress('核对完成',len(heads),len(heads)))
    return {'kind':kind,'mode':mode,'items':items}

def groups(plan):
    result={}
    for i,item in enumerate(plan['items']):
        if item['error']:continue
        h=item['head']
        k=i if plan['mode']=='separate' else (h['warehouse'],item['sha']) if plan['kind']=='original' else h['warehouse']
        result.setdefault(k,[]).append(item)
    return list(result.values())

def identity(order,wh):
    return [order['day'],order['store'],order['payload']['result']['store_name'],wh,ao.number(order['number'],wh),f'V{order["revision"]:03d}']

def write_group(repo,plan,items,target,progress,cancel):
    kind=plan['kind'];wh=items[0]['head']['warehouse'];checkpoint(cancel)
    if kind=='original':
        data=None
        for item in items:
            order=load(repo,item['head'],cancel);catalog,stored=source(repo,order,wh)
            if catalog['source_sha256']!=item['sha']:raise ValueError('原文件绑定已变化')
            data=stored[1]
        fd,tmp=tempfile.mkstemp(dir=target.parent,prefix='.original-')
        try:
            with os.fdopen(fd,'wb') as f:f.write(data)
            checkpoint(cancel)
            with closing(repo.work.connect()) as db:
                for item in items:ao.require_active(db,item['head']['oid'],wh)
            os.link(tmp,target)
        finally:Path(tmp).unlink(missing_ok=True)
        return
    union=list(dict.fromkeys(c for item in items for c in item['columns']))
    total=sum(item['count'] for item in items)
    if total+1>MAX_ROWS or len(union)+len(IDENTITY)>MAX_COLUMNS:raise ValueError('超过Excel容量，请缩小范围或分别导出；未截断')
    def rows():
        for i,item in enumerate(items):
            progress(f'读取仓单 {i+1}/{len(items)} · {wh} · {item["head"]["store"]}')
            order=load(repo,item['head'],cancel);hs,rs,n=table(repo,order,wh,kind,cancel)
            cs=columns(hs)
            if cs!=item['columns'] or n!=item['count']:raise ValueError('预览后的字段或记录数发生变化')
            prefix=identity(order,wh);pos={c:j for j,c in enumerate(cs)}
            for values in rs:
                checkpoint(cancel)
                yield prefix+[values[pos[c]] if c in pos else None for c in union]
    fd,tmp=tempfile.mkstemp(dir=target.parent,prefix='.archive-',suffix='.xlsx');os.close(fd)
    try:
        with stream_book() as book:
            ws=write_table(book,TITLES[kind],IDENTITY+[c[0] for c in union],legacy.checked_rows(rows()),cancel=cancel,progress=progress,row_count=total)
            from openpyxl.utils import get_column_letter
            ws.auto_filter.ref=f'A1:{get_column_letter(len(union)+len(IDENTITY))}{total+1}'
            ws.freeze_panes='I2'
            progress('明细已写完，正在压缩保存；完成当前安全步骤后响应取消…');book.save(tmp);checkpoint(cancel)
            with closing(repo.work.connect()) as db:
                for item in items:ao.require_active(db,item['head']['oid'],wh)
            os.link(tmp,target)
    finally:Path(tmp).unlink(missing_ok=True)

def execute(path,plan,folder,allow_partial=False,progress=lambda _:None,cancel=None):
    missing=[i for i in plan['items'] if i['error']]
    if missing and not allow_partial:raise ValueError('存在不可导出单据，请明确选择仅导出可用部分')
    if not groups(plan):raise ValueError('没有可导出单据')
    root=Path(folder)/('归档导出-'+uuid.uuid4().hex[:12]);root.mkdir(parents=True,exist_ok=False)
    result={'directory':str(root),'files':[],'failed':[],'skipped':missing,'cancelled':[],'kind':plan['kind'],'mode':plan['mode']}
    repo=repository(path);units=groups(plan)
    for index,items in enumerate(units):
        if cancel is not None and cancel.is_set():
            result['cancelled'].extend(x for unit in units[index:] for x in unit);break
        h=items[0]['head'];wh=h['warehouse'];days=[i['head']['day'] for i in items]
        name=(h['day']+safe(h.get('name',h['store']))+'-'+safe(h['store'])+f'（{TITLES[plan["kind"]]}·{wh}·V{h["revision"]:03d}）') if plan['mode']=='separate' else (min(days)+'至'+max(days)+f'（{TITLES[plan["kind"]]}汇总·{wh}·{len(items)}单）')
        suffix='.xlsx'
        if plan['kind']=='original':
            suffix=Path(items[0]['name']).suffix or '.xlsx'
            name=(name if plan['mode']=='separate' else wh+'-'+safe(Path(items[0]['name']).stem))+'-'+items[0]['sha'][:12]
        target=root/(f'{index+1:03d}-'+name+suffix)
        progress(f'文件 {index+1}/{len(units)} · '+target.name)
        try:
            write_group(repo,plan,items,target,progress,cancel)
            result['files'].append({'path':str(target),'heads':[i['head'] for i in items],'sha256':hashlib.sha256(target.read_bytes()).hexdigest()})
        except Cancelled:
            result['cancelled'].extend(x for unit in units[index:] for x in unit);break
        except Exception as exc:result['failed'].extend({**i,'error':str(exc)} for i in items)
    lines=[f'导出内容：{TITLES[plan["kind"]]}',f'完成文件：{len(result["files"])}；失败仓单：{len(result["failed"])}；未留存/不可用：{len(missing)}；取消未执行：{len(result["cancelled"])}','文件绑定下列明确修订，不重新计算。缺少的历史字段留空，不能视为0；全店共享预算不可按行累加。']
    for f in result['files']:
        lines.append('\n文件：'+Path(f['path']).name)
        for h in f['heads']:lines.append(f'  {h["day"]} / {h["store"]} / {h["warehouse"]} / {h["number"]} / V{h["revision"]:03d}')
    for state in ('skipped','failed','cancelled'):
        for item in result[state]:lines.append(state+'：'+str(key(item['head']))+' / '+str(item.get('error') or '未执行'))
    (root/'导出清单.txt').write_text('\n'.join(lines),encoding='utf-8-sig')
    return result
