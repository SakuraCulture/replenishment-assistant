"""Desktop input, calculation and export service. No network or reference workbook."""
from __future__ import annotations
import copy,json,math,os,uuid,warnings
from pathlib import Path
from datetime import datetime
from collections import Counter
from decimal import Decimal
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter
import model
import input_adapter as adapter
from one_box_exception import apply_one_box_exception
import allocation

from version_info import VERSION,MODEL_VERSION,diagnostic
DEFAULT_WAREHOUSES={'河北仓':(8.0,4.0),'山东仓':(8.0,5.0)}
class Cancelled(Exception):pass
def checkpoint(cancel):
    if cancel and cancel.is_set():raise Cancelled('已取消')
def config():return json.loads((Path(__file__).parent/'default_config.json').read_text(encoding='utf-8-sig'))
def numeric(v):return model.number(v)
def fmt(v,digits=2):
    if v is None:return '—'
    if numeric(v):return f'{v:,.{digits}f}'.rstrip('0').rstrip('.') if digits else f'{v:,.0f}'
    return str(v)
def category_of(d,m):
    if m['match']=='商品编码精确匹配':return '资料匹配'
    if m['match']=='资料未匹配' and not any('编码' in b or '重复' in b for b in m['blocking'] if b!='待补商品资料'):return '地采商品'
    return '匹配待核'
def issue_text(d):
    m=d['_meta'];issues=[s for s in m['blocking'] if not (d['_source']=='地采商品' and s=='待补商品资料')]
    return '；'.join(issues)
def sheet_name(path,kind):
    warnings.filterwarnings('ignore',message='Workbook contains no default style')
    book=openpyxl.load_workbook(path,read_only=True,data_only=True,keep_links=False)
    try:
        required={'商品编码','日销','周销','月销','仓库/门店编码'} if kind=='sales' else {'商品编码','基本售价','分类','出库箱规'}
        candidates=[]
        for ws in book.worksheets:
            ws.reset_dimensions();row=next(ws.iter_rows(values_only=True),())
            heads={str(x).strip() for x in row if x is not None}
            if required<=heads:candidates.append((ws.title,heads))
        if len(candidates)!=1:
            if not candidates:raise ValueError(('补货建议' if kind=='sales' else '商品资料')+'未找到唯一可用工作表。必需列：'+ '、'.join(sorted(required)))
            raise ValueError('发现多个可用工作表：'+ '、'.join(x[0] for x in candidates)+'；请保留本次要使用的一张表。')
        return candidates[0]
    finally:book.close()

def load_inputs(sales,products,progress=lambda s:None,cancel=None):
    sales,products=Path(sales).resolve(),Path(products).resolve()
    if sales==products:raise ValueError('补货建议和商品资料必须选择两份不同的文件。')
    for p in [sales,products]:
        if not p.is_file() or p.suffix.lower()!='.xlsx':raise ValueError(f'请选择有效的 .xlsx 文件：{p.name}')
    before={str(p):adapter.digest(p) for p in [sales,products]};checkpoint(cancel)
    progress('正在识别工作表与字段…')
    ss,sh=sheet_name(sales,'sales');ps,ph=sheet_name(products,'products')
    cfg=config();cfg['interface'].update(sales_sheet=ss,product_sheet=ps,blank_outbound_unit='one')
    cfg['sales_fields']['unit_type']='基础单位' if '基础单位' in sh else '库存单位' if '库存单位' in sh else '单位类型'
    cfg['weight_profiles']['无分类']=copy.deepcopy(cfg['weight_profiles']['快销'])
    progress('正在读取商品资料…');checkpoint(cancel)
    pr,_=adapter.read_table(products,ps,['商品编码','分类','基本售价','出库箱规'])
    warehouses={};categories=set()
    for r in pr:
        k,err=adapter.identifier(r.get('商品编码'))
        if not err:warehouses[k]=str(r.get('仓归属') or '未填写仓归属').strip()
        categories.add(str(r.get('分类')))
    for c in categories:
        if c not in cfg['category_profiles']:cfg['category_profiles'][c]='无分类'
    progress('正在按商品编码匹配销量与资料…');checkpoint(cancel)
    norm,info=adapter.prepare(sales,products,cfg)
    stores={}
    for d,m in norm:
        d['_source']=category_of(d,m)
        if m['match']=='商品编码精确匹配':d['AU']=warehouses.get(d['E'],'未填写仓归属')
        elif d['_source']=='地采商品':d['AU']='地采'
        else:d['AU']='待核'
        if d.get('O')=='/':d['O']=1
        stores.setdefault(d['B'],d['A'] or d['B'])
    checkpoint(cancel)
    if any(adapter.digest(p)!=h for p,h in before.items()):raise ValueError('读取期间源文件发生变化，请保存Excel后重新导入。')
    return {'normalized':norm,'config':cfg,'info':info,'hashes':before,'sales':str(sales),'products':str(products),'stores':stores,
            'warehouses':sorted({d['AU'] for d,m in norm if d['_source']=='资料匹配'}),'loaded_at':datetime.now().isoformat(timespec='seconds')}

def preview(loaded,store):
    rows=[]
    for d,m in loaded['normalized']:
        if d['B']!=store:continue
        x=copy.deepcopy(d);x['_meta']=copy.deepcopy(m);x.update(AS=None,BV=None,_base_qty=None,_base_amount=None,_audit={},_preview=True)
        rows.append(x)
    return rows

def run_calculation(loaded,store,warehouses,enabled,budget=None,profile=None,progress=lambda s:None,cancel=None):
    rows=[copy.deepcopy(x) for x in loaded['normalized'] if x[0]['B']==store]
    if not rows:raise ValueError('所选门店没有商品。')
    p=model.Rules(**loaded['config']['parameters']);result=[];baseline=[]
    for i,(d,m) in enumerate(rows):
        checkpoint(cancel)
        if i%300==0:progress(f'正在计算 {i:,} / {len(rows):,} 件商品…')
        if d['_source']=='资料匹配':
            depth,lead=warehouses.get(d['AU'],(-1,-1))
            if depth<0 or lead<0 or depth+lead<=0:m['blocking'].append('所属仓的备货/到货天数未配置')
            else:d.update(AG=depth,AH=lead)
        b=adapter.execute([(d,m)],p)[0]
        x,a=apply_one_box_exception(b,m,enabled=enabled)
        baseline.append(b)
        if not a['applied'] and b!=x:raise RuntimeError('隔离检查失败：非例外商品发生变化')
        if a['applied']:
            assert b['AS']==0 and x['AS']==b['BD']
            assert all(b.get(k)==x.get(k) for k in set(b)|set(x) if k not in {'AS','BV','BM','BW'})
        x=copy.deepcopy(x);x.update(_audit=a,_base_qty=b['AS'],_base_amount=b['BV'],_preview=False)
        result.append(x)
    checkpoint(cancel);progress('正在核对数量、金额与例外隔离…')
    checks=adapter.business_checks(result,p)
    if checks['failures']:raise RuntimeError('结果核验失败：'+str(checks['failures'][:3]))
    if not any(numeric(d['AS']) for d in result):raise ValueError('没有可计算商品，请检查商品资料与仓参数。')
    report={'rows':result,'baseline':baseline,'store':store,'store_name':loaded['stores'][store],
        'warehouses':warehouses,'enabled':enabled,'config':loaded['config'],'sources':loaded['hashes'],
        'created':datetime.now().isoformat(timespec='seconds'),'checks':checks,'summary':summarize(result)}
    return apply_budget(report,budget,profile,progress,cancel)

def apply_budget(result,budget=None,profile=None,progress=lambda s:None,cancel=None):
    import priority_list
    if profile is not None:priority_list.validate_profile(profile)
    if budget is not None and profile is None:raise ValueError('启用预算前请先导入有效必卖清单')
    members={d['sku']:d for d in profile['items']} if profile else {}
    for d in result['rows']:
        for k in list(d):
            if k.startswith(('_budget_','_allocated_','_holding_','_candidate_','_must_')):d.pop(k)
        d['_must_sell']=d.get('E') in members
        d['_candidate_qty']=0 if exclusion_reason(d) else quantity(d)
        d['_candidate_amount']=0 if exclusion_reason(d) else amount_value(d)
    progress('正在按门店预算分配必卖与其他商品…' if budget is not None else '正在标记必卖清单…')
    result['budget']=allocation.apply(result['rows'],budget,lambda:checkpoint(cancel)) if budget is not None else {'enabled':False}
    by_code={d['E']:d for d in result['rows']}
    audit=[]
    for sku,item in members.items():
        d=by_code.get(sku)
        state='未出现在本次门店导出' if d is None else '已排除：'+exclusion_reason(d) if exclusion_reason(d) else '地采，资料未匹配' if d['_source']=='地采商品' else '数据待核' if issue_text(d) else '有补货需求' if numeric(d.get('AS')) and d['AS']>0 else '当前无补货建议'
        audit.append({'sku':sku,'name':item.get('name',''),'state':state})
    result['must_list']={'name':profile['name'],'revision':profile['revision'],'source_sha256':profile.get('source_sha256',''),'total':len(members),
        'matched':sum(d['_must_sell'] for d in result['rows']),'missing':sum(x['state']=='未出现在本次门店导出' for x in audit)} if profile else None
    result['must_list_audit']=audit
    result['summary']=summarize(result['rows'])
    return result

def exclusion_reason(d):
    reasons=[]
    if str(d.get('S') or '').strip()=='禁用':reasons.append('商品资料禁用')
    if str(d.get('_meta',{}).get('export_status') or '').strip()=='停采':reasons.append('补货建议停采')
    return '；'.join(reasons)

def order_quantity(d):
    return 0 if exclusion_reason(d) else d.get('_allocated_qty') if d.get('_budget_active') else quantity(d)

def order_amount(d):
    return 0 if exclusion_reason(d) else d.get('_allocated_amount') if d.get('_budget_active') else amount_value(d)

def orderable(d):
    return not exclusion_reason(d) and d.get('_source')=='资料匹配' and numeric(order_quantity(d)) and order_quantity(d)>0

def order_status(d):
    if exclusion_reason(d):return '已排除'
    if d.get('_preview'):return '待计算'
    if d.get('_budget_gap',0)>0:return '预算部分分配' if numeric(order_quantity(d)) and order_quantity(d)>0 else '预算暂缓'
    if orderable(d):return '纳入补货'
    if numeric(d.get('AS')):return '无需补货'
    return '地采未计算' if d.get('_source')=='地采商品' else '待核'

def summarize(rows):
    total=lambda values:float(sum((Decimal(str(v)) for v in values if numeric(v)),Decimal(0)).quantize(Decimal('.01')))
    excluded=[d for d in rows if exclusion_reason(d)]
    return {'rows':len(rows),'local':sum(d['_source']=='地采商品' for d in rows),
        'issues':sum(bool(issue_text(d)) for d in rows),
        'calculable':sum(numeric(d.get('AS')) for d in rows),
        'positive':sum(orderable(d) for d in rows),
        'amount':total(order_amount(d) for d in rows),
        'calculated_positive':sum(numeric(d.get('AS')) and d['AS']>0 for d in rows),
        'calculated_amount':total(d.get('BV') for d in rows),
        'excluded':len(excluded),'excluded_positive':sum(numeric(d.get('AS')) and d['AS']>0 for d in excluded),
        'excluded_amount':total(d.get('BV') for d in excluded),
        'unpriced':sum(orderable(d) and not numeric(d.get('BV')) for d in rows),
        'candidate_amount':total(d.get('_candidate_amount',order_amount(d)) for d in rows),
        'must_amount':total(order_amount(d) for d in rows if d.get('_must_sell')),
        'must_candidate_amount':total(d.get('_candidate_amount',order_amount(d)) for d in rows if d.get('_must_sell')),
        'budget_deferred':sum(d.get('_budget_gap',0)>0 for d in rows),'holding_gaps':sum(d.get('_holding_gap',False) for d in rows),
        'exception':sum(bool(d.get('_audit',{}).get('applied')) for d in rows),
        'extra':round(sum(d.get('_audit',{}).get('incremental_amount',0) for d in rows),2)}

def quantity(d):
    if d['_source']=='地采商品':return '未计算'
    return d.get('AS') if numeric(d.get('AS')) else '待核' if not d.get('_preview') else '待计算'
def amount_value(d):
    if d['_source']=='地采商品':return '未计价'
    return d.get('BV') if numeric(d.get('BV')) else '待核' if not d.get('_preview') else '待计算'
def coverage(d):
    if not all(numeric(d.get(k)) for k in ['AN','AM','K','L','AS']):return None
    rate=d['AN']*d['AM']
    return (max(d['K'],0)+max(d['L'],0)+d['AS'])/rate if rate>0 else None
def explanation(d):
    if exclusion_reason(d):return '默认不补：'+exclusion_reason(d)+'；计算结果仅供参考'
    if issue_text(d):return issue_text(d)
    if d.get('_budget_gap',0)>0:return ('必卖：' if d.get('_must_sell') else '普通：')+d.get('_budget_reason','')
    if d['_source']=='地采商品':return '商品资料未匹配，按地采单列；保留销量库存，本版不补造价格或数量'
    if d.get('_preview'):return '已匹配，等待计算'
    return d.get('BM','')

COLUMNS=[('商品名称',lambda d:d.get('D')),('商品编码',lambda d:d.get('E')),('来源',lambda d:d['_source']),('所属仓',lambda d:d.get('AU')),
 ('分类',lambda d:d.get('_category')),('月销',lambda d:d.get('J')),('可用库存',lambda d:d.get('K')),('在途',lambda d:d.get('L')),
 ('箱规',lambda d:d.get('BD',d.get('O'))),('执行补货',order_quantity),('执行金额',order_amount),('例外',lambda d:'一箱放行' if d.get('_audit',{}).get('applied') else ''),
 ('补货建议状态',lambda d:d['_meta'].get('export_status')),('原因',explanation),('商品资料状态',lambda d:d.get('S')),
 ('执行状态',order_status),('排除原因',exclusion_reason),('计算建议',quantity),('计算金额',amount_value),
 ('优先清单',lambda d:'必卖' if d.get('_must_sell') else '普通'),('预算前执行量',lambda d:d.get('_candidate_qty',quantity(d))),
 ('预算前金额',lambda d:d.get('_candidate_amount',amount_value(d))),('预算缺口数量',lambda d:d.get('_budget_gap',0)),
 ('最低持有缺口',lambda d:'有销量且无可用/在途，本次资金未覆盖' if d.get('_holding_gap') else ''),('资金分配说明',lambda d:d.get('_budget_reason','未启用总预算'))]

IMPORT_HEADERS=['类型','门店编码','商品编码','来源文件','工作表','源行 / 行数','保留选择 / 状态','SHA256']
def import_audit_rows(result):
    audit=result.get('import_audit',{})
    yield ['一致重复去重数量','','','','',audit.get('identical_duplicates',0),'','']
    for f in audit.get('files',[]):yield ['读取文件','','',f['path'],'',f['rows'],f['error'] or '成功',result.get('sources',{}).get(f['path'],'')]
    for c in audit.get('conflicts',[]):
        for i,source in enumerate(c['sources']):yield ['重复冲突',c['store'],c['sku'],source['file'],source['sheet'],source['row'],'保留' if c['choice']==i else '未处理' if c['choice'] is None else '未选','']

def export_xlsx(result,path,rows=None,scope='本次全部结果',cancel=None,progress=lambda s:None):
    """Application's native offline exporter. Text identifiers/formula-looking input stay text."""
    path=Path(path).resolve();rows=result['rows'] if rows is None else rows
    if str(path) in result['sources'] or path.exists():raise ValueError('保存位置已有文件，请换一个新文件名；不会覆盖源文件或旧结果。')
    if path.suffix.lower()!='.xlsx':raise ValueError('导出文件必须为.xlsx')
    if not rows:raise ValueError('当前筛选没有商品可导出。')
    book=openpyxl.Workbook(write_only=True)
    def sheet(name,headers,data,widths=None):
        ws=book.create_sheet(name);ws.freeze_panes='C2';ws.sheet_view.showGridLines=False
        for i,h in enumerate(headers,1):ws.column_dimensions[get_column_letter(i)].width=(widths or {}).get(i,20)
        cells=[]
        for h in headers:
            c=WriteOnlyCell(ws,h);c.font=Font(name='微软雅黑',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='183F45');c.alignment=Alignment(wrap_text=True);cells.append(c)
        ws.append(cells);count=1
        for row in data:
            checkpoint(cancel);cells=[]
            for v in row:
                c=WriteOnlyCell(ws,v)
                if isinstance(v,str):c.data_type='s'
                if numeric(v):c.number_format='0.##########'
                c.font=Font(name='微软雅黑',size=10);c.alignment=Alignment(vertical='top',wrap_text=True);cells.append(c)
            ws.append(cells);count+=1
        ws.auto_filter.ref=f'A1:{get_column_letter(len(headers))}{count}'
        return ws
    summary=summarize(rows)
    info=[['软件版本',VERSION],['模型版本',MODEL_VERSION],['构建信息',json.dumps(diagnostic(),ensure_ascii=False)],['计算时间',result['created']],['门店',result['store_name']],['导出范围',scope],['一箱例外',result['enabled']],
        ['计价口径','商品资料基本售价；地采缺价不计为0'],['规则口径','原库存额度BH/BN日均；一箱例外仅后置修改最终建议'],
        ['商品数',summary['rows']],['默认补货商品',summary['positive']],['默认补货金额',summary['amount']],['地采商品',summary['local']],['数据异常商品',summary['issues']],
        ['排除规则','商品资料状态=禁用 或 补货建议状态=停采，任一命中即排除；计算保留，执行量和执行金额为0'],
        ['已排除商品',summary['excluded']],['排除前计算金额',summary['calculated_amount']],['排除部分计算金额',summary['excluded_amount']],
        ['一箱例外商品',summary['exception']],['例外新增金额',summary['extra']],['未计价补货商品',summary['unpriced']]]
    b=result.get('budget',{});ml=result.get('must_list')
    info.extend([['预算前候选金额（本导出范围）',summary['candidate_amount']],['必卖执行金额（本导出范围）',summary['must_amount']],['预算未满足商品（本导出范围）',summary['budget_deferred']],['最低持有缺口（本导出范围）',summary['holding_gaps']]])
    if b.get('enabled'):
        info.extend([['全店总预算',b['limit']],['全店已分配',b['spent']],['全店剩余预算',b['remaining']],['全店必卖分配',b['must_spent']],['全店普通分配',b['other_spent']],['全店价格或箱规待核',b['unpriced']],['分配顺序',b['policy']],['预算口径','本次新增补货金额，不含现有库存/在途价值；一箱例外也受门店总预算约束']])
    else:info.append(['门店总预算','未启用，执行口径与0.2一致'])
    if ml:info.extend([['必卖清单',ml['name']],['清单版本',ml['revision']],['清单来源SHA256',ml['source_sha256']],['清单总商品数',ml['total']],['本店命中清单数',ml['matched']],['清单未出现在本店导出',ml['missing']]])
    for p,h in result['sources'].items():info.extend([['输入文件',p],['SHA256',h]])
    sheet('运行概览',['项目','内容'],info,{1:30,2:100})
    if result.get('import_audit'):sheet('导入记录',IMPORT_HEADERS,import_audit_rows(result),{4:70})
    headers=[h for h,f in COLUMNS]+['周销','14天日均','缺货天数','基本售价','原模型建议','原模型金额','例外新增金额','预测日均','BH额度日均','BN增长额度日均','库存额度总量','按计算建议持有天数','备注','导出源行','资料源行']
    def values(d):
        r=d['AN']*d['AM'] if numeric(d.get('AN')) and numeric(d.get('AM')) else None
        return [f(d) for h,f in COLUMNS]+[d.get('I'),d.get('H'),d.get('N'),d.get('V'),d.get('_base_qty'),d.get('_base_amount'),d.get('_audit',{}).get('incremental_amount',0),r,d.get('BH'),d.get('BN'),d.get('库存额度目标量'),coverage(d),'；'.join(d['_meta']['notes']),d['_meta']['export_row'],d['_meta'].get('product_row')]
    progress('正在写入商品明细…')
    for name,subset in [('补货清单',[d for d in rows if orderable(d)]),('全部明细' if scope=='本次全部结果' else '筛选明细',rows),('预算未满足',[d for d in rows if d.get('_budget_gap',0)>0]),('已排除',[d for d in rows if exclusion_reason(d)]),('地采商品',[d for d in rows if d['_source']=='地采商品']),('待处理',[d for d in rows if issue_text(d)]),('一箱例外',[d for d in rows if d.get('_audit',{}).get('applied')])]:
        sheet(name,headers,(values(d) for d in subset),{1:55,2:26,14:85,27:50})
    if ml:sheet('全店必卖匹配',['商品编码','清单商品名称','本店匹配状态'],([x['sku'],x['name'],x['state']] for x in result['must_list_audit']),{1:26,2:65,3:42})
    params=[]
    for wh,(depth,lead) in result['warehouses'].items():params.extend([[wh,'备货天数',depth],[wh,'到货天数',lead]])
    for k,v in result['config']['parameters'].items():params.append(['模型参数',result['config']['parameter_labels'].get(k,k),v])
    for name,p in result['config']['weight_profiles'].items():
        for k,v in p.items():params.append(['分类权重',name+' / '+k,v])
    sheet('参数',['范围','参数','值'],params,{1:24,2:48,3:28})
    tmp=path.parent/('.'+path.stem+'.'+uuid.uuid4().hex+'.xlsx')
    try:
        checkpoint(cancel);progress('正在保存文件…');book.save(tmp);checkpoint(cancel)
        os.rename(tmp,path)
    finally:
        book.close()
        if tmp.exists():tmp.unlink()
    return str(path)
