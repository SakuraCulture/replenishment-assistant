"""Native single-store and multi-store Excel export, sharing writing primitives."""
import os,re,json,uuid
from contextlib import contextmanager
from pathlib import Path
from datetime import datetime
import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from result_semantics import *
from display_fields import COLUMNS,IMPORT_HEADERS,import_audit_rows,single_headers,single_values,batch_headers
from runtime_support import checkpoint
import export_progress as ep

# Immutable style objects shared by cells; never mutate after assignment.
_SINGLE_FONT=Font(name="微软雅黑",size=10)
_SINGLE_ALIGN=Alignment(vertical="top",wrap_text=True)
_HEADER_FONT={True:Font(name="微软雅黑",bold=True,color="FFFFFF"),False:Font(bold=True,color="FFFFFF")}
_HEADER_FILL={True:PatternFill("solid",fgColor="183F45"),False:PatternFill("solid",fgColor="176C63")}
_HEADER_ALIGN=Alignment(wrap_text=True)
from version_info import VERSION,MODEL_VERSION,diagnostic

@contextmanager
def stream_book():
    book=openpyxl.Workbook(write_only=True)
    try:yield book
    finally:
        # Closing a cancelled write-only sheet finalizes its open XML stream;
        # remove only this workbook's temporary XML, never user output files.
        for ws in book.worksheets:
            try:
                if not ws.closed:ws.close()
            finally:
                writer=getattr(ws,'_writer',None)
                if writer and isinstance(writer.out,str) and os.path.exists(writer.out):writer.cleanup()
        book.close()


def write_table(book,name,headers,rows,*,widths=None,style='batch',cancel=None,progress=lambda s:None,row_count=None):
    """Same safe text cells, with explicit historical single/batch formatting policies."""
    ws=book.create_sheet(name);ws.freeze_panes='C2';single=style=='single'
    if single:ws.sheet_view.showGridLines=False
    for i in range(len(headers)):
        ws.column_dimensions[get_column_letter(i+1)].width=(widths or {}).get(i+1,20) if single else 24 if i else 22
    count=0
    if row_count is None and hasattr(rows,'__len__'):row_count=len(rows)
    checkpoint(cancel)
    if row_count is not None:ep.rows(progress,name,0,row_count)
    else:progress('写入：'+name)
    for values in _with_header(headers,() if rows is None else rows):
        checkpoint(cancel)
        cells=[]
        for value in values:
            c=WriteOnlyCell(ws,value)
            if isinstance(value,str):c.data_type='s'
            if count==0:
                c.font=_HEADER_FONT[single]
                c.fill=_HEADER_FILL[single]
                if single:c.alignment=_HEADER_ALIGN
            else:
                if numeric(value):c.number_format='0.##########'
                if single:c.font=_SINGLE_FONT;c.alignment=_SINGLE_ALIGN
            cells.append(c)
        ws.append(cells);count+=1
        done=count-1
        if done and done%500==0:
            if row_count is not None:ep.rows(progress,name,done,row_count)
            else:progress(f"{name}：已写 {done:,} 行")
    ep.rows(progress,name,max(0,count-1),row_count if row_count is not None else max(0,count-1))
    if single:ws.auto_filter.ref=f'A1:{get_column_letter(len(headers))}{count}'
    return ws

def write_sheet(book,name,headers,rows,**kwargs):
    """Compatibility helper retaining batch formatting, also used by old snapshot adapters."""
    return write_table(book,name,headers,rows,**kwargs)

def _with_header(headers,rows):
    yield headers
    yield from rows


def export_xlsx(result,path,rows=None,scope='本次全部结果',cancel=None,progress=lambda s:None):
    """Application's native offline exporter. Text identifiers/formula-looking input stay text."""
    path=Path(path).resolve();raw_rows=result['rows'] if rows is None else rows;rows=working_rows(raw_rows)
    result=export_projection(result)
    if str(path) in result['sources'] or path.exists():raise ValueError('保存位置已有文件，请换一个新文件名；不会覆盖源文件或旧结果。')
    if path.suffix.lower()!='.xlsx':raise ValueError('导出文件必须为.xlsx')
    if not raw_rows:raise ValueError('当前筛选没有商品可导出。')
    with stream_book() as book:
        def sheet(name,headers,data,widths=None,row_count=None):
            return write_table(book,name,headers,data,widths=widths,style='single',cancel=cancel,progress=progress,row_count=row_count)
        summary=summarize(rows)
        info=[['软件版本',VERSION],['模型版本',MODEL_VERSION],['构建信息',json.dumps(diagnostic(),ensure_ascii=False)],['计算时间',result['created']],['门店',result['store_name']],['导出范围',scope],['一箱例外',result['enabled']],
            ['计价口径','商品资料基本售价；地采缺价不计为0'],['规则口径','原库存额度BH/BN日均；一箱例外仅后置修改最终建议'],
            ['商品数',summary['rows']],['默认补货商品',summary['positive']],['默认补货金额',summary['amount']],['地采商品',summary['local']],['数据异常商品',summary['issues']],
            ['排除规则','商品资料状态=禁用 或 补货建议状态=停采，任一命中即排除；计算保留，执行量和执行金额为0'],
            ['已排除商品',summary['excluded']],['排除前计算金额',summary['calculated_amount']],['排除部分计算金额',summary['excluded_amount']],
            ['一箱例外商品',summary['exception']],['例外新增金额',summary['extra']],['未计价补货商品',summary['unpriced']]]
        b=result.get('budget',{});ml=result.get('must_list')
        info.extend([['预算前候选金额（本导出范围）',summary['candidate_amount']],['必卖执行金额（本导出范围）',summary['must_amount']],['预算未满足商品（本导出范围）',summary['budget_deferred']],['最低持有缺口（本导出范围）',summary['holding_gaps']]])
        if b.get('enabled'):
            info.extend([['全店总预算',b['limit']],['全店已分配',b['spent']],['全店剩余预算',b['remaining']],['全店必卖分配',b['must_spent']],['全店普通分配',b['other_spent']],['全店价格或箱规待核',b['unpriced']],['分配顺序',b['policy']],['预算口径','本次新增补货金额，不含现有库存/在途价值；一箱例外也受门店总预算约束']])
        else:info.append(['门店总预算','未启用，执行口径与0.2一致'])
        if ml:info.extend([['必卖清单',ml['name']],['清单版本',ml['revision']],['清单来源SHA256',ml['source_sha256']],['清单总商品数',ml['total']],['本店命中清单数',ml['matched']],['清单未出现在本店导出',ml['missing']]])
        for p,h in result['sources'].items():info.extend([['输入文件',p],['SHA256',h]])
        if result.get('effective_settings'):info.append(['本店实际生效完整参数',json.dumps(result['effective_settings'],ensure_ascii=False)])
        if result.get('catalogs'):
            info.extend([['主从仓配置',json.dumps(result['supply_route'],ensure_ascii=False)],['仓资料快照',json.dumps(result['catalogs'],ensure_ascii=False)],['选仓口径','主仓禁用/可用数<=0/无该商品时转从仓；从仓无货仍保留需求归属，不拆单、不限量、不扣减仓库存；未知可用数不当0']])
        sheet('运行概览',['项目','内容'],info,{1:30,2:100})
        if result.get('import_audit'):sheet('导入记录',IMPORT_HEADERS,import_audit_rows(result),{4:70})
        headers=single_headers()
        progress('正在写入商品明细…')
        for name,subset in [('补货清单',[d for d in rows if orderable(d)]),('全部明细' if scope=='本次全部结果' else '筛选明细',rows),('预算未满足',[d for d in rows if d.get('_budget_gap',0)>0]),('已排除',[d for d in rows if exclusion_reason(d)]),('地采商品',[d for d in rows if d['_source']=='地采商品']),('待处理',[d for d in rows if issue_text(d)]),('一箱例外',[d for d in rows if d.get('_audit',{}).get('applied')])]:
            sheet(name,headers,(single_values(d) for d in subset),{1:55,2:26,14:85,27:50},row_count=len(subset))
        if ml:sheet('全店必卖匹配',['商品编码','清单商品名称','本店匹配状态'],([x['sku'],x['name'],x['state']] for x in result['must_list_audit']),{1:26,2:65,3:42})
        params=[]
        for wh,(depth,lead) in result['warehouses'].items():params.extend([[wh,'备货天数',depth],[wh,'到货天数',lead]])
        for k,v in result['config']['parameters'].items():params.append(['模型参数',result['config']['parameter_labels'].get(k,k),v])
        for name,p in result['config']['weight_profiles'].items():
            for k,v in p.items():params.append(['分类权重',name+' / '+k,v])
        sheet('参数',['范围','参数','值'],params,{1:24,2:48,3:28})
        tmp=path.parent/('.'+path.stem+'.'+uuid.uuid4().hex+'.xlsx')
        try:
            checkpoint(cancel);progress('数据行写入完成，正在压缩保存Excel；此阶段暂不能中断…');book.save(tmp);checkpoint(cancel)
            os.rename(tmp,path)
        finally:
            book.close()
            if tmp.exists():tmp.unlink()
        return str(path)



def export_batch(results,statuses,path,per_store=False,progress=lambda s:None,cancel=None):
    """Only current successful results count as spend; missing results are blank, never zero."""
    path=Path(path).resolve()
    if path.exists():raise ValueError('目标已存在，请选择新名称。')
    ready={s:r for s,r in results.items() if statuses.get(s,{}).get('state')=='已完成'}
    if not ready:raise ValueError('没有可导出的有效结果，请先计算。')
    if per_store:
        path.mkdir();written=[]
        try:
            for index,(s,r) in enumerate(ready.items(),1):
                prefix=f'门店 {index}/{len(ready)} · '+r['store_name'];checkpoint(cancel);progress(prefix)
                safe=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',s)[:80]
                file=path/(safe+'_'+uuid.uuid4().hex[:6]+'.xlsx')
                export_xlsx(r,file,cancel=cancel,progress=ep.prefixed(progress,prefix));written.append(str(file))
            export_batch(ready,statuses,path/'门店汇总及明细.xlsx',cancel=cancel,progress=ep.prefixed(progress,'全部门店文件完成 · 汇总及明细'))
        except Exception:
            (path/'部分导出说明.txt').write_text('导出未全部完成。已生成文件：\n'+'\n'.join(written),encoding='utf-8');raise
        return str(path)
    if path.suffix.lower()!='.xlsx':raise ValueError('汇总文件必须为.xlsx')
    excluded_keys={(s,d['E']) for s,r in ready.items() for d in r['rows'] if not work_eligible(d)}
    ready={s:export_projection(r,excluded_keys) for s,r in ready.items()}
    with stream_book() as book:
        def summary_rows():
            for store,state in statuses.items():
                r=ready.get(store);s=r['summary'] if r else {};b=r.get('budget',{}) if r else {}
                yield [store,state['name'],state['state'],state.get('error',''),b.get('limit') if b.get('enabled') else '不限预算' if r else None,s.get('candidate_amount'),s.get('amount'),s.get('must_amount'),round(s['amount']-s['must_amount'],2) if r else None,round(s['candidate_amount']-s['amount'],2) if r else None,b.get('remaining'),s.get('holding_gaps'),s.get('local'),s.get('excluded'),s.get('unpriced')]
        write_sheet(book,'门店汇总',['门店编码','门店名称','状态','提示','本次预算','预算前需求金额','本次安排金额','必卖金额','其他金额','未满足需求金额','剩余预算','最低持有缺口款数','地采款数','已排除款数','补货未计价款数'],summary_rows())
        if any(r.get('supply_route') for r in ready.values()):
            def warehouse_summary():
                for store,r in ready.items():
                    for wh in sorted({d.get('AU','') for d in r['rows']}):
                        s=summarize([d for d in r['rows'] if d.get('AU')==wh])
                        yield [store,r['store_name'],wh,s['rows'],s['calculated_amount'],s['candidate_amount'],s['amount'],s['positive']]
            write_sheet(book,'分仓需求汇总',['门店编码','门店名称','需求归属仓','商品数','计算金额','预算前可补金额','执行金额','执行商品数'],warehouse_summary())
        headers=batch_headers()
        filters={'补货清单':orderable,'全部明细':lambda d:True,'预算未满足':lambda d:d.get('_budget_gap',0)>0,'已排除':exclusion_reason,'地采待核':lambda d:d['_source']=='地采商品' or issue_text(d)}
        for name,predicate in filters.items():
            progress('写入：'+name)
            def rows():
                for store,r in ready.items():
                    for d in r['rows']:
                        checkpoint(cancel)
                        if predicate(d):yield [store,r['store_name']]+[f(d) for h,f in COLUMNS]+[d['_meta'].get('export_file'),d['_meta'].get('export_sheet'),d['_meta']['export_row']]
            write_sheet(book,name,headers,rows(),cancel=cancel,progress=progress,row_count=sum(1 for r in ready.values() for d in r['rows'] if predicate(d)))
        write_sheet(book,'必卖匹配',['门店编码','门店名称','商品编码','商品名称','匹配状态'],([s,r['store_name'],a['sku'],a['name'],a['state']] for s,r in ready.items() for a in r['must_list_audit']))
        audits={json.dumps(r.get('import_audit',{}),ensure_ascii=False,sort_keys=True):r for r in ready.values()}
        write_sheet(book,'导入记录',IMPORT_HEADERS,(row for r in audits.values() for row in import_audit_rows(r)))
        info=[['批次状态','全部完成' if len(ready)==len(statuses) else '部分完成'],['导出时间',datetime.now().isoformat(timespec='seconds')],['说明','只有已完成且未过期的门店计入金额；失败、未运行、待重算均留空。'],['计价','商品资料基本售价；地采和缺价未知金额不计为0。']]
        for s,r in ready.items():
            info.extend([[s+' / '+k,str(v)] for k,v in r.get('version',diagnostic()).items()])
            if r.get('effective_settings'):info.append([s+' / 方案与实际生效配置',json.dumps(r['effective_settings'],ensure_ascii=False)])
            info.extend([[s+' / 参数',json.dumps(r['config'],ensure_ascii=False)],[s+' / 仓参数',json.dumps(r['warehouses'],ensure_ascii=False)],[s+' / 一箱例外',r['enabled']],[s+' / 必卖清单',json.dumps(r['must_list'],ensure_ascii=False)],[s+' / 输入校验',json.dumps(r['sources'],ensure_ascii=False)]])
            if r.get('catalogs'):info.extend([[s+' / 主从仓',json.dumps(r['supply_route'],ensure_ascii=False)],[s+' / 仓资料快照',json.dumps(r['catalogs'],ensure_ascii=False)],[s+' / 选仓口径','可用数>0优先主仓；主仓禁用/无货/无商品则归属从仓，从仓无货不清零需求；不拆单、不限量、不扣减仓库存']])
        write_sheet(book,'版本及参数',['项目','内容'],info)
        tmp=path.with_name('.'+uuid.uuid4().hex+'.xlsx')
        try:
            checkpoint(cancel);progress('数据行写入完成，正在压缩保存汇总Excel；此阶段暂不能中断…');book.save(tmp);checkpoint(cancel);os.rename(tmp,path)
        finally:
            book.close()
            if tmp.exists():tmp.unlink()
        return str(path)
