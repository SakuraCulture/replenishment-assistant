"""Batch orchestration around the unchanged single-store business engine."""
from __future__ import annotations
import copy,hashlib,json,multiprocessing,os,re,tempfile,time,uuid
from collections import defaultdict
from concurrent.futures import ProcessPoolExecutor,wait,FIRST_COMPLETED
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font,PatternFill
import engine
import warehouse_catalog
from version_info import diagnostic

def intake(paths,products,progress=lambda s:None,cancel=None):
    paths=list(dict.fromkeys(str(Path(p).resolve()) for p in paths))
    catalogs=copy.deepcopy(products) if isinstance(products,dict) else None
    if catalogs:
        for wh,catalog in catalogs.items():
            warehouse_catalog.validate(catalog)
            if wh!=catalog['warehouse']:raise ValueError('仓资料标签不一致')
    elif catalogs is not None:raise ValueError('请先导入河北仓或山东仓资料')
    products=str(Path(products).resolve()) if catalogs is None else None
    if not paths:raise ValueError('请至少添加一个补货建议文件。')
    if products in paths:raise ValueError('商品资料不能同时作为补货建议。')
    product_hash=engine.adapter.digest(Path(products)) if products else None
    product_snapshot=None
    if products:
        product_sheet,_=engine.sheet_name(products,'products')
        product_snapshot,_=engine.adapter.read_table(products,product_sheet,['商品编码','基本售价','分类','出库箱规'])
    groups=defaultdict(list);files=[];stores={};headers=[];hashes={products:product_hash} if products else {}
    for path in paths:
        engine.checkpoint(cancel);progress('读取补货建议：'+Path(path).name)
        report={'path':path,'rows':0,'stores':0,'error':''};files.append(report)
        try:
            before=engine.adapter.digest(Path(path));sheet,heads=engine.sheet_name(path,'sales')
            cfg=engine.config();sf=cfg['sales_fields']
            rows,info=engine.adapter.read_table(Path(path),sheet,[sf[k] for k in ['store_code','sku','daily14','sales7','sales30','stock','inbound','stockout_days']])
            if before!=engine.adapter.digest(Path(path)):raise ValueError('读取期间文件变化，请保存后重试。')
            ambiguous=set(sf.values()) & set(info['duplicate_unused_headers'])
            if ambiguous:raise ValueError('已映射表头重复：'+ '、'.join(sorted(ambiguous)))
            staged=[];names={}
            for raw in rows:
                code,error=engine.adapter.identifier(raw.get(sf['store_code']))
                if error:raise ValueError(f"第{raw['_row']}行门店编码无效：{error}；请使用文本编码")
                sku,sku_error=engine.adapter.identifier(raw.get(sf['sku']))
                raw=copy.deepcopy(raw)
                raw['基础单位']=raw.get('基础单位',raw.get('库存单位',raw.get('单位类型')))
                values={k:v for k,v in raw.items() if k!='_row'}
                # Invalid SKU rows remain distinct and are diagnosed by the original adapter.
                key=(code,sku if not sku_error else f"待核:{path}:{raw['_row']}")
                staged.append((key,{'values':values,'path':path,'sheet':sheet,'row':raw['_row']}))
                names.setdefault(code,str(raw.get(sf['store_name']) or code))
            if not staged:raise ValueError('文件没有商品记录。')
            for key,record in staged:groups[key].append(record)
            stores.update({k:v for k,v in names.items() if k not in stores})
            for h in sorted(heads|{'基础单位'}):
                if h not in headers:headers.append(h)
            hashes[path]=before;report.update(rows=len(staged),stores=len(names))
        except engine.Cancelled:raise
        except Exception as exc:report['error']=str(exc)
    if not groups:raise ValueError('没有可读取的补货建议：\n'+'\n'.join(Path(f['path']).name+'：'+f['error'] for f in files))
    duplicates=0;conflicts=[];selected=[]
    for (store,sku),records in groups.items():
        variants=[];seen=set()
        for rec in records:
            fingerprint=json.dumps({k:v for k,v in rec['values'].items() if v is not None},sort_keys=True,ensure_ascii=False,default=str)
            if fingerprint in seen:duplicates+=1;continue
            seen.add(fingerprint);variants.append(rec)
        if len(variants)==1:selected.append(variants[0])
        else:conflicts.append({'id':len(conflicts),'store':store,'sku':sku,'variants':variants,'choice':None})
    if products and engine.adapter.digest(Path(products))!=product_hash:raise ValueError('商品资料读取期间发生变化，请重试。')
    return {'files':files,'stores':stores,'headers':headers,'hashes':hashes,'products':products,'products_snapshot':product_snapshot,'catalogs':catalogs,'selected':selected,'conflicts':conflicts,'duplicates':duplicates}

def prepare(data,progress=lambda s:None,cancel=None):
    """Materialize deduplicated source rows privately, then use the original adapter."""
    for p,h in data['hashes'].items():
        if not data.get('frozen') and engine.adapter.digest(Path(p))!=h:raise ValueError('源文件已变化，请重新检查：'+Path(p).name)
    records=list(data['selected']);blocked=defaultdict(int)
    for c in data['conflicts']:
        if c['choice'] is None:blocked[c['store']]+=1
        else:records.append(c['variants'][c['choice']])
    if not records:raise ValueError('所有记录均有冲突，请先选择保留记录。')
    with tempfile.TemporaryDirectory(prefix='replenishment-') as tmp:
        path=Path(tmp)/'batch.xlsx';book=openpyxl.Workbook(write_only=True);ws=book.create_sheet('补货建议');ws.append(data['headers'])
        for r in records:
            engine.checkpoint(cancel);cells=[]
            for h in data['headers']:
                value=r['values'].get(h);cell=WriteOnlyCell(ws,value)
                if isinstance(value,str):cell.data_type='s'
                cells.append(cell)
            ws.append(cells)
        book.save(path);book.close()
        products=data['products']
        if data.get('frozen') and not data.get('catalogs'):
            raw=data.get('products_snapshot')
            if not raw:raise ValueError('旧作业没有商品资料快照，请重新选择原资料')
            products=Path(tmp)/'master.xlsx';book=openpyxl.Workbook(write_only=True);heads=[k for k in raw[0] if k!='_row'];write_sheet(book,'商品资料',heads,([r.get(k) for k in heads] for r in raw));book.save(products);book.close()
        loaded=warehouse_catalog.prepare_warehouses(path,data['catalogs'],tmp,progress,cancel) if data.get('catalogs') else engine.load_inputs(path,products,progress,cancel)
        if data.get('frozen') and not data.get('catalogs'):
            for d,m in loaded['normalized']:
                if m.get('product_row'):m['product_row']=data['products_snapshot'][m['product_row']-2]['_row']
    for pairs in [loaded['normalized'],*loaded.get('warehouse_rows',{}).values()]:
        for d,m in pairs:
            r=records[m['export_row']-2]
            m.update(export_row=r['row'],export_file=r['path'],export_sheet=r['sheet'])
    loaded.update(hashes=dict(data['hashes']),sales=[f['path'] for f in data['files'] if not f['error']],stores=dict(data['stores']),blocked=dict(blocked),intake=data)
    for wh,c in loaded.get('catalogs',{}).items():loaded['hashes'][wh+'资料快照 '+c['source_path']+' #'+c['revision']]=c['source_sha256']
    for p,h in data['hashes'].items():
        if not data.get('frozen') and engine.adapter.digest(Path(p))!=h:raise ValueError('读取期间源文件变化，请重新检查：'+Path(p).name)
    return loaded

def load(paths,products,progress=lambda s:None,cancel=None):
    data=intake(paths,products,progress,cancel)
    # UI must be able to resolve a dataset consisting entirely of conflicts.
    if not data['selected'] and data['conflicts']:return {'intake':data,'stores':data['stores'],'normalized':[],'blocked':{s:1 for s in data['stores']},'warehouses':[],'config':engine.config(),'hashes':data['hashes'],'catalogs':{w:warehouse_catalog.metadata(c) for w,c in (data.get('catalogs') or {}).items()}}
    return prepare(data,progress,cancel)

def calculate_store(loaded,store,settings,profile):
    started=time.time()
    if loaded.get('catalogs'):loaded=warehouse_catalog.route_store(loaded,store,settings.get('supply_route'))
    if settings.get('parameter_config'):
        import parameter_schema
        loaded=parameter_schema.apply(loaded,settings['parameter_config'])
    result=engine.run_calculation(loaded,store,settings['warehouses'],settings['exception'],settings['budget'],profile)
    if loaded.get('catalogs'):result.update(catalogs=loaded['catalogs'],supply_route=loaded['supply_route'])
    result['version']=diagnostic();result['task']={'pid':os.getpid(),'started':started,'finished':time.time()}
    result['effective_settings']=copy.deepcopy(settings)
    result['import_audit']=loaded.get('import_audit',{})
    return result

def run(loaded,jobs,profile,max_workers=2,notify=lambda s,state,result,error:None,cancel=None):
    """Submit only available slots: stop never starts queued stores; running stores finish."""
    if not 1<=max_workers<=4:raise ValueError('并发门店数必须为1至4。')
    snapshot=copy.deepcopy(loaded);data=snapshot.pop('intake',None)
    if data:snapshot['import_audit']={'files':data['files'],'identical_duplicates':data['duplicates'],'conflicts':[{'store':c['store'],'sku':c['sku'],'choice':c['choice'],'sources':[{'file':v['path'],'sheet':v['sheet'],'row':v['row']} for v in c['variants']]} for c in data['conflicts']]}
    pending=iter(jobs.items());active={};out={};stopped=False
    with ProcessPoolExecutor(max_workers=max_workers,mp_context=multiprocessing.get_context('spawn')) as pool:
        def fill():
            nonlocal stopped
            stopped=bool(cancel and cancel.is_set())
            if stopped:return
            while len(active)<max_workers:
                if cancel and cancel.is_set():return
                try:store,settings=next(pending)
                except StopIteration:break
                if snapshot.get('blocked',{}).get(store):
                    notify(store,'计算失败',None,'存在重复冲突，请先处理');out[store]=None;continue
                one={**snapshot,'normalized':[x for x in snapshot['normalized'] if x[0]['B']==store]}
                if snapshot.get('warehouse_rows'):one['warehouse_rows']={w:[x for x in rows if x[0]['B']==store] for w,rows in snapshot['warehouse_rows'].items()}
                notify(store,'计算中',None,'')
                active[pool.submit(calculate_store,one,store,copy.deepcopy(settings),copy.deepcopy(profile))]=store
        fill()
        while active:
            done,_=wait(active,timeout=.2,return_when=FIRST_COMPLETED)
            for future in done:
                store=active.pop(future)
                try:
                    result=future.result();out[store]=result;notify(store,'已完成',result,'')
                except Exception as exc:out[store]=None;notify(store,'计算失败',None,str(exc))
            fill()
        for store,_ in pending:notify(store,'已停止',None,'未启动，可重新计算')
    return out

def write_sheet(book,name,headers,rows):
    ws=book.create_sheet(name);ws.freeze_panes='C2'
    from openpyxl.utils import get_column_letter
    for i in range(len(headers)):ws.column_dimensions[get_column_letter(i+1)].width=24 if i else 22
    for index,values in enumerate([headers] if rows is None else _with_header(headers,rows)):
        cells=[]
        for value in values:
            c=WriteOnlyCell(ws,value)
            if isinstance(value,str):c.data_type='s'
            if index==0:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='176C63')
            elif engine.numeric(value):c.number_format='0.##########'
            cells.append(c)
        ws.append(cells)
    return ws

def _with_header(headers,rows):
    yield headers
    yield from rows

def export_batch(results,statuses,path,per_store=False,progress=lambda s:None,cancel=None):
    """Only current successful results count as spend; missing results are blank, never zero."""
    path=Path(path).resolve()
    if path.exists():raise ValueError('目标已存在，请选择新名称。')
    ready={s:r for s,r in results.items() if statuses.get(s,{}).get('state')=='已完成'}
    if not ready:raise ValueError('没有可导出的有效结果，请先计算。')
    if per_store:
        path.mkdir();written=[]
        try:
            for s,r in ready.items():
                engine.checkpoint(cancel);progress('导出门店：'+r['store_name'])
                safe=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',s)[:80]
                file=path/(safe+'_'+uuid.uuid4().hex[:6]+'.xlsx')
                engine.export_xlsx(r,file,cancel=cancel);written.append(str(file))
            export_batch(ready,statuses,path/'门店汇总及明细.xlsx',cancel=cancel)
        except Exception:
            (path/'部分导出说明.txt').write_text('导出未全部完成。已生成文件：\n'+'\n'.join(written),encoding='utf-8');raise
        return str(path)
    if path.suffix.lower()!='.xlsx':raise ValueError('汇总文件必须为.xlsx')
    book=openpyxl.Workbook(write_only=True)
    def summary_rows():
        for store,state in statuses.items():
            r=ready.get(store);s=r['summary'] if r else {};b=r.get('budget',{}) if r else {}
            yield [store,state['name'],state['state'],state.get('error',''),b.get('limit') if b.get('enabled') else '不限预算' if r else None,s.get('candidate_amount'),s.get('amount'),s.get('must_amount'),round(s['amount']-s['must_amount'],2) if r else None,round(s['candidate_amount']-s['amount'],2) if r else None,b.get('remaining'),s.get('holding_gaps'),s.get('local'),s.get('excluded'),s.get('unpriced')]
    write_sheet(book,'门店汇总',['门店编码','门店名称','状态','提示','本次预算','预算前需求金额','本次安排金额','必卖金额','其他金额','未满足需求金额','剩余预算','最低持有缺口款数','地采款数','已排除款数','补货未计价款数'],summary_rows())
    if any(r.get('supply_route') for r in ready.values()):
        def warehouse_summary():
            for store,r in ready.items():
                for wh in sorted({d.get('AU','') for d in r['rows']}):
                    s=engine.summarize([d for d in r['rows'] if d.get('AU')==wh])
                    yield [store,r['store_name'],wh,s['rows'],s['calculated_amount'],s['candidate_amount'],s['amount'],s['positive']]
        write_sheet(book,'分仓需求汇总',['门店编码','门店名称','需求归属仓','商品数','计算金额','预算前可补金额','执行金额','执行商品数'],warehouse_summary())
    headers=['门店编码','门店名称']+[h for h,f in engine.COLUMNS]+['来源文件','来源工作表','来源行']
    filters={'补货清单':engine.orderable,'全部明细':lambda d:True,'预算未满足':lambda d:d.get('_budget_gap',0)>0,'已排除':engine.exclusion_reason,'地采待核':lambda d:d['_source']=='地采商品' or engine.issue_text(d)}
    for name,predicate in filters.items():
        progress('写入：'+name)
        def rows():
            for store,r in ready.items():
                for d in r['rows']:
                    engine.checkpoint(cancel)
                    if predicate(d):yield [store,r['store_name']]+[f(d) for h,f in engine.COLUMNS]+[d['_meta'].get('export_file'),d['_meta'].get('export_sheet'),d['_meta']['export_row']]
        write_sheet(book,name,headers,rows())
    write_sheet(book,'必卖匹配',['门店编码','门店名称','商品编码','商品名称','匹配状态'],([s,r['store_name'],a['sku'],a['name'],a['state']] for s,r in ready.items() for a in r['must_list_audit']))
    audits={json.dumps(r.get('import_audit',{}),ensure_ascii=False,sort_keys=True):r for r in ready.values()}
    write_sheet(book,'导入记录',engine.IMPORT_HEADERS,(row for r in audits.values() for row in engine.import_audit_rows(r)))
    info=[['批次状态','全部完成' if len(ready)==len(statuses) else '部分完成'],['导出时间',datetime.now().isoformat(timespec='seconds')],['说明','只有已完成且未过期的门店计入金额；失败、未运行、待重算均留空。'],['计价','商品资料基本售价；地采和缺价未知金额不计为0。']]
    for s,r in ready.items():
        info.extend([[s+' / '+k,str(v)] for k,v in r.get('version',diagnostic()).items()])
        if r.get('effective_settings'):info.append([s+' / 方案与实际生效配置',json.dumps(r['effective_settings'],ensure_ascii=False)])
        info.extend([[s+' / 参数',json.dumps(r['config'],ensure_ascii=False)],[s+' / 仓参数',json.dumps(r['warehouses'],ensure_ascii=False)],[s+' / 一箱例外',r['enabled']],[s+' / 必卖清单',json.dumps(r['must_list'],ensure_ascii=False)],[s+' / 输入校验',json.dumps(r['sources'],ensure_ascii=False)]])
        if r.get('catalogs'):info.extend([[s+' / 主从仓',json.dumps(r['supply_route'],ensure_ascii=False)],[s+' / 仓资料快照',json.dumps(r['catalogs'],ensure_ascii=False)],[s+' / 选仓口径','可用数>0优先主仓；主仓禁用/无货/无商品则归属从仓，从仓无货不清零需求；不拆单、不限量、不扣减仓库存']])
    write_sheet(book,'版本及参数',['项目','内容'],info)
    tmp=path.with_name('.'+uuid.uuid4().hex+'.xlsx')
    try:book.save(tmp);engine.checkpoint(cancel);os.rename(tmp,path)
    finally:
        book.close()
        if tmp.exists():tmp.unlink()
    return str(path)
