"""Returned merchant sheets reuse the signed layout, never the signed quantities."""
import copy,json,os,tempfile,uuid
from decimal import Decimal
from pathlib import Path
import openpyxl
from openpyxl.comments import Comment
from openpyxl.packaging.custom import StringProperty
import normal_exports as ne
from history_read import repository
from runtime_support import checkpoint,Cancelled
from export_progress import RowProgress

def returned_book(path,view,folder,progress,cancel):
    h=view['head'];doc=view['documents']['return'];order=repository(path).get(h['day'],h['store'],h['revision'],view='export')
    if order['digest']!=h['digest']:raise ValueError('签发修订校验失败')
    original=Path(folder)/('original-'+uuid.uuid4().hex+'.xlsx');ne.export(order,original,warehouse=h['warehouse'],progress=progress,cancel=cancel)
    b=openpyxl.load_workbook(original);rows={r['sku']:r for r in doc['payload']['rows']};seen=set()
    try:
        for ws in b:
            ws['Q1'].comment=Comment('补货数量为本版门店确认数；原签发建议独立保留在应用中。','补货助手')
            for n in range(2,ws.max_row+1):
                checkpoint(cancel);sku=str(ws.cell(n,5).value);r=rows.get(sku);seen.add(sku)
                ws.cell(n,17).value=float(Decimal(r['base_quantity'])) if r and r['base_quantity'] is not None else None
                ws.cell(n,17).comment=None
        for r in doc['payload']['rows']:
            if r['sku'] in seen:continue
            checkpoint(cancel);source_sheet=next((s['sheet'] for s in r['sources'] if s['sheet'] in ne.SHEETS),ne.SHEETS[-1]);ws=b[source_sheet]
            raw={k:v for k,v in zip(r.get('raw_headers',[]),r.get('raw_values',[]))};values=[raw.get(k) for k in ne.HEADERS+ne.AUDIT_HEADERS]
            values[0]=order['payload']['result']['store_name'];values[1]=h['store'];values[3]=r['name'];values[4]=r['sku'];values[16]=float(Decimal(r['base_quantity']));values[18]=h['warehouse'];values[20]=h['number'];values[21]='门店追加'
            ws.append(values)
            for i,c in enumerate(ws[ws.max_row],1):
                c._style=copy.copy(ws.cell(2 if ws.max_row>2 else 1,i)._style)
                if isinstance(c.value,str):c.data_type='s'
            ws.cell(ws.max_row,16).number_format='0.00';ws.cell(ws.max_row,18).number_format='@'
        b.custom_doc_props.append(StringProperty(name='门店回传版本',value=str(doc['version'])));b.properties.subject='门店回传单，数量为门店确认数'
        return b
    except Exception:b.close();raise

def export_returns(path,views,folder,mode,progress,cancel):
    completed=[]
    groups=[[v] for v in views] if mode=='separate' else [[v for v in views if v['head']['warehouse']==w] for w in dict.fromkeys(v['head']['warehouse'] for v in views)]
    try:
        for i,group in enumerate(groups):
            checkpoint(cancel);numbers={}
            for v in group:
                h=v['head']
                if h['number'] in numbers:raise ValueError('同一仓单多个修订不能合并为门店表，请分别导出，避免回传关联歧义')
                numbers[h['number']]=h['revision']
            with tempfile.TemporaryDirectory(prefix='.returned-',dir=folder) as temp:
                book=None
                try:
                    for v in group:
                        part=returned_book(path,v,temp,progress,cancel)
                        if book is None:book=part;continue
                        try:
                            for source in part:
                                dest=book[source.title]
                                for source_row in source.iter_rows(min_row=2):
                                    checkpoint(cancel);dest.append([c.value for c in source_row])
                                    for a,b in zip(source_row,dest[dest.max_row]):
                                        b._style=copy.copy(a._style);b.data_type=a.data_type;b.comment=copy.copy(a.comment)
                        finally:part.close()
                    if len(group)>1:
                        for name in ('补货单号','修订','门店正常单号','门店回传版本'):
                            if name in book.custom_doc_props.names:del book.custom_doc_props[name]
                        book.custom_doc_props.append(StringProperty(name='签发修订索引',value=json.dumps(numbers,ensure_ascii=False)))
                        for ws in book:ws.oddHeader.center.text=group[0]['head']['warehouse']+' 多订单门店回传单';ws['A1'].comment=Comment('每行补货单号独立关联；签发修订保存在文件属性。','补货助手')
                    name='门店回传单-'+group[0]['head']['warehouse']+'-'+(group[0]['head']['number'] if mode=='separate' else '多订单')+'-'+uuid.uuid4().hex[:8]+'.xlsx';target=Path(folder)/name
                    temp_file=Path(temp)/'returned.xlsx';progress('压缩保存门店回传单，完成后响应取消…');book.save(temp_file);checkpoint(cancel);os.link(temp_file,target);completed.append(str(target))
                finally:
                    if book:book.close()
            progress(RowProgress('已导出门店回传单',i+1,len(groups)))
        return {'files':completed,'cancelled':False,'error':None}
    except Cancelled:return {'files':completed,'cancelled':True,'error':None}
    except Exception as exc:return {'files':completed,'cancelled':False,'error':str(exc)}
