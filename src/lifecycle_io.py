"""Explicit whole-document adapters; inventory units, immutable Excel originals."""
import io
from datetime import date, datetime
from decimal import Decimal, localcontext
from pathlib import Path
import openpyxl
import purchase_io as base
from runtime_support import checkpoint
from export_progress import RowProgress

STAGES = ('return', 'payment', 'receipt')
TITLES = {'return': '门店确认单', 'payment': '门店付款单', 'receipt': '门店实际收货单'}
QUANTITIES = {'return': '门店确认数量', 'payment': '付款单商品数量', 'receipt': '实际收货数量'}
DOCNUM = {'return': '回传单号', 'payment': '付款单号', 'receipt': '收货单号'}
DOCDATE = {'return': '回传日期', 'payment': '付款日期', 'receipt': '收货日期'}
COMMON = ['补货单号', '签发修订', '门店编码', '供货仓', '商品编码', '商品名称', '数量单位', '折合库存单位系数', '换算依据']

def fields(stage):
    return COMMON + [QUANTITIES[stage], DOCNUM[stage], DOCDATE[stage]] + (['实际金额', '币种'] if stage == 'payment' else []) + ['备注']

def required(stage):
    return ['商品编码', QUANTITIES[stage], '数量单位'] + ([DOCNUM[stage], DOCDATE[stage], '实际金额', '币种'] if stage == 'payment' else [DOCNUM[stage], DOCDATE[stage]] if stage == 'receipt' else [])

def inspect(paths, progress=lambda _: None, cancel=None):
    infos = base.inspect(paths, progress, cancel)
    for f in infos:
        checkpoint(cancel)
        raw = Path(f['path']).read_bytes()
        if base.digest(raw) != f['sha']: raise ValueError('文件在读取期间变化')
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, keep_links=False)
        try: f['properties'] = {p.name: p.value for p in wb.custom_doc_props}
        finally: wb.close()
    return infos

def revision(value):
    if value in (None, ''): return None
    text = str(value).strip().lstrip('Vv')
    try: n = Decimal(text)
    except Exception: raise ValueError('签发修订须为正整数或V001') from None
    if not n.is_finite() or n <= 0 or n != n.to_integral_value(): raise ValueError('签发修订须为正整数或V001')
    return int(n)

def scan(stage, specs, progress=lambda _: None, cancel=None):
    names = fields(stage); rows = []; files = []
    for i, spec in enumerate(specs):
        checkpoint(cancel); p = Path(spec['path']); raw = p.read_bytes()
        if base.digest(raw) != spec['sha']: raise ValueError('映射后文件变化，请重新读取：' + p.name)
        files.append({k: spec[k] for k in ('path', 'name', 'sha')})
        mapping = spec['mapping']; defaults = spec.get('defaults', {})
        columns = [v for v in mapping.values() if v is not None]
        if len(columns) != len(set(columns)): raise ValueError('一个源列不能同时映射多个字段')
        for key in required(stage):
            if mapping.get(key) is None and defaults.get(key) in (None, ''): raise ValueError('缺少字段映射：' + key)
        book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False); cached = None
        try:
            ws = book[spec['sheet']]; ws.reset_dimensions(); iterator = ws.values
            if list(next(iterator, ())) != spec['headers']: raise ValueError('表头已变化')
            for n, cells in enumerate(iterator, 2):
                checkpoint(cancel)
                if not any(v not in (None, '') for v in cells): continue
                try:
                    d = {f: cells[c] if c is not None and c < len(cells) else defaults.get(f) for f in names for c in [mapping.get(f)]}
                    if any(isinstance(v, str) and v.startswith('=') for v in d.values()):
                        if cached is None: cached = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
                        values = next(cached[spec['sheet']].iter_rows(min_row=n, max_row=n, values_only=True))
                        for f, v in list(d.items()):
                            if isinstance(v, str) and v.startswith('='):
                                c = mapping[f]; d[f] = values[c] if c < len(values) else None
                                if d[f] is None: raise ValueError(f + '公式无缓存值，请在Excel计算保存后重试')
                    qty = base.number(d[QUANTITIES[stage]]); unit = str(d['数量单位'] or '').strip()
                    factor = base.number(d['折合库存单位系数'], True); basis = str(d['换算依据'] or '').strip()
                    if not unit: raise ValueError('数量单位不能为空')
                    if factor is not None and (Decimal(factor) <= 0 or not basis): raise ValueError('换算系数须大于0且提供依据')
                    if unit == '库存单位' and factor is not None and Decimal(factor) != 1: raise ValueError('库存单位的换算系数只能为1')
                    if unit != '库存单位' and factor is None: raise ValueError('非库存单位须提供明确换算系数及依据')
                    with localcontext() as ctx:
                        ctx.prec = max(28, len(Decimal(qty).as_tuple().digits) + len(Decimal(factor or 1).as_tuple().digits))
                        converted = base.number(Decimal(qty) * Decimal(factor or 1))
                    day = d[DOCDATE[stage]]
                    if isinstance(day, datetime): day = day.date().isoformat()
                    elif isinstance(day, date): day = day.isoformat()
                    elif day not in (None, ''): day = str(day).strip(); date.fromisoformat(day)
                    else: day = None
                    if stage != 'return' and day is None: raise ValueError(DOCDATE[stage] + '不能为空')
                    doc = base.code(d[DOCNUM[stage]], stage == 'return')
                    props = spec.get('properties', {}); signed = base.code(d['补货单号'], True)
                    rev = revision(d['签发修订'])
                    if props.get('修订'):
                        if props.get('补货单号') != signed: raise ValueError('单号与导出文件修订标识不一致')
                        embedded = revision(props['修订'])
                        if rev is not None and rev != embedded: raise ValueError('行修订与文件修订不一致')
                        rev = embedded
                    wh = str(d['供货仓'] or '').strip()
                    if wh and wh not in ('河北仓', '山东仓'): raise ValueError('供货仓须为河北仓或山东仓')
                    currency = str(d.get('币种') or '').strip()
                    if stage == 'payment' and not currency: raise ValueError('付款实际金额须明确币种')
                    rows.append({'sku': base.code(d['商品编码']), 'name': str(d['商品名称'] or ''), 'signed_number': signed, 'revision': rev, 'store': base.code(d['门店编码'], True), 'warehouse': wh, 'document_number': doc, 'document_date': day, 'quantity': qty, 'unit': unit, 'factor': factor, 'conversion_basis': basis, 'base_quantity': converted, 'amount': base.number(d.get('实际金额'), stage != 'payment'), 'currency': currency, 'note': d['备注'], 'source': {'sha': spec['sha'], 'file': p.name, 'sheet': ws.title, 'row': n}, 'raw_headers': spec['headers'], 'raw_values': list(cells)})
                except (ValueError, ArithmeticError) as exc: raise ValueError(f'{p.name} / {ws.title} 第{n}行：{exc}') from exc
                if n % 200 == 0: progress(RowProgress(f'读取{TITLES[stage]} {p.name} 第{n}行', i, len(specs)))
        finally:
            book.close()
            if cached: cached.close()
        if base.adapter.digest(p) != spec['sha']: raise ValueError('读取期间文件变化：' + p.name)
        progress(RowProgress(f'已校验文件 {i+1}/{len(specs)}', i+1, len(specs)))
    if not rows: raise ValueError('没有商品明细；整单明确为0时请填写数量0，不得用空文件表示')
    return {'stage': stage, 'rows': rows, 'files': list({f['sha']: f for f in files}.values())}

def template(stage, path):
    import os, tempfile
    from openpyxl.styles import Font
    notes = [('必需字段', '、'.join(required(stage))), ('关联', '补货单号为分仓子单号，签发修订为V001；缺少时必须明确选择，不猜最新版。'), ('范围', '每张仓单提交完整环节版本；不累计、不分摊。空数量拒绝，0保留；未列出的SKU不当作0。'), ('单位', '数量单位填库存单位；其它单位须填写折合库存单位系数及换算依据。'), ('金额', '付款金额填写源单各商品实际金额及币种，不使用模型价格估算。'), ('日期', '日期为YYYY-MM-DD；请提供已经付款/实际收货的业务单据。'), ('身份', '编码与单号按文本输入，保留前导零。每仓单每商品一行；重复同值去重，冲突拒绝。')]
    target = Path(path)
    if target.exists(): raise FileExistsError('目标文件已存在，不覆盖')
    book = openpyxl.Workbook(); ws = book.active; ws.title = TITLES[stage]; ws.append(fields(stage)); ws.freeze_panes = 'A2'
    for cell in ws[1]: cell.font = Font(bold=True); ws.column_dimensions[cell.column_letter].width = 24
    for n in range(2,202):
        for i, field in enumerate(fields(stage),1):
            if '编码' in field or '单号' in field or field == '签发修订': ws.cell(n,i).number_format = '@'
    help_sheet = book.create_sheet('填写说明'); help_sheet.append(['项目','说明'])
    for row in notes: help_sheet.append(row)
    help_sheet.column_dimensions['A'].width=18; help_sheet.column_dimensions['B'].width=110
    fd, tmp = tempfile.mkstemp(dir=target.parent,suffix='.xlsx'); os.close(fd)
    try: book.save(tmp); os.link(tmp,target)
    finally: book.close(); Path(tmp).unlink(missing_ok=True)
